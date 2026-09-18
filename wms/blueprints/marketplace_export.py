"""Выгрузка состава коробов перемещения в формате, который принимают
личные кабинеты маркетплейсов при заведении поставки:

- Ozon: "Состав грузовых мест" — короб WMS = одно грузовое место (ГМ).
  Ozon сам генерирует штрихкоды ГМ в своем кабинете (см. приложенный
  шаблон), поэтому их нужно вставить сюда самим, по одному на строку, в
  ТОМ ЖЕ порядке, что и короба в перемещении (см. movement.detail —
  порядок совпадает с MovementLine.id). "Артикул товара" Ozon — не наш
  внутренний SKU, а отдельная строка, которую нужно предварительно
  сопоставить со штрихкодом через OzonArticleMapping (см. ozon_mapping).
  Файл выгружается с ВТОРЫМ листом "Наши короба и ГМ" — не часть
  официального шаблона Ozon, просто сопоставление нашего короба
  (Box.box_number) и присвоенного ему грузоместа для собственного
  контроля (см. export_ozon_package_composition).
- Wildberries: проще во всем — "Баркод товара" не требует отдельного
  сопоставления (годится наш Nomenclature.barcode как есть), а "ШК
  короба" — это наш СОБСТВЕННЫЙ номер короба (Box.box_number), а не
  что-то выданное WB, поэтому запрашивать отдельный список не нужно,
  файл собирается и скачивается сразу (см. wb_package_composition).
- Ozon "заявка на поставку" (products-import-template): отдельно от
  состава по грузовым местам — одна строка на SKU с суммарным
  количеством по ВСЕМ коробам перемещения сразу (не по коробам
  отдельно), см. ozon_supply_request. Сумма количеств в обоих файлах
  должна совпадать, иначе Ozon аннулирует состав ГМ (см. инструкцию в
  самом шаблоне "Состав ГМ").
- WB "заявка на поставку" (см. чат): по аналогии с Ozon — второй,
  упрощенный файл всего с двумя колонками (Баркод, Количество), тоже
  суммарным количеством по ВСЕМ коробам перемещения сразу, см.
  wb_supply_request.

И там, и там "Срок годности" WMS не отслеживает — оставляется пустым,
заполняется вручную при необходимости, как и предусмотрено самими
шаблонами."""

import re

from flask import Blueprint, Response, abort, flash, redirect, render_template, request, url_for
from flask_login import current_user
from openpyxl import load_workbook

from ..extensions import db
from ..models import MovementDocument, MovementLine, OzonArticleMapping
from ..utils.excel_io import (
    export_ozon_package_composition,
    export_ozon_supply_request,
    export_wb_package_composition,
    export_wb_supply_request,
    timestamp_for_filename,
)
from ..utils.http import content_disposition
from .movement import _can_view_movement_document

bp = Blueprint("marketplace_export", __name__)


def _get_viewable_movement(doc_id):
    doc = MovementDocument.query.get_or_404(doc_id)
    if not _can_view_movement_document(doc):
        abort(404)
    return doc


def _cell_to_barcode(value):
    """Штрихкод в файле сопоставления приходит из Excel как число (Ozon
    отдает баркоды без ведущих нулей и без научной нотации) — приводим к
    обычной строке без ".0" на конце, а не str(float(...))."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


_BARCODE_RE = re.compile(r"^\d{6,20}$")


@bp.route("/ozon-mapping", methods=["GET", "POST"])
def ozon_mapping():
    if not current_user.is_admin:
        flash("Загружать сопоставление артикулов Ozon может только администратор", "danger")
        return redirect(url_for("main.index"))

    if request.method == "POST":
        file = request.files.get("file")
        if not file or file.filename == "":
            flash("Выберите файл xlsx", "danger")
            return redirect(url_for("marketplace_export.ozon_mapping"))

        workbook = load_workbook(file, read_only=True, data_only=True)
        worksheet = workbook.active
        updated = 0
        skipped_examples = []
        for row in worksheet.iter_rows(values_only=True):
            if not row or row[0] is None:
                continue
            barcode = _cell_to_barcode(row[0])
            article = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            # Первая колонка файла сопоставления — именно штрихкод (число из
            # 6+ цифр), а не название/артикул. Если это не похоже на
            # штрихкод — почти наверняка загружен не тот файл (например,
            # готовая "Заявка на поставку" вместо исходного списка
            # штрихкод-артикул) — не создаем мусорную запись, которая
            # никогда не найдется по реальному Nomenclature.barcode.
            if not barcode or not article or not _BARCODE_RE.match(barcode):
                if barcode and len(skipped_examples) < 5:
                    skipped_examples.append(barcode)
                continue
            mapping = OzonArticleMapping.query.filter_by(barcode=barcode).first()
            if mapping is None:
                mapping = OzonArticleMapping(barcode=barcode)
                db.session.add(mapping)
            mapping.article = article
            updated += 1
        db.session.commit()
        if updated:
            flash(f"Сопоставление артикулов Ozon обновлено: {updated} строк", "success")
        if skipped_examples:
            flash(
                "Пропущено строк, где первая колонка не похожа на штрихкод (6-20 цифр): "
                + str(len(skipped_examples)) + (" и другие" if len(skipped_examples) == 5 else "")
                + " — например: " + ", ".join(skipped_examples)
                + ". Убедитесь, что загружаете именно файл «штрихкод → артикул», "
                "а не готовую заявку на поставку.",
                "warning",
            )
        if not updated and not skipped_examples:
            flash("В файле не найдено ни одной строки со штрихкодом и артикулом", "danger")
        return redirect(url_for("marketplace_export.ozon_mapping"))

    return render_template(
        "marketplace_export/ozon_mapping.html",
        count=OzonArticleMapping.query.count(),
        preview=OzonArticleMapping.query.order_by(OzonArticleMapping.updated_at.desc()).limit(10).all(),
    )


@bp.route("/movement/<int:doc_id>/ozon", methods=["GET", "POST"])
def ozon_package_composition(doc_id):
    doc = _get_viewable_movement(doc_id)
    lines = doc.lines.order_by(MovementLine.id.asc()).all()

    if request.method == "POST":
        raw = request.form.get("cargo_barcodes", "")
        cargo_barcodes = [b.strip() for b in raw.splitlines() if b.strip()]
        if len(cargo_barcodes) != len(lines):
            flash(
                f"Штрихкодов ГМ должно быть ровно {len(lines)} (по числу коробов в "
                f"перемещении), а введено {len(cargo_barcodes)}",
                "danger",
            )
            return render_template(
                "marketplace_export/ozon_export.html", doc=doc, lines=lines, cargo_barcodes_raw=raw
            )

        rows = []
        unmapped = set()
        for line, cargo_barcode in zip(lines, cargo_barcodes):
            for item in line.box.items:
                barcode = item.nomenclature.barcode
                mapping = OzonArticleMapping.query.filter_by(barcode=barcode).first()
                if mapping is None:
                    unmapped.add(barcode)
                rows.append(
                    {
                        "barcode": barcode,
                        "article": mapping.article if mapping else "",
                        "qty": item.qty,
                        "cargo_barcode": cargo_barcode,
                        "box_number": line.box.box_number,
                        "name": item.nomenclature.name,
                    }
                )

        data = export_ozon_package_composition(rows)
        if unmapped:
            flash(
                "Не найден артикул Ozon для штрихкодов: " + ", ".join(sorted(unmapped))
                + " — колонка «Артикул товара» для них оставлена пустой, заполните вручную "
                "или дозагрузите сопоставление.",
                "warning",
            )
        fname = f"{doc.number}_ozon_gm_{timestamp_for_filename()}.xlsx"
        return Response(
            data,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": content_disposition(fname)},
        )

    return render_template(
        "marketplace_export/ozon_export.html", doc=doc, lines=lines, cargo_barcodes_raw=""
    )


@bp.route("/movement/<int:doc_id>/ozon/supply-request")
def ozon_supply_request(doc_id):
    """"Заявка на поставку" Ozon — в отличие от ozon_package_composition
    (по коробам), здесь одна строка на SKU с суммарным количеством по
    ВСЕМ коробам перемещения сразу; ШК ГМ тут не нужен, поэтому файл
    скачивается сразу, без промежуточной формы."""
    doc = _get_viewable_movement(doc_id)
    lines = doc.lines.order_by(MovementLine.id.asc()).all()

    totals = {}
    for line in lines:
        for item in line.box.items:
            entry = totals.setdefault(
                item.nomenclature_id,
                {"barcode": item.nomenclature.barcode, "name": item.nomenclature.name, "qty": 0.0},
            )
            entry["qty"] += item.qty

    rows = []
    unmapped = set()
    for entry in totals.values():
        mapping = OzonArticleMapping.query.filter_by(barcode=entry["barcode"]).first()
        if mapping is None:
            unmapped.add(entry["barcode"])
        rows.append({"article": mapping.article if mapping else "", "name": entry["name"], "qty": entry["qty"]})
    rows.sort(key=lambda r: r["name"])

    data = export_ozon_supply_request(rows)
    if unmapped:
        flash(
            "Не найден артикул Ozon для штрихкодов: " + ", ".join(sorted(unmapped))
            + " — колонка «артикул» для них оставлена пустой, заполните вручную "
            "или дозагрузите сопоставление.",
            "warning",
        )
    fname = f"{doc.number}_ozon_supply_request_{timestamp_for_filename()}.xlsx"
    return Response(
        data,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": content_disposition(fname)},
    )


@bp.route("/movement/<int:doc_id>/wb")
def wb_package_composition(doc_id):
    """В отличие от Ozon (там штрихкод ГМ генерирует сам маркетплейс),
    в "ШК короба" для WB вносится наш собственный номер короба
    (Box.box_number) — отдельный список запрашивать не нужно, файл
    скачивается сразу."""
    doc = _get_viewable_movement(doc_id)
    lines = doc.lines.order_by(MovementLine.id.asc()).all()

    rows = []
    for line in lines:
        for item in line.box.items:
            rows.append(
                {
                    "barcode": item.nomenclature.barcode,
                    "qty": item.qty,
                    "box_barcode": line.box.box_number,
                }
            )

    data = export_wb_package_composition(rows)
    fname = f"{doc.number}_wb_shk_{timestamp_for_filename()}.xlsx"
    return Response(
        data,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": content_disposition(fname)},
    )


@bp.route("/movement/<int:doc_id>/wb/supply-request")
def wb_supply_request(doc_id):
    """Второй, более простой файл для WB (см. чат) — всего два столбца,
    баркод и суммарное количество по ВСЕМ коробам перемещения сразу (как
    ozon_supply_request у Ozon), а не по коробам отдельно, как в
    wb_package_composition."""
    doc = _get_viewable_movement(doc_id)
    lines = doc.lines.order_by(MovementLine.id.asc()).all()

    totals = {}
    for line in lines:
        for item in line.box.items:
            totals[item.nomenclature.barcode] = totals.get(item.nomenclature.barcode, 0.0) + item.qty

    rows = sorted(
        ({"barcode": barcode, "qty": qty} for barcode, qty in totals.items()),
        key=lambda r: r["barcode"],
    )

    data = export_wb_supply_request(rows)
    fname = f"{doc.number}_wb_supply_request_{timestamp_for_filename()}.xlsx"
    return Response(
        data,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": content_disposition(fname)},
    )
