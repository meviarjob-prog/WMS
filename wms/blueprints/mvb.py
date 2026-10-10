"""«МВБ Логистика» — отдельный раздел на базе WMS со своим входом (/mvb/login).

Клиент оформляет заявку на передачу коробов (забор транспортной компанией
или самопривоз) для отправки на СЦ Wildberries / Ozon. При оформлении каждый
короб получает собственный штрихкод, клиент печатает этикетки 58×40 и
клеит их на короба. Дальше короба сканируются поштучно: водитель при
заборе, склад МВБ при приемке и при отгрузке на СЦ, — и клиент видит у себя
статус каждого короба.

Пользователи с ролями MVB_ROLES видят только этот раздел; пользователи WMS
(кроме администраторов) — наоборот, сюда не попадают.
"""

import math
import secrets
from datetime import date, datetime, timedelta

from flask import (
    Blueprint, Response, abort, flash, jsonify, redirect, render_template, request, session,
    url_for,
)
from flask_login import current_user, login_user, logout_user

from ..extensions import db
from ..models import (
    MVB_BOX_STATUS_LABELS, MVB_BOX_STATUS_ORDER, MVB_BOX_STATUSES, MVB_DELIVERY_METHODS,
    MVB_MARKETPLACES, MVB_ROLES, MvbCity, MvbDestination, MvbOrderLine, MVB_TRIP_STATUSES, MvbBox, MvbBoxEvent, MvbClient, MvbOrder,
    AppSetting, MVB_PRICE_KINDS, MvbPallet, MvbPickupZone, MvbPriceTier, MvbTrip, MvbTripStop, MvbVehicle, User,
)
from ..utils.http import content_disposition
from ..utils.labels_pdf import build_labels_batch_pdf, build_mvb_box_labels_pdf
from ..utils.numbering import next_number
from ..utils.timezone import MOSCOW_OFFSET

bp = Blueprint("mvb", __name__)

# Эндпоинты, доступные без входа (проверяется в require_login приложения).
# Ссылка для наемного водителя на СЦ (без учетной записи) — по токену рейса.
MVB_PUBLIC_ENDPOINTS = {
    "mvb.login", "mvb.register", "mvb.trip_public", "mvb.trip_public_arrive", "mvb.trip_public_stop",
}

MAX_BOXES_PER_ORDER = 500

# Режимы поштучного сканирования до склада: из каких статусов короб можно
# перевести в какой и каким ролям это разрешено. Администраторы могут всё.
# Дальше склада короба движутся в составе рейса (погрузка сканом, отправка
# и сдача на СЦ — см. раздел «Рейсы»).
SCAN_MODES = {
    "pickup": {
        "title": "Забор у клиента",
        "nav": "Скан забора",
        "to": "picked_up",
        "from": {"created"},
        # Забор у клиента сканирует водитель (в меню — только у него).
        "roles": {"mvb_driver", "mvb_admin"},
    },
    "receive": {
        "title": "Приемка на складе МВБ",
        "nav": "Приемка",
        "to": "received",
        # Приемка — для самопривоза и вернувшихся с СЦ; короба с забора
        # водитель сдает кнопкой «Короба сданы в МВБ» (скан забранного
        # короба тоже примет его, если водитель забыл нажать).
        "from": {"created", "picked_up", "not_delivered"},
        "roles": {"mvb_storekeeper", "mvb_admin"},
    },
}

BOX_TIMESTAMP_FIELDS = {
    "picked_up": "picked_up_at",
    "received": "received_at",
    "loaded": "loaded_at",
    "shipped": "shipped_at",
    "delivered": "delivered_at",
}


# ---------- доступ ----------


def _is_client():
    return current_user.role == "mvb_client" and not current_user.is_admin


def _is_operator():
    return current_user.is_admin or current_user.role in ("mvb_staff", "mvb_admin")


def _is_storekeeper():
    return current_user.is_admin or current_user.role in ("mvb_storekeeper", "mvb_admin")


def _is_staff():
    """Сотрудник МВБ: оператор или кладовщик."""
    return _is_operator() or _is_storekeeper()


def _is_driver():
    return current_user.role == "mvb_driver" and not current_user.is_admin


def _can_scan(mode):
    return current_user.is_admin or current_user.role in SCAN_MODES[mode]["roles"]


def _visible_orders_query():
    query = MvbOrder.query
    if _is_client():
        query = query.filter(MvbOrder.client_id == current_user.mvb_client_id)
    return query


def _orders_role_ok(read_only=False):
    """Заявки (список, создание, правка) — клиент и оператор; карточку и
    этикетки заявки смотрит еще кладовщик. У водителя свои экраны забора."""
    return _is_client() or _is_operator() or (read_only and _is_storekeeper())


def _deny_orders():
    flash("Этот раздел вам не доступен", "danger")
    return redirect(url_for("mvb.index"))


def _get_order_or_404(order_id):
    order = _visible_orders_query().filter(MvbOrder.id == order_id).first()
    if order is None:
        abort(404)
    return order


@bp.before_request
def _restrict_to_mvb_users():
    if request.endpoint in MVB_PUBLIC_ENDPOINTS or not current_user.is_authenticated:
        return None
    if not (current_user.is_admin or current_user.role in MVB_ROLES):
        flash("Раздел «МВБ Логистика» вам не доступен", "danger")
        return redirect(url_for("main.index"))
    if _is_client() and current_user.mvb_client_id and _client_login_block(current_user):
        message = _client_login_block(current_user)
        logout_user()
        flash(message, "warning")
        return redirect(url_for("mvb.login"))
    if _is_client() and not current_user.mvb_client_id:
        logout_user()
        flash("Учетная запись не привязана к клиенту — обратитесь к администратору МВБ", "danger")
        return redirect(url_for("mvb.login"))
    return None


@bp.context_processor
def _inject():
    return {
        "MVB_BOX_STATUSES": MVB_BOX_STATUSES,
        "MVB_BOX_STATUS_LABELS": MVB_BOX_STATUS_LABELS,
        "MVB_MARKETPLACES": MVB_MARKETPLACES,
        "MVB_DELIVERY_METHODS": MVB_DELIVERY_METHODS,
        "MVB_ROLES": MVB_ROLES,
        "SCAN_MODES": SCAN_MODES,
        "mvb_is_client": current_user.is_authenticated and _is_client(),
        "mvb_can_scan": (lambda mode: current_user.is_authenticated and _can_scan(mode)),
        "mvb_is_staff": current_user.is_authenticated and _is_staff(),
        "mvb_is_operator": current_user.is_authenticated and _is_operator(),
        "mvb_is_storekeeper": current_user.is_authenticated and _is_storekeeper(),
        "mvb_is_driver": current_user.is_authenticated and _is_driver(),
        "mvb_can_manage": current_user.is_authenticated and current_user.can_manage_mvb(),
        "mvb_pending_clients": (
            MvbClient.query.filter_by(approval="pending").count()
            if current_user.is_authenticated and _is_operator() else 0
        ),
        "MVB_TRIP_STATUSES": MVB_TRIP_STATUSES,
        "mvb_destinations": MvbDestination.active() if current_user.is_authenticated else [],
        "mvb_pickup_zones": MvbPickupZone.active() if current_user.is_authenticated else [],
    }


# ---------- вход ----------


@bp.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated and (current_user.is_admin or current_user.role in MVB_ROLES):
        return redirect(url_for("mvb.index"))

    if request.method == "GET":
        return render_template("mvb/login.html")

    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")
    user = User.query.filter_by(username=username).first()
    if (
        not user
        or not user.is_active_user
        or not user.check_password(password)
        or not (user.is_admin or user.role in MVB_ROLES)
    ):
        flash("Неверный логин или пароль", "danger")
        return render_template("mvb/login.html", username=username)

    blocked = _client_login_block(user)
    if blocked:
        flash(blocked, "warning")
        return render_template("mvb/login.html", username=username)

    login_user(user, remember=True)
    session["session_version"] = user.session_version or 0
    return redirect(url_for("mvb.index"))


def _client_login_block(user):
    """Почему клиент пока не может войти (регистрация не подтверждена,
    отклонена или клиент отключен); None — может."""
    if user.role != "mvb_client" or user.is_admin or user.mvb_client is None:
        return None
    client = user.mvb_client
    if client.approval == "pending":
        return "Регистрация на проверке у оператора МВБ — войти можно после подтверждения"
    if client.approval == "rejected":
        return "Регистрация отклонена — свяжитесь с МВБ Логистика"
    if not client.is_active:
        return "Учетная запись клиента отключена — свяжитесь с МВБ Логистика"
    return None


@bp.route("/register", methods=["GET", "POST"])
def register():
    """Самостоятельная регистрация клиента (название, адрес забора, телефон,
    логин, пароль); клиент ждет
    подтверждения оператора (approval=pending), до этого войти нельзя."""
    if current_user.is_authenticated and (current_user.is_admin or current_user.role in MVB_ROLES):
        return redirect(url_for("mvb.index"))
    form = {k: request.form.get(k, "").strip() for k in ("name", "address", "phone", "username")}
    if request.method == "GET":
        return render_template("mvb/register.html", form=form)

    password = request.form.get("password", "")
    errors = []
    if not form["name"]:
        errors.append("Укажите название компании")
    if not form["address"]:
        errors.append("Укажите адрес забора")
    if sum(ch.isdigit() for ch in form["phone"]) < 10:
        errors.append("Укажите телефон")
    if not form["username"]:
        errors.append("Придумайте логин")
    elif User.query.filter(db.func.lower(User.username) == form["username"].lower()).first():
        errors.append("Такой логин уже занят")
    if len(password) < 6:
        errors.append("Пароль — не короче 6 символов")
    if errors:
        for error in errors:
            flash(error, "danger")
        return render_template("mvb/register.html", form=form)

    client = MvbClient(name=form["name"], phone=form["phone"], address=form["address"], approval="pending")
    user = User(username=form["username"], full_name=form["name"], role="mvb_client", mvb_client=client)
    user.set_password(password)
    db.session.add_all([client, user])
    db.session.commit()
    flash("Заявка на регистрацию отправлена. Оператор МВБ проверит данные и подтвердит — после этого войдите со своим логином.", "success")
    return redirect(url_for("mvb.login"))


@bp.route("/logout", methods=["POST"])
def logout():
    logout_user()
    flash("Вы вышли из системы", "success")
    return redirect(url_for("mvb.login"))


@bp.route("/")
def index():
    if current_user.role == "mvb_driver" and not current_user.is_admin:
        return redirect(url_for("mvb.driver"))
    if current_user.role == "mvb_storekeeper" and not current_user.is_admin:
        return redirect(url_for("mvb.scan", mode="receive"))
    return redirect(url_for("mvb.orders"))


# ---------- заявки ----------


@bp.route("/orders")
def orders():
    if not _orders_role_ok(read_only=False):
        return _deny_orders()
    query = _visible_orders_query()
    status = request.args.get("status", "")
    if status in ("draft", "confirmed", "cancelled"):
        query = query.filter(MvbOrder.status == status)
    client_id = request.args.get("client_id", type=int)
    if client_id and not _is_client():
        query = query.filter(MvbOrder.client_id == client_id)
    need_driver_filter = (
        MvbOrder.status == "confirmed", MvbOrder.delivery_method == "pickup", MvbOrder.driver_id.is_(None),
        MvbOrder.boxes.any(MvbBox.status == "created"),
    )
    if status == "need_driver":
        query = query.filter(*need_driver_filter)
    items = query.order_by(MvbOrder.created_at.desc()).limit(300).all()
    clients = [] if _is_client() else MvbClient.query.order_by(MvbClient.name).all()
    need_driver = 0 if _is_client() else MvbOrder.query.filter(*need_driver_filter).count()
    # Раздел "Из WMS" был отдельной страницей — теперь кандидаты на передачу в
    # МВБ показываются прямо здесь, в "Заявках в работе" (см. чат), отдельная
    # страница/пункт меню убраны.
    wms_rows = _wms_candidate_rows() if _is_operator() else []
    return render_template(
        "mvb/orders.html", orders=items, clients=clients, status=status, client_id=client_id,
        need_driver=need_driver, drivers=_active_drivers() if _is_staff() else [],
        wms_rows=wms_rows,
    )


def _parse_time(value):
    value = (value or "").strip()
    if not value:
        return None
    try:
        datetime.strptime(value, "%H:%M")
    except ValueError:
        return None
    return value


MAX_DIRECTIONS = 10


def _form_lines():
    """Направления из формы: списки line_marketplace / line_destination /
    line_slot / line_boxes (по строке на направление). Форма в одно
    направление (marketplace, destination, slot_date, box_count) тоже
    принимается. Возвращает (список словарей, ошибка)."""
    form = request.form
    if form.getlist("line_marketplace"):
        raw = list(zip(
            form.getlist("line_marketplace"), form.getlist("line_destination"),
            form.getlist("line_slot"), form.getlist("line_boxes"),
        ))
        count = form.get("direction_count", type=int) or len(raw)
        raw = raw[:max(1, min(count, MAX_DIRECTIONS))]
    else:
        raw = [(form.get("marketplace", "wb"), form.get("destination", ""), form.get("slot_date", ""),
                form.get("box_count", ""))]
    lines = []
    cities = None
    for n, (marketplace, destination, slot, boxes) in enumerate(raw, start=1):
        prefix = f"Направление {n}: " if len(raw) > 1 else ""
        if marketplace not in MVB_MARKETPLACES:
            return None, prefix + "выберите маркетплейс"
        destination = (destination or "").strip()
        if not destination:
            return None, prefix + ("выберите фулфилмент" if marketplace == "ff" else "выберите город (СЦ)")
        if cities is None:
            cities = MvbDestination.active()
        if cities:
            # Пункты назначения заданы в прайсе — выбираем только из них.
            match = next(
                (d for d in cities if d.marketplace == marketplace and d.value.lower() == destination.lower()), None,
            )
            if match is None:
                return None, prefix + f"«{destination}» нет в списке пунктов назначения ({MVB_MARKETPLACES[marketplace]})"
            destination = match.value
        try:
            slot_date = date.fromisoformat((slot or "").strip())
        except ValueError:
            return None, prefix + "укажите дату слота на СЦ"
        try:
            box_count = int(boxes)
        except (TypeError, ValueError):
            return None, prefix + "укажите количество коробов"
        if box_count < 1:
            return None, prefix + "количество коробов — от 1"
        lines.append({"marketplace": marketplace, "destination": destination, "slot_date": slot_date,
                      "box_count": box_count})
    if sum(l["box_count"] for l in lines) > MAX_BOXES_PER_ORDER:
        return None, f"Всего коробов в заявке — не больше {MAX_BOXES_PER_ORDER}"
    return lines, None


def _form_lines_checked():
    lines, error = _form_lines()
    if error:
        error = error[0].upper() + error[1:]
    return lines, error


def _fill_order_from_form(order):
    """Заполняет заявку (черновик) из формы; возвращает текст ошибки или None."""
    delivery_method = request.form.get("delivery_method", "pickup")
    if delivery_method not in MVB_DELIVERY_METHODS:
        return "Выберите способ передачи коробов"
    lines, error = _form_lines_checked()
    if error:
        return error
    planned_date = None
    raw_date = request.form.get("planned_date", "").strip()
    if raw_date:
        try:
            planned_date = date.fromisoformat(raw_date)
        except ValueError:
            return "Некорректная дата"
    pickup_address = request.form.get("pickup_address", "").strip()
    if delivery_method == "pickup" and not pickup_address:
        return "Для забора укажите адрес"
    zone = None
    if delivery_method == "pickup" and MvbPickupZone.active():
        zone = db.session.get(MvbPickupZone, request.form.get("pickup_zone_id", type=int) or 0)
        if zone is None or not zone.is_active:
            return "Выберите зону забора"

    # Черновик: коробов еще нет — направления просто пересоздаем.
    order.lines = [MvbOrderLine(seq=n, **data) for n, data in enumerate(lines, start=1)]
    order.sync_from_lines()
    order.delivery_method = delivery_method
    order.pickup_address = pickup_address or None
    order.pickup_zone = zone
    order.planned_date = planned_date
    order.time_from = _parse_time(request.form.get("time_from"))
    order.time_to = _parse_time(request.form.get("time_to"))
    order.comment = request.form.get("comment", "").strip() or None
    return None


@bp.route("/orders/new", methods=["GET", "POST"])
def order_new():
    if not _orders_role_ok(read_only=False):
        return _deny_orders()
    clients = [] if _is_client() else MvbClient.query.filter_by(is_active=True).order_by(MvbClient.name).all()
    if request.method == "GET":
        client = current_user.mvb_client if _is_client() else None
        return render_template(
            "mvb/order_form.html", order=None, clients=clients,
            default_address=(client.address if client else ""), max_directions=MAX_DIRECTIONS,
        )

    if _is_client():
        client_id = current_user.mvb_client_id
    else:
        client_id = request.form.get("client_id", type=int)
        if not client_id or not db.session.get(MvbClient, client_id):
            flash("Выберите клиента", "danger")
            return render_template("mvb/order_form.html", order=None, clients=clients, form=request.form,
                                   max_directions=MAX_DIRECTIONS)

    order = MvbOrder(client_id=client_id, created_by_id=current_user.id, status="draft")
    error = _fill_order_from_form(order)
    if error:
        flash(error, "danger")
        return render_template("mvb/order_form.html", order=None, clients=clients, form=request.form,
                                   max_directions=MAX_DIRECTIONS)
    order.number = next_number("mvb_order", "MVB-", 6)
    db.session.add(order)
    db.session.commit()
    flash(f"Заявка {order.number} создана. Проверьте и нажмите «Оформить» — коробам будут присвоены штрихкоды.", "success")
    return redirect(url_for("mvb.order_detail", order_id=order.id))


@bp.route("/orders/<int:order_id>")
def order_detail(order_id):
    if not _orders_role_ok(read_only=True):
        return _deny_orders()
    order = _get_order_or_404(order_id)
    drivers = _active_drivers() if _is_staff() else []
    return render_template(
        "mvb/order_detail.html", order=order, counts=order.status_counts(), drivers=drivers,
    )


@bp.route("/orders/<int:order_id>/slot", methods=["POST"])
def order_slot(order_id):
    """Дата слота на СЦ меняется и у оформленной заявки (слот могут
    перенести), пока короба не отправлены на СЦ."""
    if not _orders_role_ok(read_only=False):
        return _deny_orders()
    order = _get_order_or_404(order_id)
    line = next((l for l in order.lines if l.id == request.form.get("line_id", type=int)), None)
    if line is None and len(order.lines) == 1:
        line = order.lines[0]
    if line is None:
        abort(400)
    if order.status == "cancelled" or any(b.status in ("shipped", "delivered") for b in line.boxes):
        flash("Слот уже не изменить — короба отправлены на СЦ", "warning")
        return redirect(url_for("mvb.order_detail", order_id=order.id))
    try:
        line.slot_date = date.fromisoformat(request.form.get("slot_date", "").strip())
    except ValueError:
        flash("Укажите дату слота на СЦ", "danger")
        return redirect(url_for("mvb.order_detail", order_id=order.id))
    order.sync_from_lines()
    db.session.commit()
    flash(f"{line.short_label()}: слот на СЦ {line.slot_date.strftime('%d.%m.%Y')}", "success")
    return redirect(url_for("mvb.order_detail", order_id=order.id))


@bp.route("/orders/<int:order_id>/edit", methods=["GET", "POST"])
def order_edit(order_id):
    if not _orders_role_ok(read_only=False):
        return _deny_orders()
    order = _get_order_or_404(order_id)
    if order.status != "draft":
        flash("Изменить можно только черновик", "warning")
        return redirect(url_for("mvb.order_detail", order_id=order.id))
    if request.method == "GET":
        return render_template("mvb/order_form.html", order=order, clients=[], max_directions=MAX_DIRECTIONS)
    error = _fill_order_from_form(order)
    if error:
        db.session.rollback()
        flash(error, "danger")
        return render_template("mvb/order_form.html", order=order, clients=[], form=request.form,
                               max_directions=MAX_DIRECTIONS)
    db.session.commit()
    flash("Заявка сохранена", "success")
    return redirect(url_for("mvb.order_detail", order_id=order.id))


@bp.route("/orders/<int:order_id>/confirm", methods=["POST"])
def order_confirm(order_id):
    """Оформление заявки: способ передачи фиксируется, каждому коробу
    присваивается собственный штрихкод «номер заявки-порядковый номер»."""
    if not _orders_role_ok(read_only=False):
        return _deny_orders()
    order = _get_order_or_404(order_id)
    if order.status != "draft":
        flash("Заявка уже оформлена", "warning")
        return redirect(url_for("mvb.order_detail", order_id=order.id))
    # У каждого направления свой счетчик коробов: при нескольких
    # направлениях штрихкод «номер заявки-направление-номер короба».
    seq = 0
    multi = len(order.lines) > 1
    for line in order.lines:
        for n in range(1, line.box_count + 1):
            seq += 1
            barcode = f"{order.number}-{line.seq}-{n:03d}" if multi else f"{order.number}-{n:03d}"
            order.boxes.append(MvbBox(seq=seq, barcode=barcode, status="created", line=line))
    order.status = "confirmed"
    order.confirmed_at = datetime.utcnow()
    _apply_prices(order)
    db.session.commit()
    flash(f"Заявка оформлена: присвоено штрихкодов — {order.box_count}. Распечатайте этикетки и наклейте на каждый короб.", "success")
    return redirect(url_for("mvb.order_detail", order_id=order.id))


@bp.route("/orders/<int:order_id>/cancel", methods=["POST"])
def order_cancel(order_id):
    if not _orders_role_ok(read_only=False):
        return _deny_orders()
    order = _get_order_or_404(order_id)
    if order.status == "cancelled":
        return redirect(url_for("mvb.order_detail", order_id=order.id))
    if any(box.status != "created" for box in order.boxes):
        flash("Нельзя отменить: часть коробов уже отсканирована", "danger")
        return redirect(url_for("mvb.order_detail", order_id=order.id))
    order.status = "cancelled"
    db.session.commit()
    flash("Заявка отменена", "success")
    return redirect(url_for("mvb.order_detail", order_id=order.id))


@bp.route("/orders/<int:order_id>/delete", methods=["POST"])
def order_delete(order_id):
    """Администратор удаляет заявку в любом статусе вместе с коробами и их
    историей: короба пропадают с паллет и из рейсов."""
    if not _require_manage():
        return redirect(url_for("mvb.order_detail", order_id=order_id))
    order = db.session.get(MvbOrder, order_id) or abort(404)
    number = order.number
    others = _orders_on_pallets({b.pallet_id for b in order.boxes}) - {order}
    db.session.delete(order)
    db.session.flush()
    _recalc_pallet_costs(others)
    db.session.commit()
    flash(f"Заявка {number} удалена", "success")
    return redirect(url_for("mvb.orders"))


@bp.route("/orders/<int:order_id>/labels.pdf")
def order_labels_pdf(order_id):
    """Этикетки 58×40 на все короба заявки, на одно направление (?line=id)
    или на выбранные короба (?seq=1,2)."""
    if not _orders_role_ok(read_only=True):
        return _deny_orders()
    order = _get_order_or_404(order_id)
    if order.status != "confirmed" or not order.boxes:
        abort(404)
    boxes = order.boxes
    seq_param = request.args.get("seq", "")
    if seq_param:
        try:
            wanted = {int(v) for v in seq_param.split(",") if v.strip()}
        except ValueError:
            abort(400)
        boxes = [b for b in boxes if b.seq in wanted]
    line_id = request.args.get("line", type=int)
    if line_id:
        boxes = [b for b in boxes if b.line_id == line_id]
        if not boxes:
            abort(404)
    # Номер короба на этикетке — внутри своего направления («3 / 12 кор.»):
    # так его считают на СЦ.
    position = {}
    for line in order.lines:
        for n, line_box in enumerate(line.boxes, start=1):
            position[line_box.id] = (n, len(line.boxes))
    entries = []
    for box in boxes:
        line = box.line
        n, total = position.get(box.id, (box.seq, len(order.boxes)))
        entries.append({
            # Коротко, как пишут на коробах: «WB Коледино», «OZON Хоругвино».
            "barcode": box.barcode, "destination": line.short_label(), "seq": n, "total": total,
            "slot": line.slot_date.strftime("%d.%m.%Y") if line.slot_date else "",
            "sender": order.client.name,
        })
    pdf = build_mvb_box_labels_pdf(entries)
    return Response(
        pdf,
        mimetype="application/pdf",
        headers={"Content-Disposition": content_disposition(f"{order.number}.pdf", "inline")},
    )


# ---------- сканирование ----------


# Заявка считается «новой» для водителя столько времени после оформления.
NEW_ORDER_HOURS = 3


def _driver_vehicle_and_load(user=None):
    """Авто водителя и сколько коробов сейчас у него в машине (забраны им,
    но еще не приняты на складе)."""
    user = user or current_user
    vehicle = (
        MvbVehicle.query.filter_by(driver_id=user.id, is_active=True)
        .order_by(MvbVehicle.id).first()
    )
    load = MvbBox.query.filter(MvbBox.status == "picked_up", MvbBox.picked_up_by_id == user.id).count()
    return vehicle, load


def _driver_active_trips(user=None):
    """Рейсы на СЦ, назначенные водителю и еще не завершенные."""
    user = user or current_user
    return (
        MvbTrip.query.filter(
            MvbTrip.driver_id == user.id,
            MvbTrip.status.in_(["assigned", "arrived", "loading", "departed"]),
        )
        .order_by(MvbTrip.planned_arrival_at)
        .all()
    )


@bp.route("/driver")
def driver():
    """Экран водителя на маршруте: сколько места в машине и лента заявок на
    забор (назначенные ему и свободные) с отметкой, влезает ли заявка;
    новые заявки подсвечиваются, страница сама обновляется."""
    candidates = (
        MvbOrder.query.filter(MvbOrder.status == "confirmed", MvbOrder.delivery_method == "pickup")
        .order_by(MvbOrder.confirmed_at.desc())
        .all()
    )
    vehicle, load = (None, 0)
    trips = []
    if _is_driver():
        vehicle, load = _driver_vehicle_and_load()
        trips = _driver_active_trips()
        # Водитель в рейсе на СЦ чужие заявки «по дороге» не берет — видит
        # только то, что оператор назначил ему самому.
        allowed = (current_user.id,) if trips else (None, current_user.id)
        candidates = [o for o in candidates if o.driver_id in allowed]
    free = (vehicle.capacity_boxes - load) if vehicle and vehicle.capacity_boxes else None
    now = datetime.utcnow()
    rows = []
    for order in candidates:
        remaining = sum(1 for b in order.boxes if b.status == "created")
        if not remaining:
            continue
        rows.append({
            "order": order,
            "remaining": remaining,
            "fits": None if free is None else remaining <= free,
            "is_new": bool(order.confirmed_at and (now - order.confirmed_at).total_seconds() < NEW_ORDER_HOURS * 3600),
            "mine": _is_driver() and order.driver_id == current_user.id,
        })
    # Сначала свои, затем новые, затем остальные по дате забора.
    rows.sort(key=lambda r: (not r["mine"], not r["is_new"], r["order"].planned_date or date.max))
    return render_template(
        "mvb/driver.html", rows=rows, trips=trips, on_trip=bool(trips), vehicle=vehicle, load=load, free=free,
        max_order_id=max((r["order"].id for r in rows), default=0),
    )


@bp.route("/driver/orders/<int:order_id>/take", methods=["POST"])
def driver_take(order_id):
    """Водитель берет свободную заявку на забор «по дороге»."""
    if not _is_driver():
        abort(403)
    order = MvbOrder.query.get_or_404(order_id)
    if order.status != "confirmed" or order.delivery_method != "pickup":
        abort(400)
    if order.driver_id not in (None, current_user.id):
        flash(f"Заявку {order.number} уже взял другой водитель", "warning")
        return redirect(url_for("mvb.driver"))
    if order.driver_id is None and _driver_active_trips():
        flash("Вы в рейсе на СЦ — заявки на забор назначает оператор", "warning")
        return redirect(url_for("mvb.driver"))
    vehicle, load = _driver_vehicle_and_load()
    remaining = sum(1 for b in order.boxes if b.status == "created")
    if vehicle and vehicle.capacity_boxes and remaining > vehicle.capacity_boxes - load:
        flash(f"Внимание: в машине свободно {max(vehicle.capacity_boxes - load, 0)} мест, а в заявке {remaining} кор.", "warning")
    order.driver_id = current_user.id
    db.session.commit()
    flash(f"Заявка {order.number} ваша — {order.pickup_address or 'адрес не указан'}", "success")
    return redirect(url_for("mvb.driver"))


@bp.route("/driver/orders/<int:order_id>/done", methods=["POST"])
def driver_pickup_done(order_id):
    """«Готово» после скана коробов на заборе: завершить можно, только
    когда отсканированы все короба заявки."""
    if not _can_scan("pickup"):
        abort(403)
    order = MvbOrder.query.get_or_404(order_id)
    if order.status != "confirmed" or order.delivery_method != "pickup":
        abort(400)
    left = [b for b in order.boxes if b.status == "created"]
    if left and request.form.get("partial"):
        # Все короба не влезли в машину: водитель забрал часть, остаток
        # возвращается в ленту забора (свободная заявка — для другого
        # водителя или для него же следующим рейсом).
        taken = len(order.boxes) - len(left)
        if not taken:
            flash("Сначала отсканируйте короба, которые забираете", "danger")
            return redirect(url_for("mvb.scan", mode="pickup", order=order.id))
        if order.driver_id == current_user.id or not _is_operator():
            order.driver_id = None
        db.session.commit()
        flash(
            f"Забрано {taken} из {len(order.boxes)} кор. у {order.client.name}. "
            f"Остаток {len(left)} кор. вернулся в заявки на забор.", "warning",
        )
        return redirect(url_for("mvb.driver"))
    if left:
        flash(
            f"Не все короба отсканированы: {len(order.boxes) - len(left)} из {len(order.boxes)}. "
            f"Осталось: {', '.join(b.barcode for b in left[:10])}{' …' if len(left) > 10 else ''}",
            "danger",
        )
        return redirect(url_for("mvb.scan", mode="pickup", order=order.id))
    order.pickup_done_at = order.pickup_done_at or datetime.utcnow()
    db.session.commit()
    flash(f"Забор {order.client.name} завершен: {len(order.boxes)} кор.", "success")
    return redirect(url_for("mvb.driver"))


@bp.route("/driver/handover", methods=["POST"])
def driver_handover():
    """«Короба сданы в МВБ»: водитель привез забранные короба на склад —
    все они переходят в «На складе МВБ» без поштучной приемки."""
    if not _is_driver():
        abort(403)
    boxes = MvbBox.query.filter(MvbBox.status == "picked_up", MvbBox.picked_up_by_id == current_user.id).all()
    if not boxes:
        flash("У вас нет забранных коробов", "warning")
        return redirect(url_for("mvb.driver"))
    now = datetime.utcnow()
    for box in boxes:
        _move_box(box, "received", now)
    db.session.commit()
    orders = sorted({box.order.number for box in boxes})
    flash(f"Сдано на склад МВБ: {len(boxes)} кор. ({', '.join(orders)})", "success")
    return redirect(url_for("mvb.driver"))


@bp.route("/driver/orders/<int:order_id>/release", methods=["POST"])
def driver_release(order_id):
    if not _is_driver():
        abort(403)
    order = MvbOrder.query.get_or_404(order_id)
    if order.driver_id == current_user.id and not any(b.status != "created" for b in order.boxes):
        order.driver_id = None
        db.session.commit()
        flash(f"Вы отказались от заявки {order.number}", "success")
    return redirect(url_for("mvb.driver"))


@bp.route("/scan/<mode>")
def scan(mode):
    if mode not in SCAN_MODES:
        abort(404)
    if not _can_scan(mode):
        flash("Этот режим сканирования вам не доступен", "danger")
        return redirect(url_for("mvb.index"))
    current = None
    if mode == "pickup" and request.args.get("order", type=int):
        current = MvbOrder.query.get(request.args.get("order", type=int))
    return render_template(
        "mvb/scan.html", mode=mode, mode_info=SCAN_MODES[mode], current=current,
        current_progress=_order_progress(current, "picked_up") if current else None,
    )


@bp.route("/scan/<mode>", methods=["POST"])
def scan_box(mode):
    """Скан одного короба: JSON {ok, message, box...}. Повторный скан уже
    переведенного короба не ошибка — просто сообщаем, что он уже учтен."""
    if mode not in SCAN_MODES:
        abort(404)
    if not _can_scan(mode):
        return jsonify(ok=False, message="Нет доступа к этому режиму"), 403
    info = SCAN_MODES[mode]
    code = (request.form.get("barcode") or (request.get_json(silent=True) or {}).get("barcode") or "").strip()
    if not code:
        return jsonify(ok=False, message="Пустой штрихкод"), 400

    box = _find_box(code)
    imported_note = None
    if box is None and mode in ("receive", "pickup"):
        # Свой короб WMS, перемещение которого еще не передано в МВБ, —
        # передаем всё перемещение сразу по скану его этикетки.
        doc = _wms_movement_for_box_code(code)
        if doc is not None:
            order, errors = _import_movement(doc, "self" if mode == "receive" else "pickup")
            if order is not None:
                imported_note = f"Перемещение WMS {doc.number} принято в МВБ как заявка {order.number}"
                box = _find_box(code)
    if box is None:
        return jsonify(ok=False, message=f"Короб {code} не найден"), 404
    order = box.order
    payload = {
        "barcode": box.barcode,
        "order": order.number,
        "order_url": url_for("mvb.order_detail", order_id=order.id),
        "client": order.client.name,
        "seq": box.line_position,
        "total": box.line_total,
        "direction": box.line.short_label() if box.line else "",
    }
    if order.status != "confirmed":
        return jsonify(ok=False, message=f"Заявка {order.number} не оформлена или отменена", **payload), 409
    if mode == "pickup" and order.delivery_method != "pickup":
        return jsonify(ok=False, message="Это самопривоз — короб принимается на складе", **payload), 409

    payload["order_id"] = order.id
    if box.status == info["to"]:
        return jsonify(ok=True, already=True, message=f"Уже отмечен: {box.status_label}",
                       status=box.status_label, **payload, **_order_progress(order, info["to"]))
    if box.status not in info["from"]:
        return jsonify(ok=False, message=f"Сейчас короб в статусе «{box.status_label}»", **payload), 409

    now = datetime.utcnow()
    returned = box.status == "not_delivered"
    _move_box(box, info["to"], now)
    if returned:
        # Вернулся с СЦ — снова доступен для погрузки в новый рейс.
        box.trip_id = box.trip_stop_id = box.pallet_id = None
        box.loaded_at = box.shipped_at = None
    load_note = None
    if mode == "pickup":
        box.picked_up_by_id = current_user.id
        if _is_driver() and order.driver_id is None:
            order.driver_id = current_user.id
    db.session.commit()
    if mode == "pickup" and _is_driver():
        vehicle, load = _driver_vehicle_and_load()
        if vehicle and vehicle.capacity_boxes:
            load_note = f"В машине {load} из {vehicle.capacity_boxes} кор."
    done = sum(
        1 for b in (box.line.boxes if box.line else order.boxes)
        if MVB_BOX_STATUS_ORDER.get(b.status, 0) >= MVB_BOX_STATUS_ORDER[info["to"]]
    )
    message = f"{box.barcode}: {box.status_label}"
    if imported_note:
        message = f"{imported_note}. {message}"
    return jsonify(
        ok=True, already=False, message=message, warning=load_note,
        status=box.status_label, done=done, **payload, **_order_progress(order, info["to"]),
    )


def _order_progress(order, status):
    """Сколько коробов заявки уже дошло до этапа status (для кнопки
    «Готово» при заборе)."""
    level = MVB_BOX_STATUS_ORDER[status]
    return {
        "order_done": sum(1 for b in order.boxes if MVB_BOX_STATUS_ORDER.get(b.status, 0) >= level),
        "order_total": len(order.boxes),
    }


# ---------- администрирование ----------


def _require_manage():
    if not current_user.can_manage_mvb():
        flash("Доступно только администратору МВБ", "danger")
        return False
    return True


@bp.route("/registrations")
def registrations():
    """Новые регистрации теперь в разделе «Клиенты»."""
    return redirect(url_for("mvb.admin_clients"))


@bp.route("/registrations/<int:client_id>/<action>", methods=["POST"])
def registration_action(client_id, action):
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    client = db.session.get(MvbClient, client_id)
    if client is None or action not in ("approve", "reject"):
        abort(404)
    client.approval = "approved" if action == "approve" else "rejected"
    client.approved_at = datetime.utcnow()
    client.approved_by_id = current_user.id
    db.session.commit()
    if action == "approve":
        flash(f"Клиент «{client.name}» подтвержден — может входить и создавать заявки", "success")
    else:
        flash(f"Регистрация «{client.name}» отклонена", "warning")
    return redirect(url_for("mvb.admin_clients"))


@bp.route("/admin/clients", methods=["GET", "POST"])
def admin_clients():
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    if request.method == "POST":
        client_id = request.form.get("client_id", type=int)
        client = db.session.get(MvbClient, client_id) if client_id else MvbClient()
        if client is None:
            abort(404)
        name = request.form.get("name", "").strip()
        if not name:
            flash("Укажите название клиента", "danger")
            return redirect(url_for("mvb.admin_clients"))
        client.name = name
        client.inn = request.form.get("inn", "").strip() or None
        client.contact_name = request.form.get("contact_name", "").strip() or None
        client.phone = request.form.get("phone", "").strip() or None
        client.address = request.form.get("address", "").strip() or None
        if client_id:
            client.is_active = request.form.get("is_active") == "1"
        db.session.add(client)
        db.session.commit()
        flash(f"Клиент «{client.name}» сохранен", "success")
        return redirect(url_for("mvb.admin_clients"))
    clients = MvbClient.query.filter(MvbClient.approval != "pending").order_by(MvbClient.name).all()
    pending = MvbClient.query.filter_by(approval="pending").order_by(MvbClient.created_at).all()
    return render_template("mvb/admin_clients.html", clients=clients, pending=pending)


@bp.route("/admin/users", methods=["GET", "POST"])
def admin_users():
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        role = request.form.get("role", "")
        client_id = request.form.get("client_id", type=int)
        if not username or len(password) < 6:
            flash("Укажите логин и пароль не короче 6 символов", "danger")
        elif role not in _assignable_roles():
            flash("Выберите роль", "danger")
        elif role == "mvb_client" and not (client_id and db.session.get(MvbClient, client_id)):
            flash("Для роли «Клиент» выберите клиента", "danger")
        elif User.query.filter_by(username=username).first():
            flash("Такой логин уже занят", "danger")
        else:
            user = User(
                username=username,
                full_name=request.form.get("full_name", "").strip() or None,
                role=role,
                is_admin=False,
                mvb_client_id=client_id if role == "mvb_client" else None,
                allowed_sections="none",
            )
            user.set_password(password)
            db.session.add(user)
            db.session.commit()
            flash(f"Пользователь «{username}» создан", "success")
        return redirect(url_for("mvb.admin_users"))
    # Водители заводятся и правятся только в «Водители» — здесь их нет.
    users = (
        User.query.filter(User.role.in_([r for r in MVB_ROLES if r != "mvb_driver"]))
        .order_by(User.username).all()
    )
    clients = MvbClient.query.filter_by(is_active=True).order_by(MvbClient.name).all()
    return render_template(
        "mvb/admin_users.html", users=users, clients=clients, roles=_assignable_roles(),
        test_accounts=session.pop("mvb_test_accounts", None),
    )


def _assignable_roles():
    """Оператор заводит всех, кроме администраторов МВБ."""
    roles = {code: label for code, label in MVB_ROLES.items() if code != "mvb_driver"}
    if current_user.can_manage_mvb():
        return roles
    return {code: label for code, label in roles.items() if code != "mvb_admin"}


TEST_CLIENT_NAME = "Тестовый клиент"
TEST_ACCOUNTS = (
    ("test_client", "mvb_client", "Тестовый клиент"),
    ("test_driver", "mvb_driver", "Тестовый водитель"),
    ("test_operator", "mvb_staff", "Тестовый оператор"),
    ("test_storekeeper", "mvb_storekeeper", "Тестовый кладовщик"),
    ("test_admin", "mvb_admin", "Тестовый администратор МВБ"),
)


@bp.route("/admin/users/test-accounts", methods=["POST"])
def admin_test_accounts():
    """По одному тестовому пользователю на каждую роль. Пароли случайные и
    показываются один раз; повторное нажатие выдает тестовым учетным
    записям новые пароли."""
    if not _require_manage():
        return redirect(url_for("mvb.index"))
    client = MvbClient.query.filter_by(name=TEST_CLIENT_NAME).first()
    if client is None:
        client = MvbClient(name=TEST_CLIENT_NAME, address="Москва, тестовый адрес забора", phone="+70000000000")
        db.session.add(client)
    client.approval = "approved"
    client.is_active = True
    db.session.flush()
    created = []
    for username, role, full_name in TEST_ACCOUNTS:
        user = User.query.filter_by(username=username).first()
        if user is not None and user.role not in MVB_ROLES:
            flash(f"Логин «{username}» занят пользователем WMS — пропущен", "warning")
            continue
        if user is None:
            user = User(username=username, is_admin=False, allowed_sections="none")
            db.session.add(user)
        user.full_name = full_name
        user.role = role
        user.mvb_client_id = client.id if role == "mvb_client" else None
        user.is_active_user = True
        password = secrets.token_urlsafe(6)
        user.set_password(password)
        user.session_version = (user.session_version or 0) + 1
        created.append({"username": username, "password": password, "role": MVB_ROLES[role]})
    db.session.commit()
    session["mvb_test_accounts"] = created
    return redirect(url_for("mvb.admin_users"))


def _get_mvb_user_or_404(user_id):
    user = db.session.get(User, user_id)
    if user is None or user.role not in MVB_ROLES:
        abort(404)
    if user.role == "mvb_admin" and not current_user.can_manage_mvb():
        abort(403)
    return user


@bp.route("/admin/users/<int:user_id>/toggle", methods=["POST"])
def admin_user_toggle(user_id):
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    user = _get_mvb_user_or_404(user_id)
    if user.id == current_user.id:
        flash("Нельзя отключить самого себя", "danger")
        return redirect(url_for("mvb.admin_users"))
    user.is_active_user = not user.is_active_user
    if not user.is_active_user:
        user.session_version = (user.session_version or 0) + 1
    db.session.commit()
    flash(f"Пользователь «{user.username}» {'включен' if user.is_active_user else 'отключен'}", "success")
    return redirect(url_for("mvb.admin_users"))


@bp.route("/admin/users/<int:user_id>/password", methods=["POST"])
def admin_user_password(user_id):
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    user = _get_mvb_user_or_404(user_id)
    password = request.form.get("password", "")
    if len(password) < 6:
        flash("Пароль не короче 6 символов", "danger")
        return redirect(url_for("mvb.admin_users"))
    user.set_password(password)
    user.session_version = (user.session_version or 0) + 1
    db.session.commit()
    flash(f"Пароль для «{user.username}» изменен", "success")
    return redirect(url_for("mvb.admin_users"))


# ---------- общие помощники этапа 2 ----------


def _require_staff():
    if not _is_staff():
        flash("Доступно только сотрудникам МВБ", "danger")
        return False
    return True


def _require_operator():
    if not _is_operator():
        flash("Доступно только оператору МВБ", "danger")
        return False
    return True


def _active_drivers():
    return (
        User.query.filter(User.role == "mvb_driver", User.is_active_user.is_(True))
        .order_by(User.full_name, User.username)
        .all()
    )


def _move_box(box, status, now):
    box.status = status
    if status in BOX_TIMESTAMP_FIELDS:
        setattr(box, BOX_TIMESTAMP_FIELDS[status], now)
    user_id = current_user.id if current_user.is_authenticated else None
    db.session.add(MvbBoxEvent(box=box, status=status, user_id=user_id, created_at=now))


def _parse_dt(value):
    value = (value or "").strip()
    if not value:
        return None
    try:
        return datetime.fromisoformat(value)
    except ValueError:
        return None


def _local_to_utc(value):
    """Время из формы вводится по Москве (как отображается везде в WMS), а
    хранится в UTC — как и остальные отметки времени."""
    dt = _parse_dt(value)
    return dt - MOSCOW_OFFSET if dt else None


def _same_direction(line, marketplace, destination):
    return line.marketplace == marketplace and (
        (line.destination or "").strip().lower() == (destination or "").strip().lower()
    )


# ---------- назначение водителя ----------


@bp.route("/orders/<int:order_id>/driver", methods=["POST"])
def order_assign_driver(order_id):
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    order = _get_order_or_404(order_id)
    driver_id = request.form.get("driver_id", type=int)
    if driver_id:
        driver = db.session.get(User, driver_id)
        if driver is None or driver.role != "mvb_driver":
            abort(400)
        order.driver_id = driver.id
        flash(f"На забор назначен водитель {driver.display_name()}", "success")
    else:
        order.driver_id = None
        flash("Водитель снят с заявки", "success")
    db.session.commit()
    if request.form.get("back") == "orders":
        return redirect(url_for("mvb.orders", status=request.form.get("status", "")))
    return redirect(url_for("mvb.order_detail", order_id=order.id))


# ---------- транспорт ----------


@bp.route("/vehicles", methods=["GET", "POST"])
def vehicles():
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    if request.method == "POST":
        vehicle_id = request.form.get("vehicle_id", type=int)
        vehicle = db.session.get(MvbVehicle, vehicle_id) if vehicle_id else MvbVehicle()
        if vehicle is None or vehicle.driver_id:
            abort(404)  # авто штатного водителя правится в «Водители»
        plate = request.form.get("plate", "").strip().upper()
        try:
            capacity = int(request.form.get("capacity_boxes") or 0)
        except ValueError:
            capacity = -1
        if not plate or capacity < 0:
            flash("Укажите госномер и вместимость (число коробов)", "danger")
            return redirect(url_for("mvb.vehicles"))
        vehicle.plate = plate
        vehicle.model = request.form.get("model", "").strip() or None
        vehicle.carrier = request.form.get("carrier", "").strip() or None
        vehicle.capacity_boxes = capacity
        if vehicle_id:
            vehicle.is_active = request.form.get("is_active") == "1"
        db.session.add(vehicle)
        db.session.commit()
        flash(f"Транспорт {vehicle.plate} сохранен", "success")
        return redirect(url_for("mvb.vehicles"))
    return render_template(
        "mvb/vehicles.html",
        vehicles=MvbVehicle.query.filter(MvbVehicle.driver_id.is_(None))
        .order_by(MvbVehicle.is_active.desc(), MvbVehicle.plate).all(),
    )


# ---------- водители ----------


def _driver_vehicle(user):
    return (
        MvbVehicle.query.filter_by(driver_id=user.id)
        .order_by(MvbVehicle.is_active.desc(), MvbVehicle.id)
        .first()
    )


def _driver_form_vehicle(vehicle):
    """Авто из формы водителя: госномер, модель и вместимость обязательны."""
    plate = request.form.get("plate", "").strip().upper()
    model = request.form.get("model", "").strip()
    try:
        capacity = int(request.form.get("capacity_boxes") or 0)
    except ValueError:
        capacity = 0
    if not plate or not model or capacity <= 0:
        return "Укажите модель авто, госномер и сколько коробов вмещает машина"
    vehicle.plate = plate
    vehicle.model = model
    vehicle.capacity_boxes = capacity
    return None


@bp.route("/drivers", methods=["GET", "POST"])
def drivers():
    """Водителей заводит оператор: ФИО, авто (модель, госномер, вместимость),
    логин и пароль, которые он выдает водителю. Авто закрепляется за водителем."""
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    if request.method == "POST":
        driver_id = request.form.get("driver_id", type=int)
        full_name = request.form.get("full_name", "").strip()
        password = request.form.get("password", "")
        if driver_id:
            user = db.session.get(User, driver_id)
            if user is None or user.role != "mvb_driver":
                abort(404)
            vehicle = _driver_vehicle(user) or MvbVehicle(driver_id=user.id)
            error = None if full_name else "Укажите ФИО водителя"
            error = error or _driver_form_vehicle(vehicle)
            if not error and password and len(password) < 6:
                error = "Пароль не короче 6 символов"
            if error:
                flash(error, "danger")
                return redirect(url_for("mvb.drivers"))
            user.full_name = full_name
            active = request.form.get("is_active") == "1"
            if password or (user.is_active_user and not active):
                user.session_version = (user.session_version or 0) + 1
            if password:
                user.set_password(password)
            user.is_active_user = active
            vehicle.is_active = active
            db.session.add(vehicle)
            db.session.commit()
            flash(f"Водитель {user.display_name()} сохранен" + (" · пароль изменен" if password else ""), "success")
            return redirect(url_for("mvb.drivers"))

        username = request.form.get("username", "").strip()
        vehicle = MvbVehicle()
        error = None if full_name else "Укажите ФИО водителя"
        error = error or _driver_form_vehicle(vehicle)
        if not error and not username:
            error = "Придумайте логин водителю"
        if not error and User.query.filter(db.func.lower(User.username) == username.lower()).first():
            error = "Такой логин уже занят"
        if not error and len(password) < 6:
            error = "Пароль не короче 6 символов"
        if error:
            flash(error, "danger")
            session["mvb_driver_form"] = {k: request.form.get(k, "") for k in ("full_name", "model", "plate", "capacity_boxes", "username")}
            return redirect(url_for("mvb.drivers"))
        user = User(username=username, full_name=full_name, role="mvb_driver", is_admin=False, allowed_sections="none")
        user.set_password(password)
        vehicle.driver = user
        db.session.add_all([user, vehicle])
        db.session.commit()
        flash(f"Водитель {full_name} добавлен. Выдайте ему логин «{username}» и пароль «{password}» — вход на странице МВБ.", "success")
        return redirect(url_for("mvb.drivers"))

    rows = [
        {"user": u, "vehicle": _driver_vehicle(u)}
        for u in User.query.filter(User.role == "mvb_driver")
        .order_by(User.is_active_user.desc(), User.full_name, User.username).all()
    ]
    return render_template(
        "mvb/drivers.html", rows=rows, form=session.pop("mvb_driver_form", {}),
        free_vehicles=MvbVehicle.query.filter(MvbVehicle.driver_id.is_(None), MvbVehicle.is_active.is_(True)).count(),
    )


# ---------- паллеты ----------


@bp.route("/pallets", methods=["GET", "POST"])
def pallets():
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    if request.method == "POST":
        marketplace = request.form.get("marketplace", "")
        if marketplace not in MVB_MARKETPLACES:
            flash("Выберите маркетплейс", "danger")
            return redirect(url_for("mvb.pallets"))
        pallet = MvbPallet(
            number=next_number("mvb_pallet", "PLT-", 6),
            marketplace=marketplace,
            destination=request.form.get("destination", "").strip() or None,
            created_by_id=current_user.id,
        )
        db.session.add(pallet)
        db.session.commit()
        return redirect(url_for("mvb.pallet_detail", pallet_id=pallet.id))
    items = MvbPallet.query.order_by(MvbPallet.created_at.desc()).limit(200).all()
    return render_template("mvb/pallets.html", pallets=items)


@bp.route("/pallets/<int:pallet_id>")
def pallet_detail(pallet_id):
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    pallet = db.session.get(MvbPallet, pallet_id) or abort(404)
    # Все короба этого направления на складе (еще не погруженные) — те, что
    # уже на паллете, подсвечиваются.
    ready = [
        b for b in _ready_boxes_query().all()
        if b.line is not None and _same_direction(b.line, pallet.marketplace, pallet.destination)
    ]
    ids = {b.id for b in ready}
    boxes = ready + [b for b in pallet.boxes if b.id not in ids]
    boxes.sort(key=lambda b: (b.order.number, b.seq))
    return render_template("mvb/pallet_detail.html", pallet=pallet, boxes=boxes)


@bp.route("/pallets/<int:pallet_id>/scan", methods=["POST"])
def pallet_scan(pallet_id):
    """Скан короба на паллету: только принятые на складе короба того же
    направления, еще не погруженные; с другой паллеты короб переносится."""
    if not _is_staff():
        return jsonify(ok=False, message="Нет доступа"), 403
    pallet = db.session.get(MvbPallet, pallet_id) or abort(404)
    code = (request.form.get("barcode") or "").strip()
    box = MvbBox.query.filter(db.func.upper(MvbBox.barcode) == code.upper()).first()
    if box is None:
        return jsonify(ok=False, message=f"Короб {code} не найден"), 404
    if box.status != "received":
        return jsonify(ok=False, message=f"Короб в статусе «{box.status_label}» — на паллету только принятые на складе"), 409
    if not _same_direction(box.line, pallet.marketplace, pallet.destination):
        return jsonify(ok=False, message=f"Другое направление: {box.line.label()}"), 409
    if box.pallet_id == pallet.id:
        return jsonify(ok=True, already=True, message=f"{box.barcode} уже на этой паллете", count=len(pallet.boxes),
                       barcode=box.barcode)
    limit = _pallet_setting(MVB_PALLET_MAX_KEY, MVB_PALLET_MAX_DEFAULT)
    if limit and len(pallet.boxes) >= limit:
        return jsonify(ok=False, message=f"Паллета {pallet.number} заполнена: максимум {limit:g} кор. Начните новую паллету"), 409
    moved_from = box.pallet.number if box.pallet else None
    affected = _orders_on_pallets([box.pallet_id, pallet.id]) | {box.order}
    box.pallet = pallet
    _recalc_pallet_costs(affected)
    db.session.commit()
    message = f"{box.barcode} → {pallet.number}" + (f" (снят с {moved_from})" if moved_from else "")
    return jsonify(ok=True, already=False, message=message, count=len(pallet.boxes), barcode=box.barcode)


MVB_PALLET_PRICE_KEY = "mvb_pallet_price"
MVB_PALLET_PRICE_DEFAULT = 500.0
MVB_PALLET_MIN_KEY = "mvb_pallet_min_boxes"
MVB_PALLET_MIN_DEFAULT = 10
MVB_PALLET_MAX_KEY = "mvb_pallet_max_boxes"
MVB_PALLET_MAX_DEFAULT = 20


def _pallet_setting(key, default):
    setting = db.session.get(AppSetting, key)
    try:
        return float(setting.value) if setting and setting.value not in (None, "") else default
    except ValueError:
        return default


def _pallet_price():
    """Цена палетирования одной паллеты (настраивается в «Прайсе»)."""
    return _pallet_setting(MVB_PALLET_PRICE_KEY, MVB_PALLET_PRICE_DEFAULT)


def _orders_on_pallets(pallet_ids):
    pallet_ids = [pid for pid in pallet_ids if pid]
    if not pallet_ids:
        return set()
    return {box.order for box in MvbBox.query.filter(MvbBox.pallet_id.in_(pallet_ids)).all()}


def _recalc_pallet_costs(orders):
    """Палетирование заявки: выставляется, только если на паллетах от
    MVB_PALLET_MIN (10) коробов заявки; цена паллеты × число полных паллет
    по MVB_PALLET_MAX (20) коробов, остаток — тоже паллета. Меньше 10
    коробов — палетирование клиенту не выставляется."""
    db.session.flush()
    price = _pallet_price()
    min_boxes = _pallet_setting(MVB_PALLET_MIN_KEY, MVB_PALLET_MIN_DEFAULT)
    max_boxes = _pallet_setting(MVB_PALLET_MAX_KEY, MVB_PALLET_MAX_DEFAULT) or 1
    for order in orders:
        if order is None:
            continue
        if order.client and order.client.is_internal:
            order.pallet_cost = None
            continue
        count = MvbBox.query.filter(MvbBox.order_id == order.id, MvbBox.pallet_id.isnot(None)).count()
        if not count:
            order.pallet_cost = None
        elif count < min_boxes:
            order.pallet_cost = 0.0
        else:
            order.pallet_cost = round(price * math.ceil(count / max_boxes), 2)


@bp.route("/pallets/<int:pallet_id>/remove/<int:box_id>", methods=["POST"])
def pallet_remove_box(pallet_id, box_id):
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    box = db.session.get(MvbBox, box_id)
    if box is None or box.pallet_id != pallet_id or box.status != "received":
        abort(400)
    affected = _orders_on_pallets([pallet_id])
    box.pallet_id = None
    _recalc_pallet_costs(affected)
    db.session.commit()
    return redirect(url_for("mvb.pallet_detail", pallet_id=pallet_id))


@bp.route("/pallets/<int:pallet_id>/delete", methods=["POST"])
def pallet_delete(pallet_id):
    """Администратор удаляет паллету; ее короба остаются на складе без
    паллеты (погруженные и отправленные — в своем рейсе)."""
    if not _require_manage():
        return redirect(url_for("mvb.pallet_detail", pallet_id=pallet_id))
    pallet = db.session.get(MvbPallet, pallet_id) or abort(404)
    affected = _orders_on_pallets([pallet.id])
    for box in list(pallet.boxes):
        box.pallet_id = None
    _recalc_pallet_costs(affected)
    number = pallet.number
    db.session.delete(pallet)
    db.session.commit()
    flash(f"Паллета {number} удалена", "success")
    return redirect(url_for("mvb.pallets"))


@bp.route("/pallets/<int:pallet_id>/label.pdf")
def pallet_label_pdf(pallet_id):
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    pallet = db.session.get(MvbPallet, pallet_id) or abort(404)
    subtitle = f"{pallet.marketplace_label} · {pallet.destination or ''} · {len(pallet.boxes)} кор."
    pdf = build_labels_batch_pdf([(pallet.number, pallet.number, subtitle)], title_font_size=12, max_img_h_ratio=0.6)
    return Response(
        pdf, mimetype="application/pdf",
        headers={"Content-Disposition": content_disposition(f"{pallet.number}.pdf", "inline")},
    )


# ---------- отправка на СЦ: готово к отправке и рейсы ----------


def _ready_boxes_query():
    return (
        MvbBox.query.join(MvbOrder)
        .filter(MvbBox.status == "received", MvbBox.trip_id.is_(None))
        .order_by(MvbBox.received_at, MvbBox.id)
    )


ACTIVE_TRIP_STATUSES = ("searching", "assigned", "arrived", "loading")


def _line_in_active_trip(line):
    """Направление уже запланировано в действующий рейс — повторно в рейс
    его не ставим (иначе одну заявку можно отдать двум машинам)."""
    return line.planned_trip is not None and line.planned_trip.status in ACTIVE_TRIP_STATUSES


def _direction_key(marketplace, destination):
    return f"{marketplace}|{(destination or '').strip()}"


def _group_key(line):
    """Группа отгрузки: направление (маркетплейс + СЦ) + дата слота."""
    slot = line.slot_date.isoformat() if line.slot_date else ""
    return f"{_direction_key(line.marketplace, line.destination)}|{slot}"


def _selected_match(selected, line):
    """Отмечена ли группа направления заявки: по полному ключу или по
    направлению целиком (все слоты)."""
    return _group_key(line) in selected or _direction_key(line.marketplace, line.destination) in selected


def _slot_sort(slot_date):
    return slot_date or date.max


def _ready_groups():
    """Принятые и еще не погруженные короба по направлениям и датам слота;
    сверху ближайшие слоты, внутри — самые давние (FIFO)."""
    groups = {}
    for box in _ready_boxes_query().all():
        line = box.line
        if _line_in_active_trip(line):
            continue
        key = _group_key(line)
        group = groups.setdefault(key, {
            "key": key, "direction": _direction_key(line.marketplace, line.destination),
            "marketplace": line.marketplace, "slot_date": line.slot_date,
            "destination": (line.destination or "").strip(), "boxes": 0,
            "oldest": box.received_at, "first_id": box.id, "pallets": set(), "clients": set(),
        })
        group["boxes"] += 1
        group["clients"].add(box.order.client.name)
        if box.pallet_id:
            group["pallets"].add(box.pallet.number)
    return sorted(groups.values(), key=lambda g: (_slot_sort(g["slot_date"]), g["oldest"], g["first_id"]))


def _ready_order_items(selected=None):
    """Готовые к отправке короба по направлениям заявок — направление
    заявки едет целиком, поэтому компонуем рейсы ими. Порядок: ближайший
    слот, затем FIFO."""
    items = {}
    for box in _ready_boxes_query().all():
        line = box.line
        if _line_in_active_trip(line):
            continue
        if selected is not None and not _selected_match(selected, line):
            continue
        key = _group_key(line)
        item = items.setdefault(line.id, {
            "order": box.order, "line": line, "key": key, "slot_date": line.slot_date,
            "marketplace": line.marketplace,
            "destination": (line.destination or "").strip(), "boxes": 0, "oldest": box.received_at,
            "first_id": box.id,
        })
        item["boxes"] += 1
    # Короба приходят в порядке приемки, поэтому первый встреченный — самый
    # давний (при равном времени — принятый раньше).
    group_first = {}
    for item in items.values():
        group_first.setdefault(item["key"], (item["oldest"], item["first_id"]))
    return sorted(
        items.values(), key=lambda i: (_slot_sort(i["slot_date"]), group_first[i["key"]], i["oldest"], i["first_id"])
    )


def _pack_orders(items, capacity):
    """Компоновка рейсов без разбиения заявок: каждая заявка целиком идет в
    первую машину с тем же слотом, где хватает места, иначе — в новую.
    Заявка больше вместимости едет отдельной машиной (сверх вместимости)."""
    bins = []
    for item in items:
        target = next((
            b for b in bins
            if b["slot_date"] == item["slot_date"] and b["boxes"] + item["boxes"] <= capacity
        ), None)
        if target is None:
            target = {"items": [], "boxes": 0, "slot_date": item["slot_date"]}
            bins.append(target)
        target["items"].append(item)
        target["boxes"] += item["boxes"]
    return bins


def _vehicle_capacities():
    return sorted({
        v.capacity_boxes for v in MvbVehicle.query.filter_by(is_active=True).all() if v.capacity_boxes
    })


@bp.route("/dispatch")
def dispatch():
    """Готово к отправке: по каждому направлению — сколько коробов ждет и
    сколько машин выбранной вместимости нужно; отмеченные направления можно
    отправить одним рейсом-маршрутом или разбить на рейсы по наполненности
    авто."""
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    rows = _ready_groups()
    capacities = _vehicle_capacities()
    capacity = request.args.get("capacity", type=int) or (capacities[-1] if capacities else 0)
    open_trips = {}
    for trip in MvbTrip.query.filter(MvbTrip.status.in_(ACTIVE_TRIP_STATUSES)).all():
        for stop in trip.stops:
            open_trips.setdefault(_direction_key(stop.marketplace, stop.destination), []).append(trip)
    items = _ready_order_items()
    for row in rows:
        row["trips"] = open_trips.get(row["direction"], [])
        row["orders"] = [i for i in items if i["key"] == row["key"]]
        row["vehicles"] = len(_pack_orders(row["orders"], capacity)) if capacity else None
    total = sum(r["boxes"] for r in rows)
    proposal = _pack_orders(items, capacity) if capacity else []
    return render_template(
        "mvb/dispatch.html", rows=rows, capacity=capacity, capacities=capacities, total=total,
        total_vehicles=len(proposal) if capacity else None, proposal=proposal,
    )


@bp.route("/trips")
def trips():
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    status = request.args.get("status", "active")
    query = MvbTrip.query
    if status == "active":
        query = query.filter(MvbTrip.status.notin_(["delivered", "cancelled"]))
    elif status in MVB_TRIP_STATUSES:
        query = query.filter(MvbTrip.status == status)
    items = query.order_by(MvbTrip.created_at.desc()).limit(300).all()
    counts = dict(
        db.session.query(MvbTrip.status, db.func.count(MvbTrip.id)).group_by(MvbTrip.status).all()
    )
    return render_template("mvb/trips.html", trips=items, status=status, counts=counts)


def _add_stop(trip, marketplace, destination):
    stop = trip.stop_for(marketplace, destination)
    if stop is None:
        stop = MvbTripStop(
            marketplace=marketplace, destination=(destination or "").strip() or None,
            planned_boxes=0, seq=(max((s.seq for s in trip.stops), default=0) + 1),
        )
        trip.stops.append(stop)
    return stop


@bp.route("/trips/new", methods=["POST"])
def trip_new():
    """Рейсы по отмеченным направлениям (в порядке FIFO).

    mode=single — один рейс-маршрут по всем точкам; mode=fill — компоновка
    по наполненности авто целыми заявками (заявка не делится между машинами,
    см. _pack_orders). Каждый рейс создается в статусе «Поиск авто», заявки
    запоминаются как запланированные в него (planned_trip)."""
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    selected = set(request.form.getlist("dir"))
    groups = [g for g in _ready_groups() if g["key"] in selected or g["direction"] in selected]
    # Направление, для которого коробов уже нет (успели погрузить), все
    # равно можно добавить точкой — без плана.
    known = {g["key"] for g in groups} | {g["direction"] for g in groups}
    for key in selected - known:
        marketplace, _, rest = key.partition("|")
        destination = rest.split("|")[0]
        if marketplace in MVB_MARKETPLACES:
            groups.append({"key": key, "marketplace": marketplace, "destination": destination, "boxes": 0})
    if not groups:
        flash("Отметьте хотя бы одно направление", "warning")
        return redirect(url_for("mvb.dispatch"))

    mode = request.form.get("mode", "single")
    capacity = request.form.get("capacity", type=int) or 0
    created = []
    if not _ready_order_items(selected) and any(
        _line_in_active_trip(box.line) and _selected_match(selected, box.line) for box in _ready_boxes_query().all()
    ):
        # Повторное нажатие / второй рейс на те же заявки — не дублируем.
        flash("Эти направления уже в рейсе — второй рейс не создан", "warning")
        return redirect(url_for("mvb.dispatch"))

    def new_trip():
        trip = MvbTrip(
            number=next_number("mvb_trip", "RS-", 6), planned_boxes=0, status="searching",
            created_by_id=current_user.id,
        )
        db.session.add(trip)
        created.append(trip)
        return trip

    items = _ready_order_items(selected)
    oversized = []

    def plan(trip, item):
        stop = _add_stop(trip, item["marketplace"], item["destination"])
        stop.planned_boxes += item["boxes"]
        trip.planned_boxes += item["boxes"]
        item["line"].planned_trip = trip

    if mode == "fill" and capacity > 0:
        for bin_ in _pack_orders(items, capacity):
            trip = new_trip()
            trip.slot_date = bin_["slot_date"]
            for item in bin_["items"]:
                plan(trip, item)
                if item["boxes"] > capacity:
                    oversized.append(f'{item["order"].number} ({item["line"].short_label()})')
    else:
        trip = new_trip()
        slots = {item["slot_date"] for item in items}
        trip.slot_date = slots.pop() if len(slots) == 1 else None
        if len(slots) > 0:
            flash("В рейсе заявки с разными датами слота — проверьте, успеете ли сдать все", "warning")
        for item in items:
            plan(trip, item)
    # Направления без коробов — точкой в первый рейс, без плана.
    first = created[0] if created else new_trip()
    for group in groups:
        if group["boxes"] == 0:
            _add_stop(first, group["marketplace"], group["destination"])
    db.session.commit()
    if oversized:
        flash(f"Заявки больше вместимости машины едут отдельной машиной целиком: {', '.join(oversized)}", "warning")
    if len(created) == 1:
        flash(f"Рейс {created[0].number} создан: {created[0].route_label()} — статус «Поиск авто»", "success")
        return redirect(url_for("mvb.trip_detail", trip_id=created[0].id))
    flash(
        f"Создано рейсов: {len(created)} ({', '.join(t.number for t in created)}) по {capacity} кор., "
        "заявки не разбиты — все в статусе «Поиск авто»", "success",
    )
    return redirect(url_for("mvb.trips", status="searching"))


def _get_trip_or_404(trip_id):
    trip = db.session.get(MvbTrip, trip_id)
    if trip is None:
        abort(404)
    if not _is_staff() and not (_is_driver() and trip.driver_id == current_user.id):
        abort(404)
    return trip


@bp.route("/trips/<int:trip_id>")
def trip_detail(trip_id):
    trip = _get_trip_or_404(trip_id)
    ready = {}
    if _is_staff():
        groups = {}
        for g in _ready_groups():
            if trip.slot_date is None or g["slot_date"] in (None, trip.slot_date):
                groups[g["direction"]] = groups.get(g["direction"], 0) + g["boxes"]
        ready = {stop.id: groups.get(_direction_key(stop.marketplace, stop.destination), 0) for stop in trip.stops}
    if _is_staff() and not trip.access_token:
        trip.access_token = secrets.token_urlsafe(16)
        db.session.commit()
    return render_template(
        "mvb/trip_detail.html", trip=trip, ready=ready,
        driver_link=url_for("mvb.trip_public", token=trip.access_token, _external=True) if _is_operator() else None,
        vehicles=MvbVehicle.query.filter_by(is_active=True).order_by(MvbVehicle.plate).all() if _is_operator() else [],
        drivers=_active_drivers() if _is_operator() else [],
    )


TRIP_EDITABLE_STATUSES = ("searching", "assigned", "arrived", "loading")


@bp.route("/trips/<int:trip_id>/stops", methods=["POST"])
def trip_add_stop(trip_id):
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    trip = _get_trip_or_404(trip_id)
    marketplace = request.form.get("marketplace", "")
    if trip.status not in TRIP_EDITABLE_STATUSES or marketplace not in MVB_MARKETPLACES:
        abort(400)
    stop = _add_stop(trip, marketplace, request.form.get("destination", ""))
    db.session.commit()
    flash(f"Точка «{stop.label()}» в маршруте", "success")
    return redirect(url_for("mvb.trip_detail", trip_id=trip.id))


def _apply_stop_result(trip, stop, action, comment):
    """Итог на точке: «Сдано на СЦ» (deliver) или «Не сдано» с причиной
    (reject — короба возвращаются на склад МВБ). Возвращает (категория,
    сообщение) для flash."""
    if trip.status != "departed" or stop.result:
        return "warning", "Эта точка уже отмечена или рейс еще не в пути"
    comment = (comment or "").strip()
    if action == "reject" and not comment:
        return "danger", "Укажите причину, почему не сдано"
    now = datetime.utcnow()
    stop.result = "delivered" if action == "deliver" else "rejected"
    stop.delivered_at = now
    stop.delivered_by_id = current_user.id if current_user.is_authenticated else None
    stop.delivery_comment = comment or None
    for box in stop.boxes:
        if box.status == "shipped":
            _move_box(box, "delivered" if action == "deliver" else "not_delivered", now)
    if all(s.result for s in trip.stops):
        trip.status = "delivered"
        trip.delivered_at = now
    db.session.commit()
    if action == "deliver":
        return "success", f"Сдан на СЦ: {stop.label()} — {len(stop.boxes)} кор."
    return "warning", f"Не сдано: {stop.label()} — {len(stop.boxes)} кор. везите обратно на склад МВБ"


@bp.route("/trips/<int:trip_id>/stops/<int:stop_id>/<action>", methods=["POST"])
def trip_stop_action(trip_id, stop_id, action):
    """Порядок точек (up/down), удаление пустой точки (remove) и итог
    водителя на точке: «Сдано на СЦ» (deliver) или «Не сдано» с причиной
    (reject — короба возвращаются на склад МВБ)."""
    trip = _get_trip_or_404(trip_id)
    stop = db.session.get(MvbTripStop, stop_id)
    if stop is None or stop.trip_id != trip.id:
        abort(404)
    now = datetime.utcnow()

    if action in ("deliver", "reject"):
        category, message = _apply_stop_result(trip, stop, action, request.form.get("comment", ""))
        flash(message, category)
        if _is_driver():
            return redirect(url_for("mvb.driver"))
        return redirect(url_for("mvb.trip_detail", trip_id=trip.id))

    if not _is_operator():
        abort(403)
    if trip.status not in TRIP_EDITABLE_STATUSES:
        abort(400)
    stops = list(trip.stops)
    index = stops.index(stop)
    if action in ("up", "down"):
        other = index - 1 if action == "up" else index + 1
        if 0 <= other < len(stops):
            stops[index].seq, stops[other].seq = stops[other].seq, stops[index].seq
    elif action == "remove":
        if stop.boxes:
            flash("На эту точку уже погружены короба — сначала отмените рейс", "danger")
            return redirect(url_for("mvb.trip_detail", trip_id=trip.id))
        db.session.delete(stop)
    else:
        abort(404)
    db.session.commit()
    return redirect(url_for("mvb.trip_detail", trip_id=trip.id))


@bp.route("/trips/<int:trip_id>/plan", methods=["POST"])
def trip_plan(trip_id):
    """Авто найдено: транспорт, водитель и плановое время подачи/погрузки."""
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    trip = _get_trip_or_404(trip_id)
    if trip.status not in TRIP_EDITABLE_STATUSES:
        abort(400)
    vehicle_id = request.form.get("vehicle_id", type=int)
    vehicle = db.session.get(MvbVehicle, vehicle_id) if vehicle_id else None
    driver_id = request.form.get("driver_id", type=int)
    trip.vehicle = vehicle
    driver = db.session.get(User, driver_id) if driver_id else (vehicle.driver if vehicle else None)
    if driver_id and (driver is None or driver.role != "mvb_driver"):
        abort(400)
    trip.driver = driver
    trip.planned_arrival_at = _local_to_utc(request.form.get("planned_arrival_at"))
    trip.planned_load_start_at = _local_to_utc(request.form.get("planned_load_start_at"))
    trip.planned_load_end_at = _local_to_utc(request.form.get("planned_load_end_at"))
    trip.comment = request.form.get("comment", "").strip() or None
    # Наемный водитель (обычно случайный) — без учетной записи.
    trip.driver_name = request.form.get("driver_name", "").strip() or None
    trip.driver_phone = request.form.get("driver_phone", "").strip() or None
    trip.car_plate = request.form.get("car_plate", "").strip().upper() or None
    trip.car_model = request.form.get("car_model", "").strip() or None
    trip.capacity_boxes = request.form.get("capacity_boxes", type=int) or None
    if not trip.access_token:
        trip.access_token = secrets.token_urlsafe(16)
    if trip.has_transport() and trip.status == "searching":
        trip.status = "assigned"
    if not trip.has_transport() and trip.status == "assigned":
        trip.status = "searching"
    db.session.commit()
    if trip.has_transport():
        flash(f"Авто {trip.transport_label()} назначено на рейс {trip.number} — отправьте водителю ссылку", "success")
    else:
        flash("Укажите госномер авто", "warning")
    return redirect(url_for("mvb.trip_detail", trip_id=trip.id))


@bp.route("/trips/<int:trip_id>/delete", methods=["POST"])
def trip_delete(trip_id):
    """Администратор удаляет рейс в любом статусе: запланированные
    направления снова ждут рейса, погруженные и еще не сданные короба
    возвращаются на склад, сданные на СЦ остаются сданными."""
    if not _require_manage():
        return redirect(url_for("mvb.trip_detail", trip_id=trip_id))
    trip = db.session.get(MvbTrip, trip_id) or abort(404)
    for line in list(trip.planned_lines):
        line.planned_trip = None
    MvbOrder.query.filter_by(planned_trip_id=trip.id).update({"planned_trip_id": None})
    for box in list(trip.boxes):
        if box.status in ("loaded", "shipped"):
            box.status = "received"
            box.loaded_at = None
            box.shipped_at = None
        box.trip_id = None
        box.trip_stop_id = None
    number = trip.number
    db.session.delete(trip)
    db.session.commit()
    flash(f"Рейс {number} удален", "success")
    return redirect(url_for("mvb.trips"))


@bp.route("/trips/<int:trip_id>/<action>", methods=["POST"])
def trip_action(trip_id, action):
    """Фактические отметки рейса. Отправка переводит погруженные короба в
    «В пути на СЦ» (пустые точки убираются из маршрута); водитель рейса
    может отметить подачу авто. Сдача — по точкам (trip_stop_action)."""
    trip = _get_trip_or_404(trip_id)
    now = datetime.utcnow()
    if not _is_staff() and action != "arrive":
        abort(403)

    if action == "arrive" and trip.status == "assigned":
        trip.status = "arrived"
        trip.arrived_at = now
    elif action == "start_loading" and trip.status in ("assigned", "arrived"):
        trip.status = "loading"
        trip.arrived_at = trip.arrived_at or now
        trip.load_started_at = now
    elif action == "depart" and trip.status in ("assigned", "arrived", "loading"):
        if not trip.boxes:
            flash("В рейсе нет погруженных коробов", "danger")
            if request.form.get("back") == "loading":
                return redirect(url_for("mvb.loading_trip", trip_id=trip.id))
            return redirect(url_for("mvb.trip_detail", trip_id=trip.id))
        trip.status = "departed"
        trip.load_finished_at = now
        trip.departed_at = now
        for stop in list(trip.stops):
            if not stop.boxes:
                db.session.delete(stop)
        for box in trip.boxes:
            if box.status == "loaded":
                _move_box(box, "shipped", now)
        _sync_wms_shipped({box.order for box in trip.boxes}, now)
    elif action == "cancel" and trip.status in TRIP_EDITABLE_STATUSES:
        if not _is_operator():
            abort(403)
        trip.status = "cancelled"
        for line in list(trip.planned_lines):
            line.planned_trip = None
        for box in list(trip.boxes):
            # погруженные короба возвращаются на склад
            box.status = "received"
            box.loaded_at = None
            box.trip_id = None
            box.trip_stop_id = None
    else:
        flash("Это действие сейчас недоступно", "warning")
        return redirect(url_for("mvb.trip_detail", trip_id=trip.id))
    db.session.commit()
    if action == "depart":
        flash(f"Погрузка завершена: рейс {trip.number} ({trip.transport_label()}) в пути, коробов {len(trip.boxes)}", "success")
    else:
        flash(f"Рейс {trip.number}: {trip.status_label}", "success")
    if _is_driver():
        return redirect(url_for("mvb.driver"))
    if request.form.get("back") == "loading":
        return redirect(url_for("mvb.loading"))
    return redirect(url_for("mvb.trip_detail", trip_id=trip.id))


@bp.route("/trips/<int:trip_id>/scan", methods=["POST"])
def trip_scan(trip_id):
    """Погрузка сканом (кладовщик): короб или паллета целиком; короб
    попадает на точку маршрута своего направления. Первый скан сам отмечает
    начало погрузки. Сверх вместимости авто — предупреждение."""
    if not _is_staff():
        return jsonify(ok=False, message="Нет доступа"), 403
    trip = _get_trip_or_404(trip_id)
    if trip.status not in ("assigned", "arrived", "loading"):
        return jsonify(ok=False, message=f"Рейс в статусе «{trip.status_label}» — погрузка закрыта"), 409
    code = (request.form.get("barcode") or "").strip()
    now = datetime.utcnow()

    pallet = MvbPallet.query.filter(db.func.upper(MvbPallet.number) == code.upper()).first()
    if pallet is not None:
        boxes = [b for b in pallet.boxes if b.status == "received" and b.trip_id is None]
        if not boxes:
            return jsonify(ok=False, message=f"На паллете {pallet.number} нет коробов к погрузке"), 409
        order, line = boxes[0].order, boxes[0].line
    else:
        box = _find_box(code)
        if box is None:
            return jsonify(ok=False, message=f"Короб или паллета {code} не найдены"), 404
        if box.trip_id == trip.id:
            return jsonify(ok=True, already=True, message=f"{box.barcode} уже в этом рейсе", count=len(trip.boxes))
        if box.status != "received" or box.trip_id is not None:
            return jsonify(ok=False, message=f"Короб в статусе «{box.status_label}»"), 409
        boxes = [box]
        order, line = box.order, box.line
    stop = trip.stop_for(line.marketplace, line.destination)
    if stop is None:
        return jsonify(
            ok=False,
            message=f"Направления «{line.label()}» нет в маршруте — добавьте точку",
        ), 409

    if trip.status != "loading":
        trip.status = "loading"
        trip.arrived_at = trip.arrived_at or now
        trip.load_started_at = now
    for box in boxes:
        box.trip = trip
        box.trip_stop = stop
        _move_box(box, "loaded", now)
    db.session.commit()
    count = len(trip.boxes)
    message = (
        f"Паллета {pallet.number}: погружено {len(boxes)} кор. → {stop.label()}" if pallet is not None
        else f"{boxes[0].barcode} погружен → {stop.label()}"
    )
    warnings = []
    capacity = trip.capacity()
    if capacity and count > capacity:
        warnings.append(f"{count} кор. больше вместимости авто ({capacity})")
    if stop.planned_boxes and len(stop.boxes) > stop.planned_boxes:
        warnings.append(f"на точку «{stop.label()}» по плану {stop.planned_boxes} кор., погружено {len(stop.boxes)}")
    if trip.slot_date and line.slot_date and line.slot_date != trip.slot_date:
        warnings.append(f"слот заявки {line.slot_date.strftime('%d.%m')}, а рейса {trip.slot_date.strftime('%d.%m')}")
    if line.planned_trip_id and line.planned_trip_id != trip.id:
        warnings.append(f"заявка {order.number} ({line.short_label()}) запланирована в рейс {line.planned_trip.number}")
    warning = ("Внимание: " + "; ".join(warnings)) if warnings else None
    return jsonify(ok=True, already=False, message=message, count=count, warning=warning, loaded=len(trip.boxes))


# ---------- ссылка для водителя на СЦ (без входа) ----------
#
# На СЦ чаще всего едут случайные (наемные) водители — их не регистрируем:
# оператор вносит ФИО/телефон/госномер в рейс и отправляет водителю ссылку
# /mvb/t/<токен рейса>, где тот отмечает подачу и итог на каждой точке.


LOADING_STATUSES = ("assigned", "arrived", "loading")


@bp.route("/loading", methods=["GET", "POST"])
def loading():
    """Погрузка для кладовщика: выбрать авто по госномеру → «Начать
    погрузку» → скан паллет и коробов → «Завершить погрузку»."""
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    if request.method == "POST":
        trip = db.session.get(MvbTrip, request.form.get("trip_id", type=int))
        if trip is None or trip.status not in LOADING_STATUSES:
            flash("Выберите авто из списка", "danger")
            return redirect(url_for("mvb.loading"))
        if trip.status != "loading":
            now = datetime.utcnow()
            trip.status = "loading"
            trip.arrived_at = trip.arrived_at or now
            trip.load_started_at = now
            db.session.commit()
        return redirect(url_for("mvb.loading_trip", trip_id=trip.id))
    trips = (
        MvbTrip.query.filter(MvbTrip.status.in_(LOADING_STATUSES))
        .order_by(MvbTrip.planned_load_start_at.is_(None), MvbTrip.planned_load_start_at, MvbTrip.id)
        .all()
    )
    return render_template("mvb/loading.html", trips=trips)


@bp.route("/loading/<int:trip_id>")
def loading_trip(trip_id):
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    trip = _get_trip_or_404(trip_id)
    if trip.status not in LOADING_STATUSES:
        flash(f"Рейс {trip.number}: {trip.status_label}", "info")
        return redirect(url_for("mvb.loading"))
    return render_template("mvb/loading_trip.html", trip=trip)


# Ссылка водителя действует, пока рейс в работе, и еще сутки после
# завершения маршрута (досмотреть/поправить отметку), потом гаснет.
TRIP_LINK_GRACE = timedelta(days=1)


def _trip_by_token_or_404(token):
    trip = MvbTrip.query.filter_by(access_token=token).first() if token else None
    if trip is None or trip.status == "cancelled":
        abort(404)
    if trip.status == "delivered" and trip.delivered_at and datetime.utcnow() - trip.delivered_at > TRIP_LINK_GRACE:
        abort(410)
    return trip


@bp.errorhandler(410)
def _link_expired(_error):
    return (
        "<!doctype html><meta charset=utf-8><meta name=viewport content='width=device-width,initial-scale=1'>"
        "<div style='font-family:sans-serif;max-width:420px;margin:15vh auto;text-align:center;padding:16px'>"
        "<div style='font-size:48px'>🚚</div><h2>Ссылка больше не действует</h2>"
        "<p style='color:#666'>Рейс завершен. Если нужна новая ссылка, обратитесь к оператору МВБ.</p></div>",
        410,
    )


@bp.route("/t/<token>")
def trip_public(token):
    trip = _trip_by_token_or_404(token)
    return render_template("mvb/trip_public.html", trip=trip, token=token)


@bp.route("/t/<token>/arrive", methods=["POST"])
def trip_public_arrive(token):
    trip = _trip_by_token_or_404(token)
    if trip.status == "assigned":
        trip.status = "arrived"
        trip.arrived_at = datetime.utcnow()
        db.session.commit()
        flash("Отмечено: авто на погрузке", "success")
    return redirect(url_for("mvb.trip_public", token=token))


@bp.route("/t/<token>/stops/<int:stop_id>/<action>", methods=["POST"])
def trip_public_stop(token, stop_id, action):
    trip = _trip_by_token_or_404(token)
    stop = db.session.get(MvbTripStop, stop_id)
    if stop is None or stop.trip_id != trip.id or action not in ("deliver", "reject"):
        abort(404)
    category, message = _apply_stop_result(trip, stop, action, request.form.get("comment", ""))
    flash(message, category)
    return redirect(url_for("mvb.trip_public", token=token))


# ---------- свои короба из WMS ----------
#
# Перемещения WMS на склады маркетплейсов (их короба уже с этикетками
# BOX-...) передаются в МВБ как заявки служебного клиента «Свои короба
# (WMS)»; штрихкод короба МВБ = штрихкод короба WMS, поэтому дальше короба
# сканируются по своим же этикеткам, ничего не переклеивая. Когда рейс МВБ
# увозит все короба перемещения, в WMS ставится отметка «Транспорт забрал».

INTERNAL_CLIENT_NAME = "Свои короба (WMS)"
FINISHED_BOX_STATUSES = {"delivered"}


def _internal_client():
    client = MvbClient.query.filter_by(is_internal=True).order_by(MvbClient.id).first()
    if client is None:
        client = MvbClient(name=INTERNAL_CLIENT_NAME, is_internal=True)
        db.session.add(client)
        db.session.flush()
    return client


def _wms_candidates_query():
    """Собранные перемещения WMS на склады WB/Ozon, которые еще не уехали
    (в WMS не отмечено «Транспорт забрал»)."""
    from ..models import MovementDocument, Warehouse

    return (
        MovementDocument.query.join(Warehouse, MovementDocument.to_warehouse_id == Warehouse.id)
        .filter(
            MovementDocument.status == "completed",
            MovementDocument.shipped_at.is_(None),
            Warehouse.marketplace.in_(["wb", "ozon"]),
        )
    )


def _active_order_for_movement(doc_id):
    line = (
        MvbOrderLine.query.join(MvbOrder, MvbOrderLine.order_id == MvbOrder.id)
        .filter(MvbOrderLine.wms_movement_id == doc_id, MvbOrder.status != "cancelled")
        .first()
    )
    if line is not None:
        return line.order
    return MvbOrder.query.filter(
        MvbOrder.wms_movement_id == doc_id, MvbOrder.status != "cancelled"
    ).first()


def _find_box(code):
    """Короб МВБ по скану: свой штрихкод МВБ или этикетка короба WMS (в т.ч.
    номер BOX-000123, введенный вручную)."""
    code = (code or "").strip()
    if not code:
        return None
    box = MvbBox.query.filter(db.func.upper(MvbBox.barcode) == code.upper()).first()
    if box is not None:
        return box
    from ..models import Box

    wms_box = Box.find_by_scanned_code(code)
    if wms_box is None:
        return None
    return (
        MvbBox.query.join(MvbOrder)
        .filter(MvbBox.wms_box_id == wms_box.id, MvbOrder.status != "cancelled")
        .order_by(MvbBox.id.desc())
        .first()
    )


def _wms_movement_for_box_code(code):
    """Перемещение WMS (из кандидатов на передачу), в котором едет короб с
    этой этикеткой и которое еще не передано в МВБ."""
    from ..models import Box, MovementDocument, MovementLine

    wms_box = Box.find_by_scanned_code(code)
    if wms_box is None:
        return None
    docs = (
        _wms_candidates_query()
        .join(MovementLine, MovementLine.document_id == MovementDocument.id)
        .filter(MovementLine.box_id == wms_box.id)
        .all()
    )
    for doc in docs:
        if _active_order_for_movement(doc.id) is None:
            return doc
    return None


def _accepted_wms_boxes(doc, errors):
    """Короба перемещения, которые можно передать в МВБ (не в работе)."""
    wms_boxes = []
    seen = set()
    for line in doc.lines:
        if line.box_id not in seen:
            seen.add(line.box_id)
            wms_boxes.append(line.box)
    if not wms_boxes:
        errors.append(f"В перемещении {doc.number} нет коробов")
    accepted = []
    for wms_box in wms_boxes:
        barcode = wms_box.barcode_value
        existing = MvbBox.query.filter(MvbBox.barcode == barcode).first()
        if existing is not None:
            if existing.status in FINISHED_BOX_STATUSES or existing.order.status == "cancelled":
                # Короб WMS используется повторно — прежнюю запись МВБ
                # архивируем, чтобы штрихкод снова был свободен.
                existing.barcode = f"{existing.barcode}#{existing.id}"
            else:
                errors.append(f"Короб {wms_box.box_number} уже в работе МВБ ({existing.order.number})")
                continue
        accepted.append(wms_box)
    return accepted


def _import_movements(docs, delivery_method="pickup", pickup_address=None):
    """Одна оформленная заявка МВБ из перемещений WMS одного склада-
    отправителя: каждое перемещение — отдельное направление заявки.
    Возвращает (заявка или None, список проблем)."""
    errors = []
    todo = []
    for doc in docs:
        if _active_order_for_movement(doc.id) is not None:
            errors.append(f"Перемещение {doc.number} уже передано в МВБ")
        else:
            todo.append(doc)
    if not todo:
        return None, errors
    senders = {doc.from_warehouse_id for doc in todo}
    if len(senders) > 1:
        return None, errors + ["В одну заявку объединяются перемещения только с одного склада-отправителя"]

    parts = []
    for doc in todo:
        accepted = _accepted_wms_boxes(doc, errors)
        if accepted:
            parts.append((doc, accepted))
    if not parts:
        return None, errors

    source = parts[0][0].from_warehouse
    now = datetime.utcnow()
    order = MvbOrder(
        number=next_number("mvb_order", "MVB-", 6),
        client_id=_internal_client().id,
        delivery_method=delivery_method if delivery_method in MVB_DELIVERY_METHODS else "pickup",
        pickup_address=pickup_address or (source.address if source else None),
        comment="Перемещения WMS " + ", ".join(doc.number for doc, _ in parts)
        + (f" со склада {source.name}" if source else ""),
        status="confirmed",
        confirmed_at=now,
        created_by_id=current_user.id if current_user.is_authenticated else None,
        wms_movement_id=parts[0][0].id if len(parts) == 1 else None,
    )
    seq = 0
    for n, (doc, accepted) in enumerate(parts, start=1):
        target = doc.to_warehouse
        line = MvbOrderLine(
            seq=n,
            marketplace=target.marketplace,
            destination=target.marketplace_city or target.name,
            box_count=len(accepted),
            # Слот на СЦ — «Дата поставки» перемещения в WMS.
            slot_date=doc.delivery_slot_date,
            wms_movement_id=doc.id,
        )
        order.lines.append(line)
        for wms_box in accepted:
            seq += 1
            order.boxes.append(MvbBox(
                seq=seq, barcode=wms_box.barcode_value, status="created",
                wms_box_id=wms_box.id, line=line,
            ))
    order.sync_from_lines()
    db.session.add(order)
    db.session.commit()
    return order, errors


def _import_movement(doc, delivery_method="pickup", pickup_address=None):
    return _import_movements([doc], delivery_method, pickup_address)


def _sync_wms_shipped(orders, now):
    """Все короба перемещения уехали рейсом МВБ → в WMS «Транспорт забрал»
    (как кнопка на странице перемещения: только если заявка на МП подана)."""
    shipped = MVB_BOX_STATUS_ORDER["shipped"]
    for order in orders:
        groups = [(line.wms_movement, line.boxes) for line in order.lines if line.wms_movement is not None]
        if not groups and order.wms_movement is not None:
            groups = [(order.wms_movement, order.boxes)]
        for doc, boxes in groups:
            if doc.shipped_at is not None or not doc.marketplace_request_created_at:
                continue
            if boxes and all(MVB_BOX_STATUS_ORDER.get(b.status, 0) >= shipped for b in boxes):
                doc.shipped_at = now


def _wms_candidate_rows():
    from ..models import MovementDocument

    docs = _wms_candidates_query().order_by(MovementDocument.completed_at.desc()).limit(300).all()
    return [
        {
            "doc": doc,
            "boxes": len({line.box_id for line in doc.lines}),
            "order": _active_order_for_movement(doc.id),
        }
        for doc in docs
    ]


@bp.route("/wms")
def wms_movements():
    """Отдельная страница не используется из меню (раздел перенесен в
    "Заявки в работе", см. чат) — маршрут оставлен для прямых ссылок и
    совместимости."""
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    return render_template("mvb/wms.html", rows=_wms_candidate_rows())


@bp.route("/wms/import", methods=["POST"])
def wms_import():
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    from ..models import MovementDocument

    ids = [int(v) for v in request.form.getlist("doc_id") if v.isdigit()]
    # filter_by здесь применился бы к присоединенной таблице складов.
    docs = _wms_candidates_query().filter(MovementDocument.id.in_(ids)).all() if ids else []
    if not docs:
        flash("Отметьте перемещения, которые передаете в МВБ", "danger")
        return redirect(url_for("mvb.wms_movements"))
    order, errors = _import_movements(docs, request.form.get("delivery_method", "pickup"))
    for error in errors:
        flash(error, "warning")
    if order is None:
        return redirect(url_for("mvb.wms_movements"))
    flash(
        f"Передано в МВБ: заявка {order.number}, направлений {len(order.lines)}, коробов {order.box_count}. "
        "Этикетки WMS остаются прежними.", "success",
    )
    return redirect(url_for("mvb.order_detail", order_id=order.id))


# ---------- прайс, стоимость заявки, отчеты ----------


def _apply_prices(order):
    """Стоимость заявки по прайсу: забор — только для способа «забор», для
    своих коробов из WMS стоимость не считается."""
    if order.client and order.client.is_internal:
        return
    if order.delivery_method != "pickup":
        order.pickup_cost = None
    elif order.pickup_zone is not None:
        # Забор по зоне — фиксированная цена за забор.
        order.pickup_cost = order.pickup_zone.price
    else:
        order.pickup_cost = MvbPriceTier.cost_for("pickup", order.box_count)
    # Отправка на СЦ — по каждому направлению, по прайсу его города.
    costs = []
    for line in order.lines or []:
        dest = MvbDestination.find(line.marketplace, line.destination)
        city = MvbCity.find(dest.city if dest else line.destination)
        costs.append(MvbPriceTier.cost_for("sc", line.box_count, city.id if city else None))
    if not order.lines:
        costs.append(MvbPriceTier.cost_for("sc", order.box_count))
    known = [c for c in costs if c is not None]
    order.sc_cost = round(sum(known), 2) if known else None


def _parse_money(value):
    value = (value or "").strip().replace(" ", "").replace(",", ".")
    if not value:
        return None
    try:
        amount = float(value)
    except ValueError:
        raise ValueError(value)
    if amount < 0:
        raise ValueError(value)
    return round(amount, 2)


@bp.route("/orders/<int:order_id>/costs", methods=["POST"])
def order_costs(order_id):
    """Оператор вносит/правит стоимость забора и отправки на СЦ или
    пересчитывает ее по прайсу."""
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    order = _get_order_or_404(order_id)
    if request.form.get("action") == "recalc":
        _apply_prices(order)
        _recalc_pallet_costs({order})
        flash("Стоимость пересчитана по прайсу", "success")
    else:
        try:
            order.pickup_cost = _parse_money(request.form.get("pickup_cost"))
            order.sc_cost = _parse_money(request.form.get("sc_cost"))
            order.pallet_cost = _parse_money(request.form.get("pallet_cost"))
        except ValueError:
            flash("Стоимость — неотрицательное число", "danger")
            return redirect(url_for("mvb.order_detail", order_id=order.id))
        flash("Стоимость сохранена", "success")
    db.session.commit()
    return redirect(url_for("mvb.order_detail", order_id=order.id))


@bp.route("/guide")
def guide():
    """Обучение МВБ по ролям — для всех вошедших пользователей МВБ."""
    return render_template("mvb/guide.html")


@bp.route("/prices", methods=["GET", "POST"])
def prices():
    """Прайс оператора: цена за короб для забора и для отправки на СЦ с
    градацией «от N коробов»."""
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    if request.method == "POST":
        action = request.form.get("action", "save")
        if action in ("dest_add", "dest_toggle", "dest_save"):
            return _prices_destination_action(action)
        if action in ("city_add", "city_toggle", "city_save"):
            return _prices_city_action(action)
        if action in ("zone_add", "zone_toggle", "zone_save"):
            return _prices_zone_action(action)
        if action == "pallet_price":
            try:
                value = _parse_money(request.form.get("pallet_price"))
            except ValueError:
                value = None
            min_boxes = request.form.get("pallet_min", type=int)
            max_boxes = request.form.get("pallet_max", type=int)
            if value is None:
                flash("Укажите цену палетирования", "danger")
            elif not min_boxes or not max_boxes or min_boxes < 1 or max_boxes < min_boxes:
                flash("Укажите порог и максимум коробов на паллете (максимум не меньше порога)", "danger")
            else:
                for key, val in ((MVB_PALLET_PRICE_KEY, value), (MVB_PALLET_MIN_KEY, min_boxes), (MVB_PALLET_MAX_KEY, max_boxes)):
                    setting = db.session.get(AppSetting, key) or AppSetting(key=key)
                    setting.value = f"{val:g}"
                    db.session.add(setting)
                db.session.commit()
                flash(f"Палетирование — {value:g} руб. за паллету до {max_boxes} кор., клиенту от {min_boxes} кор. (для новых сборов паллет)", "success")
            return redirect(url_for("mvb.prices"))
        if action == "delete":
            tier = db.session.get(MvbPriceTier, request.form.get("tier_id", type=int)) or abort(404)
            db.session.delete(tier)
            db.session.commit()
            flash("Строка прайса удалена", "success")
            return redirect(url_for("mvb.prices"))
        kind = request.form.get("kind", "")
        min_boxes = request.form.get("min_boxes", type=int)
        try:
            price = _parse_money(request.form.get("price_per_box"))
        except ValueError:
            price = None
        if kind not in MVB_PRICE_KINDS or not min_boxes or min_boxes < 1 or price is None:
            flash("Укажите «от скольких коробов» (от 1) и цену за короб", "danger")
            return redirect(url_for("mvb.prices"))
        city_id = request.form.get("city_id", type=int) if kind == "sc" else None
        if city_id and db.session.get(MvbCity, city_id) is None:
            abort(404)
        tier_id = request.form.get("tier_id", type=int)
        tier = db.session.get(MvbPriceTier, tier_id) if tier_id else None
        duplicate = MvbPriceTier.query.filter_by(kind=kind, min_boxes=min_boxes, city_id=city_id).first()
        if duplicate is not None and duplicate is not tier:
            # Одна ступень на количество: правим существующую, лишнюю удаляем.
            if tier is not None:
                db.session.delete(tier)
            tier = duplicate
        if tier is None:
            tier = MvbPriceTier(kind=kind)
            db.session.add(tier)
        tier.min_boxes = min_boxes
        tier.price_per_box = price
        tier.city_id = city_id
        db.session.commit()
        flash("Прайс сохранен", "success")
        return redirect(url_for("mvb.prices"))
    destinations = MvbDestination.query.order_by(
        MvbDestination.is_active.desc(), MvbDestination.marketplace, MvbDestination.city, MvbDestination.ff_name,
    ).all()
    cities = MvbCity.query.order_by(MvbCity.is_active.desc(), MvbCity.name).all()
    tiers = {
        "pickup": MvbPriceTier.query.filter_by(kind="pickup").order_by(MvbPriceTier.min_boxes).all(),
        "sc": MvbPriceTier.query.filter_by(kind="sc", city_id=None).order_by(MvbPriceTier.min_boxes).all(),
    }
    city_tiers = {
        c.id: MvbPriceTier.query.filter_by(kind="sc", city_id=c.id).order_by(MvbPriceTier.min_boxes).all()
        for c in cities
    }
    return render_template(
        "mvb/prices.html", tiers=tiers, kinds=MVB_PRICE_KINDS, destinations=destinations, cities=cities,
        city_tiers=city_tiers, pallet_price=_pallet_price(),
        pallet_min=_pallet_setting(MVB_PALLET_MIN_KEY, MVB_PALLET_MIN_DEFAULT),
        pallet_max=_pallet_setting(MVB_PALLET_MAX_KEY, MVB_PALLET_MAX_DEFAULT), zones=MvbPickupZone.query.order_by(MvbPickupZone.is_active.desc(), MvbPickupZone.price).all(),
    )


def _destination_from_form(dest):
    """Заполняет пункт назначения из формы; возвращает текст ошибки или None."""
    marketplace = request.form.get("marketplace", "")
    city = request.form.get("city", "").strip()
    ff_name = request.form.get("ff_name", "").strip() if marketplace == "ff" else ""
    if marketplace not in MVB_MARKETPLACES:
        return "Выберите: WB, Ozon или ФФ"
    if not city:
        return "Выберите город"
    city_row = MvbCity.find(city)
    if city_row is None:
        return f"Города «{city}» нет в справочнике — сначала добавьте его в «Города»"
    city = city_row.name
    if marketplace == "ff" and not ff_name:
        return "Укажите название фулфилмента"
    value = ff_name or city
    other = MvbDestination.find(marketplace, value)
    if other is not None and other is not dest:
        return f"«{value}» уже есть в списке ({MVB_MARKETPLACES[marketplace]})"
    dest.marketplace = marketplace
    dest.city = city
    dest.ff_name = ff_name or None
    dest.address = request.form.get("address", "").strip() or None
    return None


def _prices_destination_action(action):
    """Пункты назначения: добавить, изменить, скрыть/вернуть."""
    if action == "dest_add":
        dest = MvbDestination()
        error = _destination_from_form(dest)
        if error:
            flash(error, "danger")
        else:
            db.session.add(dest)
            db.session.commit()
            flash(f"Добавлено: {dest.label()}", "success")
        return redirect(url_for("mvb.prices"))
    dest = db.session.get(MvbDestination, request.form.get("destination_id", type=int)) or abort(404)
    if action == "dest_toggle":
        dest.is_active = not dest.is_active
        db.session.commit()
        flash(f"{dest.label()}: {'снова в списке' if dest.is_active else 'скрыт из списка'}", "success")
    else:
        error = _destination_from_form(dest)
        if error:
            db.session.rollback()
            flash(error, "danger")
        else:
            db.session.commit()
            flash(f"Сохранено: {dest.label()}", "success")
    return redirect(url_for("mvb.prices"))


def _prices_city_action(action):
    """Справочник городов: добавить, переименовать, скрыть/вернуть."""
    name = " ".join(request.form.get("name", "").split())
    if action == "city_add":
        if not name:
            flash("Укажите название города", "danger")
        elif MvbCity.find(name):
            flash(f"Город «{name}» уже есть", "danger")
        else:
            db.session.add(MvbCity(name=name))
            db.session.commit()
            flash(f"Город добавлен: {name}. Задайте ему прайс отправки ниже.", "success")
        return redirect(url_for("mvb.prices"))
    city = db.session.get(MvbCity, request.form.get("city_id", type=int)) or abort(404)
    if action == "city_toggle":
        city.is_active = not city.is_active
        db.session.commit()
        flash(f"{city.name}: {'снова в списке' if city.is_active else 'скрыт — его пункты в заявке не выбрать'}", "success")
        return redirect(url_for("mvb.prices"))
    other = MvbCity.find(name)
    if not name:
        flash("Укажите название города", "danger")
    elif other is not None and other is not city:
        flash(f"Город «{name}» уже есть", "danger")
    else:
        old = city.name
        for dest in MvbDestination.query.filter_by(city=old).all():
            dest.city = name
        city.name = name
        db.session.commit()
        flash(f"Сохранено: {name}", "success")
    return redirect(url_for("mvb.prices"))


def _prices_zone_action(action):
    """Зоны забора: добавить, изменить, скрыть/вернуть."""
    zone = None
    if action != "zone_add":
        zone = db.session.get(MvbPickupZone, request.form.get("zone_id", type=int)) or abort(404)
    if action == "zone_toggle":
        zone.is_active = not zone.is_active
        db.session.commit()
        flash(f"Зона «{zone.name}»: {'снова в списке' if zone.is_active else 'скрыта'}", "success")
        return redirect(url_for("mvb.prices"))
    name = " ".join(request.form.get("name", "").split())
    try:
        price = _parse_money(request.form.get("price"))
    except ValueError:
        price = None
    other = MvbPickupZone.query.filter(db.func.lower(MvbPickupZone.name) == name.lower()).first()
    if not name or price is None:
        flash("Укажите название зоны и цену забора", "danger")
    elif other is not None and other is not zone:
        flash(f"Зона «{name}» уже есть", "danger")
    else:
        if zone is None:
            zone = MvbPickupZone()
            db.session.add(zone)
        zone.name, zone.price = name, price
        db.session.commit()
        flash(f"Зона забора «{name}» — {price:g} руб.", "success")
    return redirect(url_for("mvb.prices"))


def _report_period():
    today = date.today()
    try:
        date_from = date.fromisoformat(request.args.get("date_from", ""))
    except ValueError:
        date_from = today.replace(day=1)
    try:
        date_to = date.fromisoformat(request.args.get("date_to", ""))
    except ValueError:
        date_to = today
    return date_from, date_to


def _report_rows(date_from, date_to):
    """Отчет по клиентам за период (даты по Москве): заявки, оформленные в
    периоде, их короба по этапам и стоимость; отдельно — сколько коробов
    клиента отправлено на СЦ в периоде (по дате отправки рейса)."""
    from datetime import time

    start = datetime.combine(date_from, time.min) - MOSCOW_OFFSET
    end = datetime.combine(date_to, time.max) - MOSCOW_OFFSET
    rows = {}

    def row_for(client):
        return rows.setdefault(client.id, {
            "client": client, "orders": 0, "boxes": 0, "picked_up": 0, "received": 0,
            "shipped": 0, "delivered": 0, "not_delivered": 0,
            "pickup_cost": 0.0, "sc_cost": 0.0, "pallet_cost": 0.0, "total": 0.0,
        })

    orders = MvbOrder.query.filter(
        MvbOrder.status == "confirmed", MvbOrder.confirmed_at >= start, MvbOrder.confirmed_at <= end,
    ).order_by(MvbOrder.created_at, MvbOrder.id).all()
    order_rows = []
    for order in orders:
        order_rows.append({
            "order": order, "date": order.created_at + MOSCOW_OFFSET if order.created_at else None,
            "boxes": len(order.boxes),
            "picked_up": sum(1 for b in order.boxes if b.picked_up_at),
            "received": sum(1 for b in order.boxes if b.received_at),
            "shipped": sum(1 for b in order.boxes if b.shipped_at),
            "delivered": sum(1 for b in order.boxes if b.status == "delivered"),
            "not_delivered": sum(1 for b in order.boxes if b.status == "not_delivered"),
            "total": order.total_cost,
        })
        row = row_for(order.client)
        row["orders"] += 1
        row["boxes"] += len(order.boxes)
        row["picked_up"] += sum(1 for b in order.boxes if b.picked_up_at)
        row["received"] += sum(1 for b in order.boxes if b.received_at)
        row["pickup_cost"] += order.pickup_cost or 0
        row["sc_cost"] += order.sc_cost or 0
        row["pallet_cost"] += order.pallet_cost or 0
        row["total"] += order.total_cost or 0

    shipped = (
        MvbBox.query.join(MvbOrder)
        .filter(MvbBox.shipped_at >= start, MvbBox.shipped_at <= end)
        .all()
    )
    for box in shipped:
        row = row_for(box.order.client)
        row["shipped"] += 1
        if box.status == "delivered":
            row["delivered"] += 1
        elif box.status == "not_delivered":
            row["not_delivered"] += 1
    result = sorted(rows.values(), key=lambda r: (-r["shipped"], -r["boxes"], r["client"].name))
    totals = {key: sum(r[key] for r in result) for key in (
        "orders", "boxes", "picked_up", "received", "shipped", "delivered", "not_delivered",
        "pickup_cost", "sc_cost", "pallet_cost", "total",
    )}
    return result, totals, order_rows


def _driver_report_rows(date_from, date_to):
    """Отчет по водителям за период (по дате скана забора picked_up_at, по
    Москве): кто из водителей сколько коробов забрал и по каким заявкам
    (см. чат)."""
    from datetime import time

    start = datetime.combine(date_from, time.min) - MOSCOW_OFFSET
    end = datetime.combine(date_to, time.max) - MOSCOW_OFFSET

    boxes = (
        MvbBox.query.join(MvbOrder)
        .filter(MvbBox.picked_up_at.isnot(None), MvbBox.picked_up_at >= start, MvbBox.picked_up_at <= end)
        .all()
    )
    drivers = {u.id: u for u in User.query.filter(User.id.in_({b.picked_up_by_id for b in boxes if b.picked_up_by_id}))}

    order_rows = {}
    for box in boxes:
        key = (box.picked_up_by_id, box.order_id)
        row = order_rows.setdefault(key, {
            "driver": drivers.get(box.picked_up_by_id), "order": box.order, "boxes": 0, "picked_up_at": None,
        })
        row["boxes"] += 1
        if row["picked_up_at"] is None or box.picked_up_at > row["picked_up_at"]:
            row["picked_up_at"] = box.picked_up_at + MOSCOW_OFFSET

    def driver_name(row):
        return row["driver"].display_name() if row["driver"] else "Без водителя"

    rows = sorted(order_rows.values(), key=lambda r: (driver_name(r), r["picked_up_at"] or datetime.min))
    driver_totals = {}
    for r in rows:
        name = driver_name(r)
        driver_totals[name] = driver_totals.get(name, 0) + r["boxes"]
    totals = sorted(driver_totals.items(), key=lambda kv: (-kv[1], kv[0]))
    return rows, totals


REPORT_COLUMNS = [
    ("orders", "Заявок"), ("boxes", "Коробов в заявках"), ("picked_up", "Забрано"),
    ("received", "Принято на складе"), ("shipped", "Отправлено на СЦ"), ("delivered", "Сдано на СЦ"),
    ("not_delivered", "Не сдано"), ("pickup_cost", "Забор, руб."), ("sc_cost", "Отправка на СЦ, руб."),
    ("pallet_cost", "Палетирование, руб."), ("total", "Итого, руб."),
]


# Отчет по заявкам: у каждой заявки — дата заявки (создания, по Москве).
REPORT_ORDER_COLUMNS = [
    ("boxes", "Коробов"), ("picked_up", "Забрано"), ("received", "Принято на складе"),
    ("shipped", "Отправлено на СЦ"), ("delivered", "Сдано на СЦ"), ("not_delivered", "Не сдано"),
    ("total", "Стоимость, руб."),
]


@bp.route("/reports")
def reports():
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    date_from, date_to = _report_period()
    rows, totals, order_rows = _report_rows(date_from, date_to)
    # Отчет по водителям — тот же период, та же страница "Отчет" (см. чат).
    driver_rows, driver_totals = _driver_report_rows(date_from, date_to)
    return render_template(
        "mvb/reports.html", rows=rows, totals=totals, date_from=date_from, date_to=date_to,
        columns=REPORT_COLUMNS, order_rows=order_rows, order_columns=REPORT_ORDER_COLUMNS,
        driver_rows=driver_rows, driver_totals=driver_totals,
    )


@bp.route("/reports.xlsx")
def reports_xlsx():
    if not _require_operator():
        return redirect(url_for("mvb.index"))
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    date_from, date_to = _report_period()
    rows, totals, order_rows = _report_rows(date_from, date_to)
    wb = Workbook()
    ws = wb.active
    ws.title = "По клиентам"
    ws.append([f"МВБ Логистика — отчет по клиентам с {date_from:%d.%m.%Y} по {date_to:%d.%m.%Y}"])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["Клиент", "ИНН"] + [label for _, label in REPORT_COLUMNS])
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row["client"].name, row["client"].inn or ""] + [row[key] for key, _ in REPORT_COLUMNS])
    ws.append(["Итого", ""] + [totals[key] for key, _ in REPORT_COLUMNS])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 32
    for col in "CDEFGHIJKL":
        ws.column_dimensions[col].width = 16

    ws = wb.create_sheet("По заявкам")
    ws.append(["Дата заявки", "Заявка", "Клиент", "Направления"] + [label for _, label in REPORT_ORDER_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in order_rows:
        order = row["order"]
        ws.append([row["date"], order.number, order.client.name, order.directions_label]
                  + [row[key] for key, _ in REPORT_ORDER_COLUMNS])
        ws.cell(row=ws.max_row, column=1).number_format = "DD.MM.YYYY HH:MM"
    for col, width in zip("ABCD", (17, 14, 28, 44)):
        ws.column_dimensions[col].width = width

    driver_rows, driver_totals = _driver_report_rows(date_from, date_to)
    ws = wb.create_sheet("По водителям")
    ws.append(["Водитель", "Заявка", "Клиент", "Дата забора", "Коробов"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in driver_rows:
        order = row["order"]
        ws.append([
            row["driver"].display_name() if row["driver"] else "Без водителя",
            order.number, order.client.name,
            row["picked_up_at"].strftime("%d.%m.%Y %H:%M") if row["picked_up_at"] else "",
            row["boxes"],
        ])
    ws.append([])
    ws.append(["Итого по водителям:"])
    for name, boxes in driver_totals:
        ws.append([name, boxes])
    for cell in ws[ws.max_row - len(driver_totals)]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 18

    buffer = io.BytesIO()
    wb.save(buffer)
    fname = f"mvb_report_{date_from:%Y%m%d}_{date_to:%Y%m%d}.xlsx"
    return Response(
        buffer.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": content_disposition(fname)},
    )
