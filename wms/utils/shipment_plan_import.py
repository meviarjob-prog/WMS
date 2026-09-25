"""Разбор файла плана отгрузок (выгрузка из общей таблицы MEVIAR).

Формат листов заранее не фиксирован жестко — заголовок с городами каждый
раз может съехать на строку/колонку, а сама таблица заливается заново раз
в 2 недели под новой датой в названии листа. Поэтому вместо фиксированных
номеров строк/колонок здесь идет поиск по смыслу: строка-заголовок находится
по ячейке "Баркод", дальше колонки-города определяются по тому, что их
подпись не похожа на служебную ("остатки", "отгружен / в пути" и т.п.).
"""

import re
import unicodedata
from datetime import date

from openpyxl import load_workbook

_PERIOD_START_RE = re.compile(r"(?<!\d)(\d{1,2})\.(\d{1,2})(?!\d)")


def extract_period_start(sheet_name, today=None):
    """Дата начала периода плана из названия листа ("...от 27.08" или
    "...-27.08" -> 27 августа). Год не указан в файле — берем текущий, а если получившаяся
    дата вышла в будущем больше чем на месяц (переход через Новый год,
    например план от 28.12 гружен уже в январе) — откатываем на год назад."""
    match = _PERIOD_START_RE.search(sheet_name or "")
    if not match:
        return None
    day, month = int(match.group(1)), int(match.group(2))
    today = today or date.today()
    try:
        result = date(today.year, month, day)
    except ValueError:
        return None
    if (result - today).days > 31:
        result = date(today.year - 1, month, day)
    return result

# Подписи служебных/общих колонок (не города): и суб-колонки под городом
# (остаток/факт маркетплейса), и общие метаданные товара, которые могут
# встретиться правее штрихкода (GTIN, кол-во в коробе, план продаж и т.п.).
# Если подпись колонки не содержит ни одного из этих кусков — считаем ее
# названием нового города. Список специально с запасом (шире, чем нужно
# для текущего файла) — раз в 2 недели заливается новый файл, и лишняя
# устойчивость к формулировкам не помешает.
_SUBCOLUMN_MARKERS = (
    "остат", "отгруж", "путь", "факт", "план", "%", "gtin", "sku",
    "коробе", "производств", "хватает", "всего", "раскладк", "продаж",
    "дней", "поставк", "склад", "приорит", "коммент", "примечан", "закупщ",
)

# Колонка с комментарием закупщиков — расположение в файле заранее не
# известно (может быть как слева от штрихкода, рядом с артикулом/размером,
# так и отдельной колонкой правее всех городов), поэтому ищем по всей
# строке заголовка, а не только с одной стороны.
_COMMENT_KEYWORDS = ("коммент", "примечан", "закупщ")

_SHEET_ALIASES = {
    "ozon": ("озон",),
    "wb": ("вб", "wb"),
}

_CITY_ALIASES = {
    "екб": "Екатеринбург",
    "екатеринбург": "Екатеринбург",
    "спб": "Санкт-Петербург",
    "питер": "Санкт-Петербург",
    "санкт петербург": "Санкт-Петербург",
    "санкт-петербург": "Санкт-Петербург",
    "москва": "Москва",
}
_MARKETPLACE_PREFIX_RE = re.compile(
    r"^(?:озон|ozon|вб|wb)\s*(?::|[-–—])?\s*", re.IGNORECASE
)


def canonical_marketplace_city(value):
    """Единый ключ города для импорта плана и обратной записи факта."""
    text = unicodedata.normalize("NFKC", str(value or "")).replace("ё", "е")
    text = _MARKETPLACE_PREFIX_RE.sub("", text)
    text = re.sub(r"\s+", " ", text).strip(" \t,;:")
    key = re.sub(r"\s*[-–—]\s*", "-", text.casefold())
    if key in _CITY_ALIASES:
        return _CITY_ALIASES[key]
    moscow = re.fullmatch(r"москва\s*([12])", key)
    if moscow:
        return f"Москва {moscow.group(1)}"
    return "-".join(part.capitalize() for part in key.split("-"))


def _norm(value):
    return str(value).strip() if value is not None else ""


def _find_plan_sheets(wb, marketplace):
    """Все листы этого маркетплейса, а не только первый найденный — иначе
    при появлении в выгрузке второго листа "Распределение..." для того же
    маркетплейса (например, под отдельную категорию) он молча терялся бы.

    Лист, чье название не называет явно ни одну площадку (например,
    "Распределение Свитеры-27.08" — сводит ОБЕ площадки в одну таблицу с
    отдельными колонками штрихкода и городов на каждую, см.
    _parse_combined_marketplace_sheet), тоже считается кандидатом сразу для
    обеих площадок — по имени не различить, какая площадка на нем есть,
    это решает разбор по содержимому (parse_plan_sheet пробует оба
    парсера и просто ничего не находит для той площадки, которой на листе
    нет)."""
    own_markers = _SHEET_ALIASES[marketplace]
    other_markers = [
        m for mp, markers in _SHEET_ALIASES.items() if mp != marketplace for m in markers
    ]
    result = []
    for name in wb.sheetnames:
        lower = name.lower()
        if "распределение" not in lower:
            continue
        if any(m in lower for m in own_markers) or not any(m in lower for m in other_markers):
            result.append(name)
    return result


def _find_header_row(ws, max_scan_rows=40):
    """Возвращает (row_idx, barcode_col) — строку и колонку с "Баркод"."""
    max_row = min(ws.max_row, max_scan_rows)
    for r in range(1, max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = _norm(ws.cell(row=r, column=c).value).lower()
            if "баркод" in v:
                return r, c
    return None, None


def _find_label_col(ws, header_row, barcode_col, keyword):
    for c in range(1, barcode_col):
        v = _norm(ws.cell(row=header_row, column=c).value).lower()
        if keyword in v:
            return c
    return None


def _find_comment_col(ws, header_row, max_col=None):
    """Колонка с комментарием закупщиков — см. _COMMENT_KEYWORDS."""
    max_col = max_col or ws.max_column
    for c in range(1, max_col + 1):
        v = _norm(ws.cell(row=header_row, column=c).value).lower()
        if any(keyword in v for keyword in _COMMENT_KEYWORDS):
            return c
    return None


def _find_priority_col(ws, header_row, max_col=None):
    """Колонка "Приоритет" — как и комментарий, может стоять где угодно
    относительно штрихкода (обычно среди служебных колонок правее, см.
    "приорит" в _SUBCOLUMN_MARKERS — именно поэтому она туда и попала: не
    спутать со столбцом-городом), поэтому ищем по всей строке заголовка, а
    не только левее штрихкода, как артикул/размер."""
    max_col = max_col or ws.max_column
    for c in range(1, max_col + 1):
        v = _norm(ws.cell(row=header_row, column=c).value).lower()
        if "приорит" in v:
            return c
    return None


def _find_city_columns(ws, header_row, start_col, end_col=None):
    """[(plan_col, city_name, fact_col), ...] — plan_col это первая колонка
    группы города (план по количеству); fact_col — колонка "отгружен / в
    пути" из той же группы (уже фактически отгружено на момент выгрузки
    плана), если она есть, иначе None. Остальные служебные колонки группы
    (например "остатки" у ОЗОН) пропускаем — они не нужны.

    start_col/end_col — диапазон колонок для поиска: на обычном листе это
    все, что правее штрихкода (end_col не задан — до конца листа); на
    листе с обеими площадками сразу (см. _parse_combined_marketplace_sheet)
    — только колонки конкретной группы площадки, иначе города одной
    площадки задвоились бы с городами другой."""
    if end_col is None:
        end_col = ws.max_column
    raw = []
    for c in range(start_col, end_col + 1):
        text = _norm(ws.cell(row=header_row, column=c).value)
        raw.append((c, text))

    cities = []
    i = 0
    while i < len(raw):
        col, text = raw[i]
        if not text:
            i += 1
            continue
        lower = text.lower()
        if any(marker in lower for marker in _SUBCOLUMN_MARKERS):
            i += 1
            continue
        # Нашли новую колонку-город — ищем "отгружен / в пути" среди
        # следующих колонок этой же группы, пока не встретим следующий город.
        fact_col = None
        j = i + 1
        while j < len(raw):
            next_col, next_text = raw[j]
            next_lower = next_text.lower()
            if next_text and not any(m in next_lower for m in _SUBCOLUMN_MARKERS):
                break  # это уже следующий город
            if "отгруж" in next_lower or "путь" in next_lower:
                fact_col = next_col
            j += 1
        cities.append((col, text, fact_col))
        i += 1
    return cities


def _to_barcode_str(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _to_qty(value):
    if value is None or value == "":
        return None
    try:
        qty = float(value)
    except (TypeError, ValueError):
        return None
    return qty if qty > 0 else None


def _to_priority(value):
    """Значение колонки "Приоритет" — целое число (обычно 0/1/2, см. чат),
    либо None, если ячейка пуста или не парсится как число (нетронутая
    строка без приоритета не должна падать на нечитаемом значении)."""
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _to_fact_qty(value):
    """В отличие от _to_qty, ноль/пусто — это законное "еще не отгружено",
    а не повод пропустить город (в отличие от плана, где qty=0 не создает
    строку вовсе)."""
    if value is None or value == "":
        return 0.0
    try:
        qty = float(value)
    except (TypeError, ValueError):
        return 0.0
    return qty if qty > 0 else 0.0


class ParsedPlan:
    def __init__(self, sheet_name):
        self.sheet_name = sheet_name
        self.cities = []  # list[str] в порядке появления
        self.rows = []  # list[dict]: barcode, article, size, city, qty, fact


def _parse_one_sheet(ws):
    """Разбирает один лист стандартного формата (один "Баркод" на все
    города). Возвращает (cities, rows) либо (None, None), если на листе не
    нашлось строки-заголовка с "Баркод"."""
    header_row, barcode_col = _find_header_row(ws)
    if header_row is None:
        return None, None

    article_col = _find_label_col(ws, header_row, barcode_col, "артикул")
    size_col = _find_label_col(ws, header_row, barcode_col, "размер")
    comment_col = _find_comment_col(ws, header_row)
    priority_col = _find_priority_col(ws, header_row)
    city_columns = _find_city_columns(ws, header_row, barcode_col + 1)

    cities = [name for _, name, _ in city_columns]
    rows = []

    for r in range(header_row + 1, ws.max_row + 1):
        barcode = _to_barcode_str(ws.cell(row=r, column=barcode_col).value)
        if not barcode:
            continue  # строки-подытоги ("ВСЕГО", "КАРДИГАНЫ" и т.п.) без штрихкода

        article = _norm(ws.cell(row=r, column=article_col).value) if article_col else ""
        size = _norm(ws.cell(row=r, column=size_col).value) if size_col else ""
        comment = _norm(ws.cell(row=r, column=comment_col).value) if comment_col else ""
        priority = _to_priority(ws.cell(row=r, column=priority_col).value) if priority_col else None

        for col, city, fact_col in city_columns:
            qty = _to_qty(ws.cell(row=r, column=col).value)
            fact = _to_fact_qty(ws.cell(row=r, column=fact_col).value) if fact_col else 0.0
            # Раньше здесь пропускалась вся строка, если план (qty) пуст —
            # но реальный файл в ~1% ячеек показывает факт по городу, для
            # которого план не проставлен (0 или пусто). Пропуск такой
            # строки тихо терял уже отгруженное количество из общей суммы
            # факта (сумма расходилась с "Факт отгружено 14 д." в файле).
            # Плюс раньше факт обрезался по плану той же строки (fact =
            # min(fact, qty)) — из-за той же причины (факт местами больше
            # плана на конкретной паре товар-город) это тоже теряло часть
            # уже отгруженного. remaining_qty() и так не уходит в минус,
            # поэтому обрезать факт не нужно — тянем оба значения как есть.
            if qty is None and fact <= 0:
                continue
            rows.append(
                {
                    "barcode": barcode,
                    "article": article,
                    "size": size,
                    "city": city,
                    "qty": qty or 0.0,
                    "fact": fact,
                    "comment": comment,
                    "priority": priority,
                    "_source_row": r,
                }
            )

    rows = _exclude_blocks_outside_control_totals(
        ws, header_row, barcode_col, city_columns, rows
    )
    for row in rows:
        row.pop("_source_row", None)
    return cities, rows


def _exclude_blocks_outside_control_totals(
    ws, header_row, barcode_col, city_columns, rows
):
    """Сверяет детализацию обычного листа с его верхней итоговой строкой.

    В рабочей Google Таблице встречаются вложенные блоки: у блока есть
    собственный подытог и строки со штрихкодами, но родительский итог листа
    этот блок не включает. Простое суммирование всех штрихкодов тогда
    завышает план. Если превышение целиком совпадает с одним таким блоком,
    исключаем его строки — итог WMS становится равен контрольной строке
    источника, а остальные SKU остаются без изменений.
    """
    control_row = header_row + 1
    if _to_barcode_str(ws.cell(row=control_row, column=barcode_col).value):
        return rows

    cities = [city for _col, city, _fact_col in city_columns]
    control = {
        city: _to_qty(ws.cell(row=control_row, column=col).value) or 0.0
        for col, city, _fact_col in city_columns
    }
    if not any(control.values()):
        return rows

    detailed = {city: 0.0 for city in cities}
    for row in rows:
        detailed[row["city"]] += row["qty"]
    excess = {city: detailed[city] - control[city] for city in cities}
    tolerance = 1e-6
    if any(value < -tolerance for value in excess.values()) or not any(
        value > tolerance for value in excess.values()
    ):
        return rows

    rows_by_source = {}
    for row in rows:
        rows_by_source.setdefault(row["_source_row"], []).append(row)

    # Кандидат — подытог без штрихкода, непосредственно после которого
    # идут товарные строки до следующего подытога/заголовка.
    for candidate_row in range(control_row + 1, ws.max_row + 1):
        if _to_barcode_str(ws.cell(row=candidate_row, column=barcode_col).value):
            continue
        if not any(
            (_to_qty(ws.cell(row=candidate_row, column=col).value) or 0.0) > 0
            for col, _city, _fact_col in city_columns
        ):
            continue

        source_rows = []
        next_row = candidate_row + 1
        while next_row <= ws.max_row and _to_barcode_str(
            ws.cell(row=next_row, column=barcode_col).value
        ):
            if next_row in rows_by_source:
                source_rows.append(next_row)
            next_row += 1
        if not source_rows:
            continue

        block = {city: 0.0 for city in cities}
        for source_row in source_rows:
            for row in rows_by_source[source_row]:
                block[row["city"]] += row["qty"]
        if all(abs(block[city] - excess[city]) <= tolerance for city in cities):
            excluded = set(source_rows)
            return [row for row in rows if row["_source_row"] not in excluded]

    return rows


# Заголовки колонки штрихкода на листах, где ОБЕ площадки сведены в одну
# таблицу (см. _parse_combined_marketplace_sheet) — там нет общей колонки
# "Баркод", у каждой площадки своя, например "ШК ВБ (Горсани)"/"ШК Ozon".
_MARKETPLACE_BARCODE_MARKERS = {
    "wb": ("шк вб", "штрихкод вб"),
    "ozon": ("шк ozon", "шк озон", "штрихкод ozon", "штрихкод озон"),
}


def _find_marketplace_barcode_col(ws, marketplace, max_scan_rows=40):
    """Аналог _find_header_row, но ищет колонку штрихкода конкретной
    площадки по ее собственной подписи, а не общее "Баркод" — для листов
    вида _parse_combined_marketplace_sheet."""
    markers = _MARKETPLACE_BARCODE_MARKERS[marketplace]
    max_row = min(ws.max_row, max_scan_rows)
    for r in range(1, max_row + 1):
        for c in range(1, ws.max_column + 1):
            text = _norm(ws.cell(row=r, column=c).value).lower()
            if any(m in text for m in markers):
                return r, c
    return None, None


def _find_group_start_columns(ws, max_scan_rows=5):
    """{marketplace: column} по объединенным заголовкам групп городов
    ("ВБ (2 склада)"/"Озон (2 склада)") в строках НАД шапкой таблицы —
    только для листов с обеими площадками сразу. Ищем подпись с "склад",
    чтобы не зацепить, например, "Артикул ВБ (Горсани)" в самой шапке."""
    found = {}
    max_row = min(ws.max_row, max_scan_rows)
    for r in range(1, max_row + 1):
        for c in range(1, ws.max_column + 1):
            text = _norm(ws.cell(row=r, column=c).value).lower()
            if not text or "склад" not in text:
                continue
            for marketplace, markers in _SHEET_ALIASES.items():
                if marketplace not in found and any(m in text for m in markers):
                    found[marketplace] = c
    return found


def _parse_combined_marketplace_sheet(ws, marketplace):
    """Разбирает лист, где ОБЕ площадки сведены в одну таблицу — раздельная
    колонка штрихкода на каждую площадку ("ШК ВБ (Горсани)"/"ШК Ozon") и
    раздельная группа городов под объединенным заголовком ("ВБ (2
    склада)"/"Озон (2 склада)"), например реальный лист "Распределение
    Свитеры-27.08". В отличие от _parse_one_sheet, вызывается отдельно для
    каждой площадки — один такой лист дает вклад в план и ОЗОН, и ВБ.

    None, None — на листе нет ни колонки штрихкода, ни группы городов этой
    площадки (либо лист не такого формата вовсе, либо на нем есть данные
    только по другой площадке)."""
    header_row, barcode_col = _find_marketplace_barcode_col(ws, marketplace)
    if header_row is None:
        return None, None

    group_columns = _find_group_start_columns(ws)
    start_col = group_columns.get(marketplace)
    if start_col is None:
        return None, None
    later_group_starts = [
        c for mp, c in group_columns.items() if mp != marketplace and c > start_col
    ]
    end_col = min(later_group_starts) - 1 if later_group_starts else ws.max_column

    article_col = _find_label_col(ws, header_row, barcode_col, "артикул")
    size_col = _find_label_col(ws, header_row, barcode_col, "размер")
    comment_col = _find_comment_col(ws, header_row, end_col)
    priority_col = _find_priority_col(ws, header_row, end_col)
    city_columns = _find_city_columns(ws, header_row, start_col, end_col)

    cities = [name for _, name, _ in city_columns]
    rows = []

    for r in range(header_row + 1, ws.max_row + 1):
        barcode = _to_barcode_str(ws.cell(row=r, column=barcode_col).value)
        if not barcode:
            continue

        article = _norm(ws.cell(row=r, column=article_col).value) if article_col else ""
        size = _norm(ws.cell(row=r, column=size_col).value) if size_col else ""
        comment = _norm(ws.cell(row=r, column=comment_col).value) if comment_col else ""
        priority = _to_priority(ws.cell(row=r, column=priority_col).value) if priority_col else None

        for col, city, fact_col in city_columns:
            qty = _to_qty(ws.cell(row=r, column=col).value)
            fact = _to_fact_qty(ws.cell(row=r, column=fact_col).value) if fact_col else 0.0
            if qty is None and fact <= 0:
                continue
            rows.append(
                {
                    "barcode": barcode,
                    "article": article,
                    "size": size,
                    "city": city,
                    "qty": qty or 0.0,
                    "fact": fact,
                    "comment": comment,
                    "priority": priority,
                }
            )

    return cities, rows


def parse_plan_sheet(file_stream, marketplace):
    """Возвращает ParsedPlan либо None, если не найдено ни одного подходящего
    листа. Если листов, подходящих этому маркетплейсу, несколько (например,
    основная выгрузка плюс отдельная под какую-то категорию) — объединяет их
    все в один план, а не берет только первый найденный."""
    wb = load_workbook(file_stream, data_only=True)
    sheet_names = _find_plan_sheets(wb, marketplace)
    if not sheet_names:
        return None

    plan = ParsedPlan(", ".join(sheet_names))
    seen_cities = set()
    matched_any = False
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        period_start = extract_period_start(sheet_name)
        cities, rows = _parse_one_sheet(ws)
        if cities is None:
            # Не обычный формат (нет общей колонки "Баркод") — пробуем формат
            # с обеими площадками на одном листе (см.
            # _parse_combined_marketplace_sheet). Ничего не нашлось и там —
            # лист либо не такого формата вовсе, либо на нем данные только
            # для другой площадки (см. _find_plan_sheets).
            cities, rows = _parse_combined_marketplace_sheet(ws, marketplace)
        if cities is None:
            continue
        for row in rows:
            # Дата относится именно к листу-источнику. Это важно, когда в
            # одной Google Таблице одновременно есть несколько листов
            # «Распределение» с разными датами.
            row["period_start"] = period_start
        matched_any = True
        for city in cities:
            if city not in seen_cities:
                seen_cities.add(city)
                plan.cities.append(city)
        plan.rows.extend(rows)

    if not matched_any:
        return None
    return plan
