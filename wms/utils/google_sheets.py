"""Прямой обмен WMS с Google Таблицей плана отгрузок."""

import io
import os
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from ..models import MovementDocument, Nomenclature, ShipmentPlanLine
from .shipment_plan_import import (
    _find_city_columns,
    _find_group_start_columns,
    _find_header_row,
    _find_marketplace_barcode_col,
    _find_plan_sheets,
    _to_barcode_str,
)


OUTPUT_SHEET_TITLE = "WMS — перемещения"
OUTPUT_HEADERS = [
    "Обновлено WMS",
    "Маркетплейс",
    "Склад назначения",
    "Штрихкод",
    "Артикул",
    "Товар",
    "В пути",
    "Принято",
    "Всего отгружено",
]


def google_sheets_configured(app):
    return bool(
        app.config.get("GOOGLE_SHEETS_SPREADSHEET_ID")
        and os.path.isfile(app.config.get("GOOGLE_SERVICE_ACCOUNT_FILE", ""))
    )


def distribution_sheet_titles(sheet_titles):
    """Все листы с признаком «Распределение», без ограничения по дате."""
    return [title for title in sheet_titles if "распределение" in title.casefold()]


def _service(app):
    try:
        from google.oauth2 import service_account
        from googleapiclient.discovery import build
    except ImportError as exc:
        raise RuntimeError("Не установлены библиотеки Google API") from exc

    credentials = service_account.Credentials.from_service_account_file(
        app.config["GOOGLE_SERVICE_ACCOUNT_FILE"],
        scopes=["https://www.googleapis.com/auth/spreadsheets"],
    )
    return build("sheets", "v4", credentials=credentials, cache_discovery=False)


def _a1_sheet(title):
    return "'" + title.replace("'", "''") + "'"


def load_distribution_workbook(app):
    """Читает значения подходящих листов через Sheets API и собирает
    совместимую с существующим импортом книгу в памяти."""
    if not google_sheets_configured(app):
        raise RuntimeError("Google Таблица не настроена")
    service = _service(app)
    spreadsheet_id = app.config["GOOGLE_SHEETS_SPREADSHEET_ID"]
    metadata = service.spreadsheets().get(
        spreadsheetId=spreadsheet_id,
        fields="sheets.properties(title,index)",
    ).execute()
    titles = distribution_sheet_titles(
        [s["properties"]["title"] for s in metadata.get("sheets", [])]
    )
    if not titles:
        raise RuntimeError("Нет листов с признаком «Распределение»")

    response = service.spreadsheets().values().batchGet(
        spreadsheetId=spreadsheet_id,
        ranges=[_a1_sheet(title) for title in titles],
        majorDimension="ROWS",
        valueRenderOption="UNFORMATTED_VALUE",
        dateTimeRenderOption="FORMATTED_STRING",
    ).execute()

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, value_range in zip(titles, response.get("valueRanges", [])):
        worksheet = workbook.create_sheet(title=title)
        for row in value_range.get("values", []):
            worksheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream, titles


def resolve_sheet_title(app, spreadsheet_id, sheet_gid):
    """Название листа по его gid (числу после "gid=" в ссылке на таблицу) —
    надежнее, чем хранить название листа текстом: его могут переименовать,
    а gid не меняется. Возвращает None, если лист с таким gid не найден."""
    service = _service(app)
    metadata = service.spreadsheets().get(
        spreadsheetId=spreadsheet_id,
        fields="sheets.properties(sheetId,title)",
    ).execute()
    for sheet in metadata.get("sheets", []):
        if str(sheet["properties"]["sheetId"]) == str(sheet_gid):
            return sheet["properties"]["title"]
    return None


def read_sheet_table(app, spreadsheet_id, sheet_title):
    """Читает один лист как простую таблицу «заголовок + строки» — первая
    непустая строка считается заголовком, дальше каждая строка отдается
    структурой {название_колонки: значение}. В отличие от
    load_distribution_workbook (заточен под план отгрузок — многоуровневые
    заголовки, колонки-города), здесь формат листа заранее не предполагается
    вообще — только "таблица с шапкой", подходит для любого простого
    списка (см. production_orders_import)."""
    if not google_sheets_configured(app):
        raise RuntimeError("Google Таблица не настроена")
    service = _service(app)
    response = service.spreadsheets().values().get(
        spreadsheetId=spreadsheet_id,
        range=_a1_sheet(sheet_title),
        valueRenderOption="UNFORMATTED_VALUE",
        dateTimeRenderOption="FORMATTED_STRING",
    ).execute()
    values = response.get("values", [])
    header_row_idx = None
    for idx, row in enumerate(values):
        if any(str(cell).strip() for cell in row):
            header_row_idx = idx
            break
    if header_row_idx is None:
        return [], []
    headers = [str(cell).strip() for cell in values[header_row_idx]]
    rows = []
    for raw_row in values[header_row_idx + 1 :]:
        if not any(str(cell).strip() for cell in raw_row if cell is not None):
            continue
        row = {}
        for col_idx, header in enumerate(headers):
            if not header:
                continue
            row[header] = raw_row[col_idx] if col_idx < len(raw_row) else None
        rows.append(row)
    return headers, rows


def _movement_totals():
    totals = defaultdict(
        lambda: {
            "in_transit": 0.0,
            "received": 0.0,
            "warehouse": None,
            "nomenclature": None,
        }
    )
    documents = MovementDocument.query.filter_by(status="completed").all()
    for document in documents:
        warehouse = document.to_warehouse
        if not warehouse.marketplace:
            continue

        expected = defaultdict(float)
        nomenclature_by_id = {}
        for line in document.lines:
            for item in line.box.items:
                expected[item.nomenclature_id] += item.qty
                nomenclature_by_id[item.nomenclature_id] = item.nomenclature

        actual = dict(expected)
        if document.received_at is not None:
            for discrepancy in document.discrepancies:
                actual[discrepancy.nomenclature_id] = discrepancy.received_qty

        for nomenclature_id, expected_qty in expected.items():
            nomenclature = nomenclature_by_id.get(nomenclature_id) or Nomenclature.query.get(
                nomenclature_id
            )
            if nomenclature is None:
                continue
            key = (warehouse.id, nomenclature.id)
            totals[key]["warehouse"] = warehouse
            totals[key]["nomenclature"] = nomenclature
            if document.received_at is None:
                totals[key]["in_transit"] += expected_qty
            else:
                totals[key]["received"] += actual.get(nomenclature_id, expected_qty)
    return totals


def received_wms_totals():
    """Фактически принятые количества по складу и товару."""
    return {key: value["received"] for key, value in _movement_totals().items()}


def build_wms_movement_rows():
    """Агрегирует только факт WMS. Повторный экспорт всегда дает тот же
    результат, поэтому сетевой повтор не способен задвоить количество."""
    totals = _movement_totals()

    # Если в плане есть более подходящий артикул, используем его вместо
    # внутреннего SKU номенклатуры.
    article_by_key = {
        (line.plan.marketplace, line.warehouse.marketplace_city, line.barcode): line.article
        for line in ShipmentPlanLine.query.all()
        if line.article
    }
    updated_at = datetime.now().strftime("%d.%m.%Y %H:%M:%S")
    rows = []
    sortable = []
    for quantities in totals.values():
        warehouse = quantities["warehouse"]
        nomenclature = quantities["nomenclature"]
        sortable.append(
            (
                warehouse.marketplace,
                warehouse.marketplace_city or warehouse.name,
                nomenclature.barcode or "",
                quantities,
            )
        )
    for marketplace, city, barcode, quantities in sorted(sortable, key=lambda row: row[:3]):
        nomenclature = quantities["nomenclature"]
        sku, name = nomenclature.sku or "", nomenclature.name
        article = article_by_key.get((marketplace, city, barcode)) or sku
        in_transit = quantities["in_transit"]
        received = quantities["received"]
        rows.append(
            [
                updated_at,
                marketplace.upper(),
                city,
                barcode,
                article,
                name,
                in_transit,
                received,
                in_transit + received,
            ]
        )
    return rows


def _fact_ranges_for_sheet(worksheet, marketplace, totals):
    """Готовит обновления колонок «отгружено», не трогая строки итогов."""
    header_row, barcode_col = _find_header_row(worksheet)
    if header_row is not None:
        city_columns = _find_city_columns(worksheet, header_row, barcode_col + 1)
    else:
        header_row, barcode_col = _find_marketplace_barcode_col(worksheet, marketplace)
        if header_row is None:
            return []
        group_columns = _find_group_start_columns(worksheet)
        start_col = group_columns.get(marketplace)
        if start_col is None:
            return []
        later = [
            col for mp, col in group_columns.items() if mp != marketplace and col > start_col
        ]
        end_col = min(later) - 1 if later else worksheet.max_column
        city_columns = _find_city_columns(worksheet, header_row, start_col, end_col)

    ranges = []
    for _plan_col, city, fact_col in city_columns:
        if fact_col is None:
            continue
        values = []
        for row_number in range(header_row + 1, worksheet.max_row + 1):
            barcode = _to_barcode_str(worksheet.cell(row=row_number, column=barcode_col).value)
            value = totals.get((marketplace, city.casefold(), barcode), 0.0) if barcode else None
            values.append([value])
        if values:
            column = get_column_letter(fact_col)
            ranges.append(
                {
                    "range": (
                        f"{_a1_sheet(worksheet.title)}!{column}{header_row + 1}:"
                        f"{column}{worksheet.max_row}"
                    ),
                    "values": values,
                }
            )
    return ranges


def write_distribution_facts(app, workbook_stream):
    """Пишет «в пути + принято» в колонки «отгружено» исходных листов."""
    rows = build_wms_movement_rows()
    totals = {
        (row[1].casefold(), row[2].casefold(), str(row[3])): row[8]
        for row in rows
    }
    workbook_stream.seek(0)
    workbook = load_workbook(workbook_stream, data_only=True)
    updates = []
    for marketplace in ("ozon", "wb"):
        for sheet_name in _find_plan_sheets(workbook, marketplace):
            updates.extend(_fact_ranges_for_sheet(workbook[sheet_name], marketplace, totals))
    if not updates:
        return 0
    response = _service(app).spreadsheets().values().batchUpdate(
        spreadsheetId=app.config["GOOGLE_SHEETS_SPREADSHEET_ID"],
        body={"valueInputOption": "RAW", "data": updates},
    ).execute()
    return response.get("totalUpdatedCells", 0)


def _ensure_output_sheet(service, spreadsheet_id):
    metadata = service.spreadsheets().get(
        spreadsheetId=spreadsheet_id,
        fields="sheets.properties(sheetId,title)",
    ).execute()
    for sheet in metadata.get("sheets", []):
        if sheet["properties"]["title"] == OUTPUT_SHEET_TITLE:
            return sheet["properties"]["sheetId"]
    response = service.spreadsheets().batchUpdate(
        spreadsheetId=spreadsheet_id,
        body={
            "requests": [
                {
                    "addSheet": {
                        "properties": {
                            "title": OUTPUT_SHEET_TITLE,
                            "gridProperties": {"frozenRowCount": 1},
                        }
                    }
                }
            ]
        },
    ).execute()
    return response["replies"][0]["addSheet"]["properties"]["sheetId"]


def write_wms_movement_sheet(app):
    service = _service(app)
    spreadsheet_id = app.config["GOOGLE_SHEETS_SPREADSHEET_ID"]
    _ensure_output_sheet(service, spreadsheet_id)
    rows = build_wms_movement_rows()
    values = [OUTPUT_HEADERS] + rows
    service.spreadsheets().values().update(
        spreadsheetId=spreadsheet_id,
        range=f"{_a1_sheet(OUTPUT_SHEET_TITLE)}!A1:I{len(values)}",
        valueInputOption="RAW",
        body={"values": values},
    ).execute()
    # Старые хвостовые строки очищаем только после успешной записи новых.
    service.spreadsheets().values().clear(
        spreadsheetId=spreadsheet_id,
        range=f"{_a1_sheet(OUTPUT_SHEET_TITLE)}!A{len(values) + 1}:I",
        body={},
    ).execute()
    return len(rows)
