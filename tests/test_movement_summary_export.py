"""Сводный экспорт перемещений (/movement/export-summary.xlsx) — одна
строка на документ целиком с количеством коробов и суммарным количеством
товара, в отличие от export_all (там строка на каждый товар в каждом
коробе)."""

import io
from datetime import datetime

import openpyxl

from wms.extensions import db
from wms.models import (
    Box,
    BoxItem,
    MovementDocument,
    MovementLine,
    MovementReceiptDiscrepancy,
    Nomenclature,
    Warehouse,
)


def _make_item(barcode, name="Товар для сводки"):
    item = Nomenclature(sku=barcode, barcode=barcode, name=name, unit="шт")
    db.session.add(item)
    db.session.commit()
    return item


def _make_movement_with_boxes(
    number,
    box_specs,
    sender_code="WH-MSUM-A",
    receiver_code="WH-MSUM-B",
    receiver_marketplace=None,
    **doc_kwargs,
):
    """box_specs — список (qty,) на короб, по одному товару в каждом."""
    sender = Warehouse.query.filter_by(code=sender_code).first() or Warehouse(code=sender_code, name="Отправитель")
    receiver = Warehouse.query.filter_by(code=receiver_code).first() or Warehouse(code=receiver_code, name="Получатель")
    receiver.marketplace = receiver_marketplace
    if sender.id is None:
        db.session.add(sender)
    if receiver.id is None:
        db.session.add(receiver)
    db.session.commit()

    doc = MovementDocument(
        number=number, from_warehouse_id=sender.id, to_warehouse_id=receiver.id, status="completed", **doc_kwargs
    )
    db.session.add(doc)
    db.session.commit()

    for i, qty in enumerate(box_specs):
        item = _make_item(f"777{number}{i}")
        box = Box(box_number=f"BOX-{number}-{i}", warehouse_id=sender.id, status="open")
        db.session.add(box)
        db.session.commit()
        db.session.add(BoxItem(box_id=box.id, nomenclature_id=item.id, qty=qty))
        db.session.add(MovementLine(document_id=doc.id, box_id=box.id, from_warehouse_id=sender.id))
    db.session.commit()
    return doc


def _read_xlsx_rows(data):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    return [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]


def test_summary_export_has_one_row_per_document_with_box_and_item_counts(db, client_logged_in):
    doc = _make_movement_with_boxes("MSUM-1", [3, 5, 2])

    resp = client_logged_in.get("/movement/export-summary.xlsx")
    rows = _read_xlsx_rows(resp.data)

    row = next(r for r in rows if r[0] == doc.number)
    assert row[10] == 3  # кол-во коробов
    assert row[11] == 10  # суммарное кол-во товара в коробах (3+5+2)
    assert row[12] == 0  # без заявки документ еще не входит в факт плана


def test_summary_export_headers(db, client_logged_in):
    _make_movement_with_boxes("MSUM-2", [1])
    resp = client_logged_in.get("/movement/export-summary.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == [
        "Номер документа",
        "Дата",
        "Статус",
        "Склад-источник",
        "Маркетплейс",
        "Склад-назначение",
        "№ заявки МП",
        "Заявка на МП создана",
        "Дата отправки",
        "Принято на складе",
        "Кол-во коробов",
        "Кол-во в коробах",
        "Кол-во отгружено",
    ]


def test_summary_export_includes_warehouses_marketplace_and_request_number(db, client_logged_in):
    doc = _make_movement_with_boxes(
        "MSUM-3", [4], receiver_marketplace="wb", marketplace_request_number="REQ-99"
    )

    resp = client_logged_in.get("/movement/export-summary.xlsx")
    rows = _read_xlsx_rows(resp.data)
    row = next(r for r in rows if r[0] == doc.number)

    assert row[3] == doc.from_warehouse.name
    assert row[4] == "ВБ"
    assert row[5] == doc.to_warehouse.name
    assert row[6] == "REQ-99"


def test_summary_export_marks_ozon_destination(db, client_logged_in):
    doc = _make_movement_with_boxes("MSUM-4", [2], receiver_code="WH-MSUM-OZON", receiver_marketplace="ozon")

    resp = client_logged_in.get("/movement/export-summary.xlsx")
    rows = _read_xlsx_rows(resp.data)
    row = next(r for r in rows if r[0] == doc.number)

    assert row[4] == "ОЗОН"


def test_summary_export_separates_boxed_and_actually_received_qty(db, client_logged_in):
    now = datetime.utcnow()
    doc = _make_movement_with_boxes(
        "MSUM-5",
        [10],
        receiver_code="WH-MSUM-ACTUAL",
        receiver_marketplace="wb",
        completed_at=now,
        marketplace_request_number="REQ-MSUM-5",
        marketplace_request_created_at=now,
        shipped_at=now,
        received_at=now,
    )
    box_item = doc.lines.first().box.items.first()
    db.session.add(
        MovementReceiptDiscrepancy(
            document_id=doc.id,
            nomenclature_id=box_item.nomenclature_id,
            expected_qty=10,
            received_qty=7,
        )
    )
    db.session.commit()

    resp = client_logged_in.get("/movement/export-summary.xlsx")
    row = next(r for r in _read_xlsx_rows(resp.data) if r[0] == doc.number)

    assert row[7] == "Да"
    assert row[8]
    assert row[9]
    assert row[11] == 10
    assert row[12] == 7


def test_summary_export_excludes_no_documents_and_response_is_xlsx(db, client_logged_in):
    resp = client_logged_in.get("/movement/export-summary.xlsx")
    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_summary_export_has_daily_shipments_sheet(db, client_logged_in):
    """Вторая страница файла — отгрузки по дням: суммарно коробов и товара
    за каждый день отгрузки (shipped_at), документы без отметки об отгрузке
    в эту страницу не попадают (см. чат)."""
    day1 = datetime(2026, 9, 1, 10, 0)
    day2 = datetime(2026, 9, 2, 11, 0)
    _make_movement_with_boxes("MSUM-D1", [3, 2], receiver_code="WH-MSUM-D1", shipped_at=day1)
    _make_movement_with_boxes("MSUM-D2", [5], receiver_code="WH-MSUM-D2", shipped_at=day1)
    _make_movement_with_boxes("MSUM-D3", [4], receiver_code="WH-MSUM-D3", shipped_at=day2)
    _make_movement_with_boxes("MSUM-D4-NOSHIP", [10], receiver_code="WH-MSUM-D4")

    resp = client_logged_in.get("/movement/export-summary.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(resp.data))

    assert wb.sheetnames == ["Перемещения (сводно)", "Отгрузки по дням"]
    ws2 = wb["Отгрузки по дням"]
    header = [c.value for c in ws2[1]]
    assert header == ["Дата отгрузки", "Кол-во коробов", "Кол-во товара, шт"]

    rows = {row[0]: (row[1], row[2]) for row in ws2.iter_rows(min_row=2, values_only=True) if row[0]}
    assert rows["2026-09-01"] == (3, 10)  # 2 короба + 1 короб = 3; (3+2)+5 = 10 шт
    assert rows["2026-09-02"] == (1, 4)
    assert len(rows) == 2  # неотгруженный документ (MSUM-D4) не попал


def test_summary_export_date_range_filters_by_shipped_at(db, client_logged_in):
    early = datetime(2026, 8, 1, 10, 0)
    inside = datetime(2026, 9, 10, 10, 0)
    late = datetime(2026, 10, 1, 10, 0)
    doc_early = _make_movement_with_boxes("MSUM-R1", [1], receiver_code="WH-MSUM-R1", shipped_at=early)
    doc_inside = _make_movement_with_boxes("MSUM-R2", [1], receiver_code="WH-MSUM-R2", shipped_at=inside)
    doc_late = _make_movement_with_boxes("MSUM-R3", [1], receiver_code="WH-MSUM-R3", shipped_at=late)

    resp = client_logged_in.get("/movement/export-summary.xlsx?date_from=2026-09-01&date_to=2026-09-30")
    rows = _read_xlsx_rows(resp.data)
    numbers = {r[0] for r in rows}

    assert doc_inside.number in numbers
    assert doc_early.number not in numbers
    assert doc_late.number not in numbers


def test_summary_export_date_to_is_inclusive_of_whole_day(db, client_logged_in):
    end_of_day = datetime(2026, 9, 30, 23, 59)
    doc = _make_movement_with_boxes("MSUM-R4", [1], receiver_code="WH-MSUM-R4", shipped_at=end_of_day)

    resp = client_logged_in.get("/movement/export-summary.xlsx?date_from=2026-09-01&date_to=2026-09-30")
    numbers = {r[0] for r in _read_xlsx_rows(resp.data)}

    assert doc.number in numbers


def test_summary_export_without_dates_includes_everything_like_before(db, client_logged_in):
    doc_no_ship = _make_movement_with_boxes("MSUM-R5", [1], receiver_code="WH-MSUM-R5")
    doc_shipped = _make_movement_with_boxes(
        "MSUM-R6", [1], receiver_code="WH-MSUM-R6", shipped_at=datetime(2020, 1, 1)
    )

    resp = client_logged_in.get("/movement/export-summary.xlsx")
    numbers = {r[0] for r in _read_xlsx_rows(resp.data)}

    assert doc_no_ship.number in numbers
    assert doc_shipped.number in numbers


def test_list_page_has_summary_export_date_range_fields(db, client_logged_in):
    html = client_logged_in.get("/movement/").get_data(as_text=True)
    assert 'name="date_from"' in html
    assert 'name="date_to"' in html
