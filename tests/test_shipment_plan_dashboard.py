"""«В пути» в плане — перемещения, которые транспорт забрал с даты листа."""

import io

from openpyxl import load_workbook

from wms.extensions import db
from wms.models import (
    Box,
    BoxItem,
    MovementDocument,
    MovementLine,
    Nomenclature,
    ReceivingDocument,
    ReceivingLine,
    ShipmentPlan,
    ShipmentPlanLine,
    Warehouse,
)


def _setup(planned_qty=30):
    sender = Warehouse(code="WH-D1", name="Склад-отправитель")
    city = Warehouse(code="WH-D2", name="ОЗОН: Город", marketplace="ozon", marketplace_city="Город")
    db.session.add_all([sender, city])
    db.session.commit()

    item = Nomenclature(sku="SKU-D1", barcode="7770000001", name="Товар", unit="шт")
    db.session.add(item)
    db.session.commit()

    plan = ShipmentPlan(marketplace="ozon")
    db.session.add(plan)
    db.session.commit()
    line = ShipmentPlanLine(
        plan_id=plan.id,
        warehouse_id=city.id,
        nomenclature_id=item.id,
        barcode=item.barcode,
        article="ART-1",
        planned_qty=planned_qty,
        fulfilled_qty=0,
    )
    db.session.add(line)
    db.session.commit()
    return sender, city, item


def _ship_box(sender, city, item, qty, box_number, client, mark_request=True):
    box = Box(box_number=box_number, warehouse_id=sender.id, status="open")
    db.session.add(box)
    db.session.commit()
    db.session.add(BoxItem(box_id=box.id, nomenclature_id=item.id, qty=qty))
    db.session.commit()

    doc = MovementDocument(number=f"PER-{box_number}", from_warehouse_id=sender.id, to_warehouse_id=city.id)
    db.session.add(doc)
    db.session.commit()
    db.session.add(
        MovementLine(document_id=doc.id, box_id=box.id, from_warehouse_id=box.warehouse_id, from_cell_id=box.cell_id)
    )
    db.session.commit()

    client.post(f"/movement/{doc.id}/complete")
    if mark_request:
        doc.marketplace_request_number = f"REQ-{box_number}"
        db.session.commit()
        client.post(f"/movement/{doc.id}/mark-marketplace-request")
        client.post(f"/movement/{doc.id}/mark-shipped")
    return doc


def test_dashboard_top_summary_shows_in_transit(db, client_logged_in):
    sender, city, item = _setup(planned_qty=30)
    doc = _ship_box(sender, city, item, qty=10, box_number="BOX-000001", client=client_logged_in)

    doc = MovementDocument.query.get(doc.id)
    assert doc.status == "completed"
    assert doc.received_at is None

    resp = client_logged_in.get("/shipment-plan/")
    html = resp.get_data(as_text=True)
    idx = html.find("Город")
    assert "10" in html[idx : idx + 400]


def test_city_in_transit_includes_shipped_sku_missing_from_current_plan(
    db, client_logged_in
):
    """Городская сумма — все отгрузки, а не только совпавшие строки плана."""
    from datetime import date, timedelta

    sender, city, planned_item = _setup(planned_qty=30)
    plan = ShipmentPlan.query.filter_by(marketplace="ozon").first()
    plan.period_start = date.today() - timedelta(days=1)
    unplanned_item = Nomenclature(
        sku="SKU-NOT-IN-PLAN",
        barcode="7770000999",
        name="Товар вне актуального плана",
        unit="шт",
    )
    db.session.add(unplanned_item)
    db.session.commit()

    _ship_box(
        sender,
        city,
        planned_item,
        qty=10,
        box_number="BOX-PLANNED-CITY-TOTAL",
        client=client_logged_in,
    )
    _ship_box(
        sender,
        city,
        unplanned_item,
        qty=824,
        box_number="BOX-UNPLANNED-CITY-TOTAL",
        client=client_logged_in,
        mark_request=True,
    )

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    city_idx = html.find("<td>Город</td>")
    city_snippet = html[city_idx : city_idx + 300]
    assert ">834<" in city_snippet
    top_idx = html.find("в пути")
    assert "<b>834</b>" in html[top_idx : top_idx + 100]


def test_picking_list_shows_plan_and_in_transit_separately(db, client_logged_in):
    """В таблице остается остаток плана, товар в пути показан в скобках.
    Защита от лишней отправки применяется отдельно в маршрутизации."""
    sender, city, item = _setup(planned_qty=30)
    _ship_box(sender, city, item, qty=10, box_number="BOX-000002", client=client_logged_in)

    resp = client_logged_in.get("/shipment-plan/")
    html = resp.get_data(as_text=True)
    idx = html.find("ART-1")
    snippet = html[idx : idx + 3000]

    assert ">30<" in snippet
    assert "(10)" in snippet


def test_completed_movement_without_transport_pickup_is_not_in_transit(db, client_logged_in):
    sender, city, item = _setup(planned_qty=30)
    _ship_box(
        sender,
        city,
        item,
        qty=10,
        box_number="BOX-NO-REQUEST",
        client=client_logged_in,
        mark_request=False,
    )

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    idx = html.find("ART-1")
    snippet = html[idx : idx + 3000]
    assert ">30<" in snippet
    assert "(10)" not in snippet


def test_top_summary_shows_in_transit_per_marketplace_and_total(db, client_logged_in):
    sender, city, item = _setup(planned_qty=30)
    _ship_box(sender, city, item, qty=10, box_number="BOX-SUM-1", client=client_logged_in)

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    assert "Сводка" in html
    idx = html.find("В пути")
    snippet = html[idx : idx + 600]
    assert "ОЗОН" in snippet
    assert ">10<" in snippet
    assert "Итого" in snippet


def test_top_summary_shows_total_stock(db, client_logged_in):
    from wms.models import UnplacedStock

    sender, city, item = _setup(planned_qty=30)
    UnplacedStock.add(sender.id, item.id, 25)
    db.session.commit()

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    idx = html.find("На складе:")
    assert "25" in html[idx : idx + 100]


def test_top_summary_shows_total_production_since_period_start(db, client_logged_in):
    from datetime import date, timedelta

    from wms.models import ProductionRecord, User

    sender, city, item = _setup(planned_qty=30)
    plan = ShipmentPlan.query.filter_by(marketplace="ozon").first()
    plan.period_start = date.today() - timedelta(days=3)
    db.session.commit()

    worker = User(username="prod-worker", role="production")
    worker.set_password("x")
    db.session.add(worker)
    db.session.commit()

    # В периоде плана — должно засчитаться.
    db.session.add(
        ProductionRecord(user_id=worker.id, nomenclature_id=item.id, work_date=date.today())
    )
    # До начала периода — не должно засчитаться.
    db.session.add(
        ProductionRecord(
            user_id=worker.id, nomenclature_id=item.id, work_date=date.today() - timedelta(days=10)
        )
    )
    db.session.commit()

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    idx = html.find("На производстве")
    assert ">1<" in html[idx : idx + 200]


def test_marketplace_header_uses_only_completed_movement_fact(db, client_logged_in):
    from datetime import date, timedelta

    sender, city, item = _setup(planned_qty=30)
    plan = ShipmentPlan.query.filter_by(marketplace="ozon").first()
    plan.period_start = date.today() - timedelta(days=3)
    line = ShipmentPlanLine.query.first()
    line.fulfilled_qty = 5
    db.session.commit()
    _ship_box(sender, city, item, qty=10, box_number="BOX-SUM-2", client=client_logged_in)

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    # Сохраненный старый счетчик 5 не участвует: факт WMS — 10 отправлено.
    idx = html.find("в пути")
    snippet = html[idx : idx + 200]
    assert "<b>10</b>" in snippet
    assert "Пока нет отгрузок" not in html
    assert "Выполнено" not in html

    city_idx = html.find("<td>Город</td>")
    city_snippet = html[city_idx : city_idx + 500]
    assert "Все завершенные перемещения с 00:01 даты листа" in html
    assert ">10<" in city_snippet


def test_dashboard_keeps_sent_qty_after_marketplace_receives_less(
    db, client_logged_in
):
    sender, city, item = _setup(planned_qty=30)
    doc = _ship_box(
        sender,
        city,
        item,
        qty=6,
        box_number="BOX-RECEIVED-FACT",
        client=client_logged_in,
    )
    client_logged_in.post(
        f"/movement/{doc.id}/receive",
        data={f"qty_{item.id}": "4"},
    )
    line = ShipmentPlanLine.query.first()
    line.fulfilled_qty = 0
    db.session.commit()

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    idx = html.find("в пути")
    # Отправлено 6; недовоз 2 ведется отдельно и не уменьшает отгрузку WMS.
    assert "<b>6</b>" in html[idx : idx + 200]


def test_excel_export_matches_dashboard_table_and_keeps_transit_separate(db, client_logged_in):
    sender, city, item = _setup(planned_qty=30)
    line = ShipmentPlanLine.query.first()
    line.fulfilled_qty = 5
    db.session.commit()
    _ship_box(sender, city, item, qty=10, box_number="BOX-XLSX-1", client=client_logged_in)

    response = client_logged_in.get("/shipment-plan/export.xlsx")

    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data), data_only=True)
    sheet = workbook["План отгрузок"]
    assert "A1:A2" in {str(cell_range) for cell_range in sheet.merged_cells.ranges}
    assert sheet["A1"].value == "Артикул"
    assert sheet["E1"].value == "Общий план"
    assert sheet["F1"].value == "Не хватает по плану"
    assert sheet["G1"].value == "На разбраковке"
    assert sheet["J1"].value == "ОЗОН"
    assert sheet["J2"].value == "Город"
    assert sheet["K1"].value == "Комментарий закупщиков"
    assert sheet["A3"].value == "Итого (1 поз.)"
    assert sheet["E3"].value == 30
    assert sheet["I3"].value == 10
    assert sheet["J3"].value == 30
    assert sheet["A4"].value == "ART-1"
    assert sheet["J4"].value == "30 (10)"


def test_picking_list_shows_total_planned_and_shortfall_columns(db, client_logged_in):
    """«Общий план» и «Не хватает по плану» — построчно по товару, сумма по
    всем городам обоих маркетплейсов. «Не хватает» = план минус то, что уже
    в пути (не меньше нуля), а не просто остаток плана без учета транзита."""
    sender, city, item = _setup(planned_qty=30)
    _ship_box(sender, city, item, qty=10, box_number="BOX-TOTALS-1", client=client_logged_in)

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    assert "Общий план" in html
    assert "Не хватает по плану" in html
    idx = html.find("ART-1")
    snippet = html[idx : idx + 3000]
    assert ">30<" in snippet  # Общий план
    assert ">20<" in snippet  # Не хватает: 30 - 10 в пути


def test_shortfall_column_floors_at_zero_when_transit_covers_plan(db, client_logged_in):
    sender, city, item = _setup(planned_qty=10)
    _ship_box(sender, city, item, qty=10, box_number="BOX-TOTALS-2", client=client_logged_in)

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    idx = html.find("ART-1")
    snippet = html[idx : idx + 3000]
    assert ">10<" in snippet  # Общий план и «в пути» совпадают
    assert "text-muted\">0<" in snippet  # Не хватает: max(10 - 10, 0) = 0


def test_picking_list_shows_buyer_comment_column(db, client_logged_in):
    """Комментарий показан как поле ввода (можно править прямо в WMS), а не
    просто текстом — со значением из buyer_comment."""
    sender, city, item = _setup(planned_qty=30)
    line = ShipmentPlanLine.query.first()
    line.buyer_comment = "Поставка задерживается на неделю"
    db.session.commit()

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    assert "Комментарий закупщиков" in html
    assert 'value="Поставка задерживается на неделю"' in html


def test_comment_column_is_positioned_right_after_barcode(db, client_logged_in):
    sender, city, item = _setup(planned_qty=30)

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    barcode_idx = html.index(">Штрихкод<")
    comment_idx = html.index("Комментарий закупщиков")
    total_planned_idx = html.index("Общий план")
    assert barcode_idx < comment_idx < total_planned_idx


def test_update_comment_endpoint_saves_buyer_comment(db, client_logged_in):
    sender, city, item = _setup(planned_qty=30)

    response = client_logged_in.post(
        f"/shipment-plan/comment/{item.barcode}",
        data={"comment": "Уточнить у поставщика"},
    )

    assert response.status_code == 302
    line = ShipmentPlanLine.query.first()
    assert line.buyer_comment == "Уточнить у поставщика"


def test_update_comment_endpoint_clears_comment_on_empty_input(db, client_logged_in):
    sender, city, item = _setup(planned_qty=30)
    line = ShipmentPlanLine.query.first()
    line.buyer_comment = "Старый комментарий"
    db.session.commit()

    client_logged_in.post(f"/shipment-plan/comment/{item.barcode}", data={"comment": "   "})

    line = ShipmentPlanLine.query.first()
    assert line.buyer_comment is None


def test_update_comment_endpoint_updates_all_lines_sharing_barcode(db, client_logged_in):
    """Один штрихкод может встречаться в нескольких строках плана (разные
    города и площадки) — комментарий один на товар, должен обновиться сразу
    во всех, иначе "потеряется" при показе другого направления того же
    товара."""
    sender, city, item = _setup(planned_qty=30)
    other_city = Warehouse(
        code="WH-D3", name="ВБ: Город2", marketplace="wb", marketplace_city="Город2"
    )
    db.session.add(other_city)
    db.session.commit()
    wb_plan = ShipmentPlan(marketplace="wb")
    db.session.add(wb_plan)
    db.session.commit()
    db.session.add(
        ShipmentPlanLine(
            plan_id=wb_plan.id,
            warehouse_id=other_city.id,
            nomenclature_id=item.id,
            barcode=item.barcode,
            article="ART-1",
            planned_qty=10,
        )
    )
    db.session.commit()

    client_logged_in.post(
        f"/shipment-plan/comment/{item.barcode}", data={"comment": "Общий комментарий"}
    )

    comments = {
        line.buyer_comment
        for line in ShipmentPlanLine.query.filter_by(barcode=item.barcode).all()
    }
    assert comments == {"Общий комментарий"}


def test_update_comment_endpoint_unknown_barcode_flashes_error(db, client_logged_in):
    response = client_logged_in.post(
        "/shipment-plan/comment/0000000000000", data={"comment": "test"}, follow_redirects=True
    )

    assert "не найден" in response.get_data(as_text=True)


def test_picking_list_has_totals_row_summing_columns(db, client_logged_in):
    """Строка "Итого" сразу под заголовками — просто сумма по столбцам
    (На разбраковке/Готово к отгрузке/В пути и по каждому городу)."""
    from wms.models import UnplacedStock

    sender, city, item = _setup(planned_qty=30)
    UnplacedStock.add(sender.id, item.id, 8)
    db.session.commit()
    _ship_box(sender, city, item, qty=10, box_number="BOX-TOT-1", client=client_logged_in)

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    idx = html.find("Итого (")
    assert idx != -1
    snippet = html[idx : idx + 800]
    assert "1 поз." in snippet
    assert ">8<" in snippet  # На разбраковке
    assert ">30<" in snippet  # Остаток плана; 10 в пути показаны отдельно
    assert ">10<" in snippet  # В пути


def test_picking_list_has_v_puti_column(db, client_logged_in):
    sender, city, item = _setup(planned_qty=30)
    _ship_box(sender, city, item, qty=10, box_number="BOX-TOT-2", client=client_logged_in)

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    idx = html.find("В пути")
    assert idx != -1
    idx2 = html.find("ART-1")
    row_snippet = html[idx2 : idx2 + 1200]
    assert "10" in row_snippet


def test_totals_row_not_hidden_by_search_filter(db, client_logged_in):
    sender, city, item = _setup(planned_qty=30)
    _ship_box(sender, city, item, qty=10, box_number="BOX-TOT-3", client=client_logged_in)

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)

    assert "picking-totals-row" in html


def test_picking_list_keeps_item_when_plan_is_fully_in_transit(db, client_logged_in):
    """Даже когда весь план уже едет, строка остается видимой как
    «план (в пути)», но маршрутизация больше не предлагает этот город."""
    sender, city, item = _setup(planned_qty=10)
    _ship_box(sender, city, item, qty=10, box_number="BOX-000003", client=client_logged_in)

    resp = client_logged_in.get("/shipment-plan/")
    html = resp.get_data(as_text=True)
    assert "ART-1" in html
    idx = html.find("ART-1")
    snippet = html[idx : idx + 3000]
    assert ">10<" in snippet
    assert "(10)" in snippet


def test_picking_list_shows_receiving_on_recount_and_sorting_as_unplaced(db, client_logged_in):
    """С момента "Отправить на пересчет" и до завершения приемки товар еще
    не попал в UnplacedStock (см. receiving.complete), но физически уже
    принят на складе — должен считаться "на разбраковке" наравне с обычным
    неразмещенным остатком, а не пропадать из плана как будто его нет."""
    sender, city, item = _setup(planned_qty=30)
    # Пересчет/разбраковка доступны только приемкам из накладной (см.
    # ReceivingDocument.is_from_invoice_import) — иначе send-to-recount
    # откажет.
    doc = ReceivingDocument(
        number="REC-PLAN-1",
        warehouse_id=sender.id,
        supplier="ИП Тестов",
        invoice_file_name="накладная.xlsx",
    )
    db.session.add(doc)
    db.session.commit()
    db.session.add(ReceivingLine(document_id=doc.id, nomenclature_id=item.id, qty=12))
    db.session.commit()

    client_logged_in.post(f"/receiving/{doc.id}/send-to-recount")

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)
    idx = html.find("ART-1")
    snippet = html[idx : idx + 3000]
    assert ">12<" in snippet  # На разбраковке

    client_logged_in.post(f"/receiving/{doc.id}/send-to-sorting")

    line = ReceivingLine.query.filter_by(document_id=doc.id).first()
    client_logged_in.post(f"/receiving/{doc.id}/lines/{line.id}/update-defect", data={"defect_qty": "2"})

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)
    idx = html.find("ART-1")
    snippet = html[idx : idx + 3000]
    assert ">10<" in snippet  # 12 - 2 брака = 10 годного "на разбраковке"


def test_picking_list_ignores_invoice_receiving_still_in_draft(db, client_logged_in):
    """Заявленное в накладной количество еще не является поступившим:
    черновик не должен попадать в колонку «На разбраковку» или остаток."""
    sender, city, item = _setup(planned_qty=30)
    doc = ReceivingDocument(
        number="REC-PLAN-2",
        warehouse_id=sender.id,
        supplier="ИП Тестов",
        invoice_file_name="накладная.xlsx",
    )
    db.session.add(doc)
    db.session.commit()
    db.session.add(ReceivingLine(document_id=doc.id, nomenclature_id=item.id, qty=7))
    db.session.commit()

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)
    idx = html.find("ART-1")
    snippet = html[idx : idx + 3000]
    assert ">7<" not in snippet


def test_picking_list_counts_only_confirmed_invoice_lines_on_recount(db, client_logged_in):
    """На пересчете показываем фактически подтвержденное количество, а не
    все заявленные поставщиком позиции."""
    sender, city, item = _setup(planned_qty=30)
    other = Nomenclature(sku="SKU-D2", barcode="7770000002", name="Не поступил", unit="шт")
    db.session.add(other)
    db.session.commit()
    doc = ReceivingDocument(
        number="REC-PLAN-FACT",
        warehouse_id=sender.id,
        supplier="ИП Тестов",
        invoice_file_name="накладная.xlsx",
        status="recounting",
    )
    db.session.add(doc)
    db.session.commit()
    db.session.add_all(
        [
            ReceivingLine(
                document_id=doc.id,
                nomenclature_id=item.id,
                qty=6,
                expected_qty=10,
                confirmed=True,
            ),
            ReceivingLine(
                document_id=doc.id,
                nomenclature_id=other.id,
                qty=9,
                expected_qty=9,
                confirmed=False,
            ),
        ]
    )
    db.session.commit()

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)
    idx = html.find("ART-1")
    snippet = html[idx : idx + 3000]
    assert ">6<" in snippet


def test_picking_list_ignores_plain_draft_receiving_without_invoice(db, client_logged_in):
    """Обычная приемка в короба (не из накладной) в draft еще может быть не
    досчитана кладовщиком — в отличие от накладной, тут нет отдельного шага
    подтверждения количества, поэтому в "на разбраковке"/остаток она не
    попадает до самого завершения приемки (см. receiving.complete)."""
    sender, city, item = _setup(planned_qty=30)
    doc = ReceivingDocument(
        number="REC-PLAN-3",
        warehouse_id=sender.id,
        supplier="ИП Тестов",
    )
    db.session.add(doc)
    db.session.commit()
    db.session.add(ReceivingLine(document_id=doc.id, nomenclature_id=item.id, qty=9))
    db.session.commit()

    html = client_logged_in.get("/shipment-plan/").get_data(as_text=True)
    idx = html.find("ART-1")
    snippet = html[idx : idx + 3000]
    assert ">9<" not in snippet
