"""Проверка непринятых перемещений на Ozon (receipt_audit): совпавшие
проводим, недостачи исправляем (вернуть короб / поправить количество),
излишки, нераспознанное и «недостача здесь ↔ излишек там» — на ручную проверку."""

import io
import json
import re
from datetime import datetime
from html import unescape

import openpyxl

from wms.extensions import db
from wms.models import (
    Box, BoxItem, MovementCargoPlace, MovementDocument, MovementLine, MovementReceiptDiscrepancy,
    Nomenclature, OzonArticleMapping, User, Warehouse,
)


def _setup_world():
    sender = Warehouse(code="WH-AU0", name="Основной")
    ozon = Warehouse(code="WH-AU1", name="ОЗОН: Тверь", marketplace="ozon", marketplace_city="Тверь")
    db.session.add_all([sender, ozon])
    db.session.commit()
    items = []
    for i, art in enumerate(("ШАПКА-А", "СВИТЕР-Б")):
        item = Nomenclature(sku=f"AU-{i}", barcode=f"41000000{i}", name=f"Товар {art}", unit="шт")
        db.session.add(item)
        items.append(item)
    db.session.commit()
    for item, art in zip(items, ("ШАПКА-А", "СВИТЕР-Б")):
        db.session.add(OzonArticleMapping(barcode=item.barcode, article=art))
    db.session.commit()
    return sender, ozon, items


def _make_doc(number, sender, ozon, boxes, with_places=True, request="2000069157404", gm_start=1):
    """boxes: список {item: qty}; ГМ-коды 1000000000000<N>."""
    doc = MovementDocument(
        number=number, from_warehouse_id=sender.id, to_warehouse_id=ozon.id, status="completed",
        completed_at=datetime.utcnow(), shipped_at=datetime.utcnow(),
        marketplace_request_created_at=datetime.utcnow(), marketplace_request_number=request,
    )
    db.session.add(doc)
    db.session.commit()
    result = []
    for i, content in enumerate(boxes):
        box = Box(box_number=f"BOX-{number}-{i}", warehouse_id=ozon.id, status="open")
        db.session.add(box)
        db.session.commit()
        for item, qty in content.items():
            db.session.add(BoxItem(box_id=box.id, nomenclature_id=item.id, qty=qty))
        db.session.add(MovementLine(document_id=doc.id, box_id=box.id, from_warehouse_id=sender.id))
        gm = f"10000000000{gm_start + i:05d}"
        if with_places:
            db.session.add(MovementCargoPlace(document_id=doc.id, box_id=box.id, box_number=box.box_number,
                                              cargo_barcode=gm, items_json="[]"))
        result.append((box, gm))
    db.session.commit()
    return doc, result


ARTICLES = {"41000000" + "0": "ШАПКА-А", "41000000" + "1": "СВИТЕР-Б"}


def _report(rows):
    """rows: (ГМ, товар, заявлено, принято). Формат отчёта Ozon: лист по ГМ."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчёт по поставке"
    ws.append(["Принято ГМ / Заявлено ГМ", "x"])
    ws.append([])
    ws.append(["Ozon ID", "ШК", "Артикул товара"])
    ws2 = wb.create_sheet("Отчёт по ГМ")
    ws2.append(["Принято ГМ / Заявлено ГМ", "x"])
    ws2.append([])
    ws2.append(["ШК ГМ", "ГМ", "Тип ГМ", "Ozon ID", "ШК товара", "Артикул товара", "Наименование товара",
                "Заявлено (шт.)", "Принято на склад (шт.)"])
    for gm, item, declared, accepted in rows:
        ws2.append([gm, "Коробка", "Моно", 1, item.barcode, ARTICLES[item.barcode], item.name, declared, accepted])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _check(client, files):
    return client.post(
        "/movement-audit/check",
        data={"reports": [(buf, name) for buf, name in files]},
        content_type="multipart/form-data",
    )


def _plan(html):
    return json.loads(unescape(re.search(r'name="plan" value="([^"]*)"', html).group(1)))


def test_matching_document_is_received_as_is(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    doc, boxes = _make_doc("AU-1", sender, ozon, [{hat: 10}, {sweater: 5}])
    html = _check(client_logged_in, [(_report([
        (boxes[0][1], hat, 10, 10), (boxes[1][1], sweater, 5, 5),
    ]), "supply-2000069157404-acceptance-report.xlsx")]).get_data(as_text=True)
    assert "Совпало: 1" in html and "Исправим автоматически: 0" in html

    client_logged_in.post("/movement-audit/apply", data={"plan": json.dumps(_plan(html))})
    doc = MovementDocument.query.get(doc.id)
    assert doc.received_at is not None and doc.receipt_returned_boxes is None
    assert MovementReceiptDiscrepancy.query.count() == 0


def test_shortage_with_missing_box_returns_it_and_trims_other(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    doc, boxes = _make_doc("AU-2", sender, ozon, [{hat: 10}, {sweater: 5}, {hat: 8}])
    html = _check(client_logged_in, [(_report([
        (boxes[0][1], hat, 10, 7),      # пришёл, но не хватает 3
        (boxes[2][1], hat, 8, 8),       # ровно
        # boxes[1] (свитера) в отчёте нет — не доехал
    ]), "x.xlsx")]).get_data(as_text=True)
    assert "Исправим автоматически: 1" in html
    assert "Вернуть на склад-отправитель коробов: <b>1</b>" in html

    client_logged_in.post("/movement-audit/apply", data={"plan": json.dumps(_plan(html))})
    doc = MovementDocument.query.get(doc.id)
    assert doc.received_at is not None and doc.receipt_returned_boxes == boxes[1][0].box_number
    assert Box.query.get(boxes[1][0].id).warehouse_id == sender.id           # вернулся
    assert BoxItem.query.filter_by(box_id=boxes[0][0].id).first().qty == 7     # поправили количество
    assert doc.lines.count() == 2


def test_surplus_goes_to_manual_and_is_not_applied(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    doc, boxes = _make_doc("AU-3", sender, ozon, [{hat: 10}])
    html = _check(client_logged_in, [(_report([(boxes[0][1], hat, 10, 12)]), "x.xlsx")]).get_data(as_text=True)
    assert "Ручная проверка: 1" in html and "Принято больше отправленного" in html
    assert 'name="plan"' not in html     # применять нечего
    assert MovementDocument.query.get(doc.id).received_at is None


def test_shortage_here_and_same_item_surplus_elsewhere_is_manual_for_both(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    a, boxes_a = _make_doc("AU-4A", sender, ozon, [{hat: 10}], request="1000000000001", gm_start=1)
    b, boxes_b = _make_doc("AU-4B", sender, ozon, [{hat: 10}], request="1000000000002", gm_start=11)
    html = _check(client_logged_in, [
        (_report([(boxes_a[0][1], hat, 10, 6)]), "a.xlsx"),    # здесь не хватило 4
        (_report([(boxes_b[0][1], hat, 10, 14)]), "b.xlsx"),   # а тут лишние 4
    ]).get_data(as_text=True)
    assert "Ручная проверка: 2" in html
    assert "Возможна ошибка при отправке" in html and "в том же количестве" in html
    assert 'name="plan"' not in html


def test_shortage_matches_excess_in_recently_received_document(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    old, _ = _make_doc("AU-5OLD", sender, ozon, [{hat: 10}], request="1000000000003", gm_start=21)
    old.received_at = datetime.utcnow()
    db.session.add(MovementReceiptDiscrepancy(document_id=old.id, nomenclature_id=hat.id, expected_qty=10, received_qty=15))
    new, boxes = _make_doc("AU-5NEW", sender, ozon, [{hat: 10}], request="1000000000004", gm_start=31)
    db.session.commit()
    html = _check(client_logged_in, [(_report([(boxes[0][1], hat, 10, 5)]), "n.xlsx")]).get_data(as_text=True)
    assert "AU-5OLD" in html and "Ручная проверка: 1" in html


def test_mapping_inferred_by_declared_content_when_gm_not_saved(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    doc, boxes = _make_doc("AU-6", sender, ozon, [{hat: 10}, {sweater: 5}], with_places=False)
    html = _check(client_logged_in, [(_report([
        ("1000000000999001", hat, 10, 10), ("1000000000999002", sweater, 5, 5),
    ]), "unrelated-name.xlsx")]).get_data(as_text=True)
    assert "Совпало: 1" in html and "по составу" in html


def test_report_for_unknown_supply_is_listed_as_unrecognized(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    doc, boxes = _make_doc("AU-7", sender, ozon, [{hat: 10}], with_places=False)
    html = _check(client_logged_in, [(_report([("1000000000888001", sweater, 99, 99)]), "other.xlsx")]).get_data(as_text=True)
    assert "Отчёт не опознан: 1" in html and "Без отчёта: 1" in html


def test_missing_request_mark_blocks_auto_posting(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    doc, boxes = _make_doc("AU-8", sender, ozon, [{hat: 10}])
    doc.marketplace_request_created_at = None
    db.session.commit()
    html = _check(client_logged_in, [(_report([(boxes[0][1], hat, 10, 10)]), "x.xlsx")]).get_data(as_text=True)
    assert "Ручная проверка: 1" in html and "номер заявки на МП" in html


def test_apply_refuses_if_boxes_changed_after_check(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    doc, boxes = _make_doc("AU-9", sender, ozon, [{hat: 10}])
    html = _check(client_logged_in, [(_report([(boxes[0][1], hat, 10, 10)]), "x.xlsx")]).get_data(as_text=True)
    BoxItem.query.filter_by(box_id=boxes[0][0].id).first().qty = 4     # состав изменили между проверкой и применением
    db.session.commit()
    resp = client_logged_in.post("/movement-audit/apply", data={"plan": json.dumps(_plan(html))}, follow_redirects=True)
    assert "Пропущено" in resp.get_data(as_text=True)
    assert MovementDocument.query.get(doc.id).received_at is None


def test_only_admin_can_open(db, client):
    worker = User(username="au_user", full_name="Не админ", role="warehouse")
    worker.set_password("pw12345")
    db.session.add(worker)
    db.session.commit()
    with client.session_transaction() as sess:
        sess["_user_id"] = str(worker.id)
        sess["_fresh"] = True
    assert client.get("/movement-audit/").status_code == 404
    assert client.post("/movement-audit/apply", data={"plan": "[]"}).status_code == 404


def test_list_shows_button_for_admin(db, client_logged_in):
    assert "Проверка по отчётам (файлы)" in client_logged_in.get("/movement/").get_data(as_text=True)


def _plan_html(client, files):
    return _check(client, files).get_data(as_text=True)


def test_declared_keeps_rows_with_zero_accepted(db):
    """Ревью: заявленное считалось только по принятым строкам, и ГМ с
    непринятой позицией сопоставлялся не с тем коробом."""
    from wms.utils.acceptance_act import parse_act

    sender, ozon, (hat, sweater) = _setup_world()
    buf = _report([("1000000000777001", hat, 5, 5), ("1000000000777001", sweater, 3, 0)])
    boxes, _total, declared = parse_act(buf.getvalue(), prefer_article=True, with_declared=True)
    assert len(declared["1000000000777001"]) == 2           # заявлены обе позиции
    assert sum(boxes["1000000000777001"].values()) == 5      # принята только одна


def test_declared_composition_with_unaccepted_item_maps_to_the_right_box(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    # два короба: {шапка 5} и {шапка 5, свитер 3}; ГМ второго принят только по шапкам
    doc, boxes = _make_doc("AU-10", sender, ozon, [{hat: 5}, {hat: 5, sweater: 3}], with_places=False)
    html = _plan_html(client_logged_in, [(_report([
        ("1000000000666001", hat, 5, 5),
        ("1000000000666002", hat, 5, 5), ("1000000000666002", sweater, 3, 0),
    ]), "n.xlsx")])
    assert "Исправим автоматически: 1" in html
    plan = _plan(html)[0]["payload"]
    assert plan[boxes[1][0].box_number] == {str(hat.id): 5.0}     # свитеры во втором коробе не приняты
    assert plan[boxes[0][0].box_number] == {str(hat.id): 5.0}


def test_apply_ignores_duplicate_document_in_plan(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    doc, boxes = _make_doc("AU-11", sender, ozon, [{hat: 10}])
    html = _plan_html(client_logged_in, [(_report([(boxes[0][1], hat, 10, 6)]), "x.xlsx")])
    item = _plan(html)[0]
    client_logged_in.post("/movement-audit/apply", data={"plan": json.dumps([item, item])})
    assert MovementDocument.query.get(doc.id).received_at is not None
    assert BoxItem.query.filter_by(box_id=boxes[0][0].id).first().qty == 6


def test_apply_rejects_nan_and_infinite_quantities(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    doc, boxes = _make_doc("AU-12", sender, ozon, [{hat: 10}])
    number = boxes[0][0].box_number
    for bad in ("NaN", "Infinity", "-1"):
        raw = '[{"doc_id": %d, "payload": {"%s": {"%d": %s}}}]' % (doc.id, number, hat.id, bad)
        client_logged_in.post("/movement-audit/apply", data={"plan": raw})
        assert MovementDocument.query.get(doc.id).received_at is None
        assert BoxItem.query.filter_by(box_id=boxes[0][0].id).first().qty == 10


def test_identical_boxes_with_inferred_mapping_go_to_manual_when_one_is_missing(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    doc, boxes = _make_doc("AU-13", sender, ozon, [{hat: 10}, {hat: 10}], with_places=False)
    html = _plan_html(client_logged_in, [(_report([
        ("1000000000555001", hat, 10, 10),
        ("1000000000555002", hat, 10, 0),      # второй такой же короб не принят
    ]), "n.xlsx")])
    assert "Ручная проверка: 1" in html and "одинаковые" in html
    assert MovementDocument.query.get(doc.id).received_at is None


def test_request_number_must_match_a_full_number_in_file_name(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    doc, boxes = _make_doc("AU-14", sender, ozon, [{hat: 10}], with_places=False, request="2000069")
    # в имени файла другое, более длинное число, в котором «2000069» — лишь начало
    html = _plan_html(client_logged_in, [(_report([("1000000000444001", sweater, 99, 99)]),
                                          "supply-2000069157404-acceptance-report.xlsx")])
    assert "Отчёт не опознан: 1" in html


def test_two_documents_with_same_request_number_are_not_guessed(db, client_logged_in):
    sender, ozon, (hat, sweater) = _setup_world()
    _make_doc("AU-15A", sender, ozon, [{hat: 10}], with_places=False, request="2000069157404", gm_start=1)
    _make_doc("AU-15B", sender, ozon, [{sweater: 5}], with_places=False, request="2000069157404", gm_start=11)
    html = _plan_html(client_logged_in, [(_report([("1000000000333001", hat, 7, 7)]),
                                          "supply-2000069157404-acceptance-report.xlsx")])
    assert "Отчёт не опознан: 1" in html


def test_one_failing_document_does_not_stop_the_others(db, client_logged_in, monkeypatch):
    from wms.blueprints import receipt_audit as ra

    sender, ozon, (hat, sweater) = _setup_world()
    a, boxes_a = _make_doc("AU-16A", sender, ozon, [{hat: 10}], request="1000000000011", gm_start=1)
    b, boxes_b = _make_doc("AU-16B", sender, ozon, [{sweater: 5}], request="1000000000012", gm_start=11)
    html = _plan_html(client_logged_in, [
        (_report([(boxes_a[0][1], hat, 10, 10)]), "a.xlsx"),
        (_report([(boxes_b[0][1], sweater, 5, 5)]), "b.xlsx"),
    ])
    real = ra._receive_from_act

    def flaky(doc, payload):
        if doc.number == "AU-16A":
            raise RuntimeError("сбой")
        return real(doc, payload)

    monkeypatch.setattr(ra, "_receive_from_act", flaky)
    resp = client_logged_in.post("/movement-audit/apply", data={"plan": json.dumps(_plan(html))}, follow_redirects=True)
    assert "AU-16A (ошибка при проведении" in resp.get_data(as_text=True)
    assert MovementDocument.query.get(a.id).received_at is None
    assert MovementDocument.query.get(b.id).received_at is not None
