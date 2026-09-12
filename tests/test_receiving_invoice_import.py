"""Загрузка приходной накладной из 1С (см. utils/receiving_invoice_import.py
и receiving.import_invoice_form/confirm_invoice/confirm_line): номер
накладной становится номером приемки, поставщик заводится в справочник
автоматически, товары/количество подставляются из файла, а кладовщик
сверяет их на упрощенной мобильной форме (список + количество + галочка)."""

import io

import openpyxl
import pytest

from wms.extensions import db
from wms.models import Nomenclature, ReceivingDocument, ReceivingLine, Supplier, UnplacedStock, Warehouse
from wms.utils.receiving_invoice_import import InvoiceParseError, parse_invoice


DEFAULT_ROW_NAME = "Кардиган бежевый MEVIAR Kids / Шапки (50-54)"


def _build_invoice_xlsx(
    number="1706",
    supplier_line="ИП Кииков Мурат Борисович, ИНН 091701566682, тел.: 89283882790",
    rows=(("НФ-00003575", DEFAULT_ROW_NAME, 80),),
    with_barcode_column=False,
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([f"Приходная накладная № {number} от 10 сентября 2026 г."])
    ws.append(["(Поступление от поставщика) Ш-009  СКЛАД2 ШОССЕЙНАЯ167"])
    ws.append([])
    ws.append([f"Поставщик:   {supplier_line}"])
    ws.append(["Контактное лицо:"])
    ws.append([])
    ws.append(["Покупатель:", 'Общество с ограниченной ответственностью "ТОР-ФАРМ"'])
    ws.append([])
    header = ["№", "Код", "Товары", "Мест", "Количество", "Цена", "Сумма"]
    if with_barcode_column:
        header.append("Штрихкод")
    ws.append(header)
    for i, row in enumerate(rows, start=1):
        code, name, qty = row[0], row[1], row[2]
        line = [i, code, name, None, qty, 220.0, qty * 220.0]
        if with_barcode_column:
            line.append(row[3] if len(row) > 3 else "")
        ws.append(line)
    ws.append([])
    ws.append(["Итого:", "", "", "", "", "", sum(row[2] for row in rows) * 220.0])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_parse_invoice_extracts_header_and_rows():
    file_stream = _build_invoice_xlsx()

    invoice = parse_invoice(file_stream)

    assert invoice.invoice_number == "1706"
    assert invoice.supplier_name.startswith("ИП Кииков Мурат Борисович")
    assert invoice.supplier_inn == "091701566682"
    assert invoice.supplier_phone == "89283882790"
    assert len(invoice.rows) == 1
    assert invoice.rows[0]["code"] == "НФ-00003575"
    assert invoice.rows[0]["qty"] == 80
    assert invoice.rows[0]["barcode"] == ""


def test_parse_invoice_extracts_barcode_column_when_present():
    """Колонки со штрихкодом в реальной выгрузке пока нет, но пользователь
    предупредил, что может ее добавить — парсер должен подхватить ее
    автоматически, если она появится, не требуя ее обязательного наличия."""
    file_stream = _build_invoice_xlsx(
        rows=(("НФ-00003575", DEFAULT_ROW_NAME, 80, "4601234567890"),),
        with_barcode_column=True,
    )

    invoice = parse_invoice(file_stream)

    assert invoice.rows[0]["barcode"] == "4601234567890"


def test_parse_invoice_multiple_rows_stops_at_itogo():
    file_stream = _build_invoice_xlsx(
        rows=(
            ("НФ-00003575", "Товар 1", 80),
            ("НФ-00003576", "Товар 2", 300),
        )
    )

    invoice = parse_invoice(file_stream)

    assert len(invoice.rows) == 2
    assert invoice.rows[1]["qty"] == 300


def test_parse_invoice_raises_on_unrecognized_file():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Просто случайный файл", "без нужной структуры"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    with pytest.raises(InvoiceParseError):
        parse_invoice(buf)


def _build_invoice_html_xls(
    number="1706",
    supplier_line="ИП Кииков Мурат Борисович, ИНН 091701566682, тел.: 89283882790",
    rows=(("НФ-00003575", "Кардиган бежевый MEVIAR Kids / Шапки (50-54)", 80),),
):
    """1С очень часто отдает "выгрузку в Excel" на самом деле HTML-таблицей
    с расширением .xls/.xlsx — openpyxl такой файл открыть не может вообще
    (это не ZIP), поэтому парсер должен распознавать и такой формат тоже."""
    rows_html = "".join(
        f"<tr><td>{i}</td><td>{code}</td><td>{name}</td><td></td><td>{qty}</td><td>220</td><td>{qty * 220}</td></tr>"
        for i, (code, name, qty) in enumerate(rows, start=1)
    )
    html = f"""
    <html><body>
    <table>
      <tr><td colspan="7">Приходная накладная № {number} от 10 сентября 2026 г.</td></tr>
      <tr><td colspan="7">(Поступление от поставщика) Ш-009  СКЛАД2 ШОССЕЙНАЯ167</td></tr>
      <tr><td colspan="7">Поставщик:   {supplier_line}</td></tr>
      <tr><td colspan="7">Контактное лицо:</td></tr>
      <tr><td>№</td><td>Код</td><td>Товары</td><td>Мест</td><td>Количество</td><td>Цена</td><td>Сумма</td></tr>
      {rows_html}
      <tr><td>Итого:</td><td></td><td></td><td></td><td></td><td></td><td>{sum(q for _, _, q in rows) * 220}</td></tr>
    </table>
    </body></html>
    """
    return io.BytesIO(html.encode("utf-8"))


def test_parse_invoice_handles_html_disguised_as_excel():
    """Реальная причина "Internal Server Error" при загрузке — 1С шлет HTML
    под видом .xlsx, а openpyxl.load_workbook падает не InvoiceParseError,
    а низкоуровневым исключением (не ZIP-файл)."""
    file_stream = _build_invoice_html_xls()

    invoice = parse_invoice(file_stream)

    assert invoice.invoice_number == "1706"
    assert invoice.supplier_inn == "091701566682"
    assert len(invoice.rows) == 1
    assert invoice.rows[0]["qty"] == 80


def test_parse_invoice_raises_clean_error_on_garbage_file():
    """Совсем не Excel и не HTML — тоже не должно валиться необработанным
    исключением, только понятной InvoiceParseError."""
    buf = io.BytesIO(b"\x00\x01\x02 not a real file at all \xff\xfe")

    with pytest.raises(InvoiceParseError):
        parse_invoice(buf)


def test_upload_with_html_disguised_invoice_does_not_500(db, client_logged_in):
    warehouse = _make_warehouse()
    item = Nomenclature(sku="НФ-00003575", barcode="8880000099", name=DEFAULT_ROW_NAME, unit="шт")
    db.session.add(item)
    db.session.commit()

    resp = client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (_build_invoice_html_xls(), "invoice.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert resp.status_code == 200
    doc = ReceivingDocument.query.filter_by(number="1706").first()
    assert doc is not None
    assert doc.lines.count() == 1


def test_upload_with_garbage_file_shows_flash_not_500(db, client_logged_in):
    warehouse = _make_warehouse()

    resp = client_logged_in.post(
        "/receiving/import-invoice",
        data={
            "warehouse_id": warehouse.id,
            "file": (io.BytesIO(b"\x00\x01\x02 garbage \xff\xfe"), "invoice.xlsx"),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert resp.status_code == 200
    assert "Не удалось" in resp.get_data(as_text=True)


def _make_warehouse():
    wh = Warehouse(code="WH-INV1", name="Основной склад для накладных")
    db.session.add(wh)
    db.session.commit()
    return wh


def test_upload_creates_document_with_invoice_number_and_supplier(db, client_logged_in):
    warehouse = _make_warehouse()
    item = Nomenclature(sku="НФ-00003575", barcode="8880000001", name=DEFAULT_ROW_NAME, unit="шт")
    db.session.add(item)
    db.session.commit()

    file_stream = _build_invoice_xlsx()
    resp = client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (file_stream, "invoice.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert resp.status_code == 200
    doc = ReceivingDocument.query.filter_by(number="1706").first()
    assert doc is not None
    assert doc.warehouse_id == warehouse.id
    assert doc.supplier_id is not None

    supplier = Supplier.query.get(doc.supplier_id)
    assert supplier.inn == "091701566682"

    line = doc.lines.first()
    assert line.nomenclature_id == item.id
    assert line.qty == 80
    assert line.expected_qty == 80
    assert line.confirmed is False

    # Редиректит сразу на мобильную форму сверки, а не на обычную детальную.
    assert resp.request.path == f"/receiving/{doc.id}/confirm"


def test_upload_saves_order_number_and_invoice_file(db, client_logged_in):
    """№ заявки вносится вручную (в файле 1С его нет), а сам файл сохраняется
    для скачивания позже — см. receiving.download_invoice_file."""
    warehouse = _make_warehouse()
    item = Nomenclature(sku="НФ-00003575", barcode="8880000011", name=DEFAULT_ROW_NAME, unit="шт")
    db.session.add(item)
    db.session.commit()

    file_stream = _build_invoice_xlsx()
    raw_bytes = file_stream.getvalue()
    client_logged_in.post(
        "/receiving/import-invoice",
        data={
            "warehouse_id": warehouse.id,
            "order_number": "ЗАЯВКА-42",
            "file": (io.BytesIO(raw_bytes), "invoice.xlsx"),
        },
        content_type="multipart/form-data",
    )

    doc = ReceivingDocument.query.filter_by(number="1706").first()
    assert doc.order_number == "ЗАЯВКА-42"
    assert doc.invoice_file_name == "invoice.xlsx"
    assert doc.invoice_file_data == raw_bytes


def test_download_invoice_file_returns_saved_bytes(db, client_logged_in):
    warehouse = _make_warehouse()
    item = Nomenclature(sku="НФ-00003575", barcode="8880000012", name=DEFAULT_ROW_NAME, unit="шт")
    db.session.add(item)
    db.session.commit()

    file_stream = _build_invoice_xlsx()
    raw_bytes = file_stream.getvalue()
    client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (io.BytesIO(raw_bytes), "invoice.xlsx")},
        content_type="multipart/form-data",
    )
    doc = ReceivingDocument.query.filter_by(number="1706").first()

    resp = client_logged_in.get(f"/receiving/{doc.id}/invoice-file")

    assert resp.status_code == 200
    assert resp.data == raw_bytes


def test_download_invoice_file_missing_flashes_and_redirects(db, client_logged_in):
    warehouse = _make_warehouse()
    doc = ReceivingDocument(number="NO-FILE-1", warehouse_id=warehouse.id)
    db.session.add(doc)
    db.session.commit()

    resp = client_logged_in.get(f"/receiving/{doc.id}/invoice-file", follow_redirects=True)

    assert "не сохранен" in resp.get_data(as_text=True)


def test_upload_reuses_existing_supplier_by_inn(db, client_logged_in):
    warehouse = _make_warehouse()
    existing = Supplier(name="Старое название", inn="091701566682", phone="000")
    db.session.add(existing)
    db.session.commit()
    item = Nomenclature(sku="НФ-00003575", barcode="8880000002", name=DEFAULT_ROW_NAME, unit="шт")
    db.session.add(item)
    db.session.commit()

    client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (_build_invoice_xlsx(), "invoice.xlsx")},
        content_type="multipart/form-data",
    )

    assert Supplier.query.filter_by(inn="091701566682").count() == 1
    doc = ReceivingDocument.query.filter_by(number="1706").first()
    assert doc.supplier_id == existing.id


def test_upload_rejects_duplicate_invoice_number(db, client_logged_in):
    warehouse = _make_warehouse()
    item = Nomenclature(sku="НФ-00003575", barcode="8880000003", name=DEFAULT_ROW_NAME, unit="шт")
    db.session.add(item)
    db.session.commit()

    client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (_build_invoice_xlsx(), "invoice.xlsx")},
        content_type="multipart/form-data",
    )
    resp = client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (_build_invoice_xlsx(), "invoice.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert "уже была загружена раньше" in resp.get_data(as_text=True)
    assert ReceivingDocument.query.filter_by(number="1706").count() == 1


def test_upload_reports_unmatched_codes_but_keeps_matched(db, client_logged_in):
    warehouse = _make_warehouse()
    item = Nomenclature(sku="НФ-00003575", barcode="8880000004", name="Кардиган тест 4", unit="шт")
    db.session.add(item)
    db.session.commit()

    file_stream = _build_invoice_xlsx(
        rows=(
            ("НФ-00003575", "Кардиган тест 4", 80),
            ("НФ-НЕИЗВЕСТНЫЙ", "Товар без совпадения", 5),
        )
    )
    resp = client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (file_stream, "invoice.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    html = resp.get_data(as_text=True)
    assert "Не найдено в номенклатуре по названию/штрихкоду" in html
    doc = ReceivingDocument.query.filter_by(number="1706").first()
    assert doc.lines.count() == 1


def test_confirm_line_updates_qty_and_confirmed_flag(db, client_logged_in):
    warehouse = _make_warehouse()
    item = Nomenclature(sku="НФ-00003575", barcode="8880000005", name=DEFAULT_ROW_NAME, unit="шт")
    db.session.add(item)
    db.session.commit()
    client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (_build_invoice_xlsx(), "invoice.xlsx")},
        content_type="multipart/form-data",
    )
    doc = ReceivingDocument.query.filter_by(number="1706").first()
    line = doc.lines.first()

    resp = client_logged_in.post(
        f"/receiving/{doc.id}/lines/{line.id}/confirm",
        json={"qty": 78, "confirmed": True},
    )

    data = resp.get_json()
    assert data["ok"] is True
    assert data["confirmed_count"] == 1
    assert data["total_count"] == 1

    line = ReceivingLine.query.get(line.id)
    assert line.qty == 78
    assert line.confirmed is True
    assert line.expected_qty == 80  # исходное количество из накладной не трогаем


def test_completing_confirmed_invoice_credits_unplaced_stock(db, client_logged_in):
    warehouse = _make_warehouse()
    item = Nomenclature(sku="НФ-00003575", barcode="8880000006", name=DEFAULT_ROW_NAME, unit="шт")
    db.session.add(item)
    db.session.commit()
    client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (_build_invoice_xlsx(), "invoice.xlsx")},
        content_type="multipart/form-data",
    )
    doc = ReceivingDocument.query.filter_by(number="1706").first()
    line = doc.lines.first()
    client_logged_in.post(f"/receiving/{doc.id}/lines/{line.id}/confirm", json={"qty": 79, "confirmed": True})

    client_logged_in.post(f"/receiving/{doc.id}/send-to-recount")
    client_logged_in.post(f"/receiving/{doc.id}/send-to-sorting")
    client_logged_in.post(f"/receiving/{doc.id}/complete")

    stock = UnplacedStock.query.filter_by(warehouse_id=warehouse.id, nomenclature_id=item.id).first()
    assert stock is not None
    assert stock.qty == 79


def test_mobile_confirm_invoice_page_can_drive_full_recount_sorting_flow(db, client_logged_in):
    """Раньше на мобильной сверке "По накладной" не было кнопок пересчета/
    разбраковки вообще — "Завершить приемку" была доступна только из
    черновика напрямую. Кладовщик должен уметь пройти весь процесс, не
    уходя на десктопную страницу приемки."""
    from wms.models import SupplierReturn

    warehouse = _make_warehouse()
    item = Nomenclature(sku="НФ-00003575", barcode="8880000007", name=DEFAULT_ROW_NAME, unit="шт")
    db.session.add(item)
    db.session.commit()
    client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (_build_invoice_xlsx(), "invoice.xlsx")},
        content_type="multipart/form-data",
    )
    doc = ReceivingDocument.query.filter_by(number="1706").first()
    line = doc.lines.first()
    client_logged_in.post(f"/receiving/{doc.id}/lines/{line.id}/confirm", json={"qty": 80, "confirmed": True})

    resp = client_logged_in.post(
        f"/receiving/{doc.id}/send-to-recount", data={"next": "confirm"}
    )
    assert resp.status_code == 302
    assert f"/receiving/{doc.id}/confirm" in resp.headers["Location"]
    assert ReceivingDocument.query.get(doc.id).status == "recounting"

    # Пересчет разошелся — поправляем кол-во прямо со сверки.
    client_logged_in.post(f"/receiving/{doc.id}/lines/{line.id}/confirm", json={"qty": 75, "confirmed": True})
    assert ReceivingLine.query.get(line.id).qty == 75

    resp = client_logged_in.post(
        f"/receiving/{doc.id}/send-to-sorting", data={"next": "confirm"}
    )
    assert f"/receiving/{doc.id}/confirm" in resp.headers["Location"]
    assert ReceivingDocument.query.get(doc.id).status == "sorting"

    resp = client_logged_in.post(
        f"/receiving/{doc.id}/lines/{line.id}/update-defect",
        data={"defect_qty": "5", "next": "confirm"},
    )
    assert f"/receiving/{doc.id}/confirm" in resp.headers["Location"]

    resp = client_logged_in.post(f"/receiving/{doc.id}/complete", data={"next": "confirm"})
    assert f"/receiving/{doc.id}/confirm" in resp.headers["Location"]

    doc = ReceivingDocument.query.get(doc.id)
    assert doc.status == "completed"
    stock = UnplacedStock.query.filter_by(warehouse_id=warehouse.id, nomenclature_id=item.id).first()
    assert stock.qty == 70  # 75 - 5 брака
    ret = SupplierReturn.query.filter_by(receiving_document_id=doc.id).first()
    assert ret is not None
    assert ret.qty == 5


def test_confirm_invoice_page_has_add_unlisted_item_button(db, client_logged_in):
    """На мобильной сверке по накладной должна быть возможность добавить
    товар, которого нет в самой накладной — например, поставщик привез
    что-то незаявленное."""
    warehouse = _make_warehouse()
    item = Nomenclature(sku="НФ-00003575", barcode="8880000007", name=DEFAULT_ROW_NAME, unit="шт")
    db.session.add(item)
    db.session.commit()
    client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (_build_invoice_xlsx(), "invoice.xlsx")},
        content_type="multipart/form-data",
    )
    doc = ReceivingDocument.query.filter_by(number="1706").first()

    html = client_logged_in.get(f"/receiving/{doc.id}/confirm").get_data(as_text=True)

    assert "Добавить товар" in html


def test_confirm_invoice_page_shows_document_author(db, client_logged_in, admin_user):
    """Мобильная сверка по накладной должна показывать, кто создал
    документ приемки — это не видно на узком экране больше нигде."""
    warehouse = _make_warehouse()
    item = Nomenclature(sku="НФ-00003575", barcode="8880000010", name=DEFAULT_ROW_NAME, unit="шт")
    db.session.add(item)
    db.session.commit()
    client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (_build_invoice_xlsx(), "invoice.xlsx")},
        content_type="multipart/form-data",
    )
    doc = ReceivingDocument.query.filter_by(number="1706").first()

    html = client_logged_in.get(f"/receiving/{doc.id}/confirm").get_data(as_text=True)

    assert admin_user.display_name() in html


def test_add_unlisted_item_appears_on_confirm_invoice_page_without_expected_qty(db, client_logged_in):
    warehouse = _make_warehouse()
    item = Nomenclature(sku="НФ-00003575", barcode="8880000008", name=DEFAULT_ROW_NAME, unit="шт")
    extra_item = Nomenclature(sku="НФ-EXTRA", barcode="8880000009", name="Незаявленный товар", unit="шт")
    db.session.add_all([item, extra_item])
    db.session.commit()
    client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (_build_invoice_xlsx(), "invoice.xlsx")},
        content_type="multipart/form-data",
    )
    doc = ReceivingDocument.query.filter_by(number="1706").first()

    resp = client_logged_in.post(
        f"/receiving/{doc.id}/lines/add",
        data={"nomenclature_id": extra_item.id, "qty": 3, "next": "confirm"},
        follow_redirects=True,
    )

    assert resp.request.path == f"/receiving/{doc.id}/confirm"
    added_line = ReceivingLine.query.filter_by(document_id=doc.id, nomenclature_id=extra_item.id).first()
    assert added_line is not None
    assert added_line.qty == 3
    assert added_line.expected_qty is None  # не из накладной — сверять не с чем
    assert "Незаявленный товар" in resp.get_data(as_text=True)


def test_upload_matches_by_barcode_when_available_even_if_name_differs(db, client_logged_in):
    """1С «Код» — внутренний артикул поставщика, а не sku в номенклатуре, и
    может не совпадать вообще ни с чем. Если в файле есть штрихкод, он
    сильнее названия (названия в системе и в накладной могут отличаться
    по формулировке, штрихкод — нет)."""
    warehouse = _make_warehouse()
    item = Nomenclature(
        sku="ANY-SKU-1", barcode="4601234567890", name="Совсем другое название в системе", unit="шт"
    )
    db.session.add(item)
    db.session.commit()

    file_stream = _build_invoice_xlsx(
        rows=(("НФ-00003575", DEFAULT_ROW_NAME, 80, "4601234567890"),),
        with_barcode_column=True,
    )
    resp = client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (file_stream, "invoice.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert resp.status_code == 200
    doc = ReceivingDocument.query.filter_by(number="1706").first()
    assert doc is not None
    line = doc.lines.first()
    assert line.nomenclature_id == item.id


def test_upload_matches_by_name_case_insensitively(db, client_logged_in):
    warehouse = _make_warehouse()
    item = Nomenclature(sku="ANY-SKU-2", barcode="8880000007", name=DEFAULT_ROW_NAME.upper(), unit="шт")
    db.session.add(item)
    db.session.commit()

    resp = client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (_build_invoice_xlsx(), "invoice.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert resp.status_code == 200
    doc = ReceivingDocument.query.filter_by(number="1706").first()
    assert doc.lines.first().nomenclature_id == item.id


def test_parse_invoice_finds_supplier_several_columns_to_the_right():
    """На реальном файле реквизиты поставщика оказались не в соседней
    колонке от подписи "Поставщик:" (как в синтетической форме выше), а в
    объединенной ячейке на несколько колонок правее — печатная форма 1С
    щедро объединяет ячейки печатной сетки. Раньше парсер проверял только
    c+1 и не находил реквизиты вообще."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([None, "Приходная накладная № 1333 от 4 августа 2026 г."])
    ws.append([])
    ws.append([None, "Поставщик:", None, None, None, None, "ИП Эркенова Фатима Манафовна,  ИНН 090902545962,  тел.: 89283960999"])
    ws.append([])
    ws.append([None, "№", None, "Код", None, None, None, None, "Товары", None, None, None, None, None, None, None, None, None, None, None, None, None, None, "Штрихкод", None, None, None, None, None, None, None, None, None, None, None, None, "Количество"])
    ws.append([None, "1", None, "НФ-00003474", None, None, None, None, "Товар А", None, None, None, None, None, None, None, None, None, None, None, None, None, None, "2049731406376", None, None, None, None, None, None, None, None, None, None, None, None, 310])
    ws.append([])
    ws.append(["Всего наименований 1"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    invoice = parse_invoice(buf)

    assert invoice.supplier_name == "ИП Эркенова Фатима Манафовна"
    assert invoice.supplier_inn == "090902545962"
    assert invoice.supplier_phone == "89283960999"
    assert len(invoice.rows) == 1
    assert invoice.rows[0]["barcode"] == "2049731406376"


def test_parse_invoice_extracts_barcode_from_trailing_digits_in_name():
    """Часть накладных не выделяет штрихкод отдельной колонкой вообще, а
    печатает его в конце той же ячейки с названием через пробел — см.
    реальный пример "Кардиган_айла_беж ... (42-46) 2051137627033". Без
    отдельной колонки "Штрихкод" сопоставление шло только по названию, а
    название в номенклатуре может не совпадать буквально — штрихкод из
    хвоста строки надежнее."""
    file_stream = _build_invoice_xlsx(
        rows=(("НФ-00003575", f"{DEFAULT_ROW_NAME} 2051137627033", 80),)
    )

    invoice = parse_invoice(file_stream)

    assert invoice.rows[0]["barcode"] == "2051137627033"
    assert invoice.rows[0]["name"] == DEFAULT_ROW_NAME


def test_parse_invoice_ignores_short_trailing_number_in_name():
    """Короткие числа в конце названия (например размер без скобок) не
    похожи на реальный штрихкод (8-14 цифр) — не должны приниматься за
    него по ошибке."""
    file_stream = _build_invoice_xlsx(rows=(("НФ-00003575", "Свитер размер 54", 80),))

    invoice = parse_invoice(file_stream)

    assert invoice.rows[0]["barcode"] == ""
    assert invoice.rows[0]["name"] == "Свитер размер 54"


def test_parse_invoice_does_not_scrape_name_when_barcode_column_exists_but_empty():
    """Если в файле явно есть колонка "Штрихкод", пустая ячейка в этой
    колонке означает, что штрихкода у товара действительно нет — не нужно
    угадывать его из хвоста названия, даже если тот оканчивается цифрами."""
    file_stream = _build_invoice_xlsx(
        rows=(("НФ-00003575", f"{DEFAULT_ROW_NAME} 2051137627033", 80, ""),),
        with_barcode_column=True,
    )

    invoice = parse_invoice(file_stream)

    assert invoice.rows[0]["barcode"] == ""
    assert invoice.rows[0]["name"] == f"{DEFAULT_ROW_NAME} 2051137627033"


def test_upload_matches_by_barcode_extracted_from_name_suffix(db, client_logged_in):
    """Сквозной сценарий: накладная без колонки "Штрихкод", штрихкод —
    хвост из цифр в названии; название в WMS-номенклатуре отличается
    формулировкой (не совпало бы напрямую), но штрихкод сходится."""
    warehouse = _make_warehouse()
    item = Nomenclature(
        sku="ANY-SKU-3", barcode="2051137627033", name="Совсем другая формулировка в системе", unit="шт"
    )
    db.session.add(item)
    db.session.commit()

    file_stream = _build_invoice_xlsx(
        rows=(("НФ-00003575", f"{DEFAULT_ROW_NAME} 2051137627033", 80),)
    )
    resp = client_logged_in.post(
        "/receiving/import-invoice",
        data={"warehouse_id": warehouse.id, "file": (file_stream, "invoice.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert resp.status_code == 200
    doc = ReceivingDocument.query.filter_by(number="1706").first()
    assert doc is not None
    assert doc.lines.first().nomenclature_id == item.id


def test_parse_invoice_date_not_truncated_by_letter_ge_in_month_name():
    """Регекс даты раньше исключал ЛЮБУЮ букву "г"/"Г" до конца строки —
    ломалось на названиях месяцев, содержащих "г" (например, "августа"),
    обрубая дату до "4 ав". Теперь ищем буквальное "г." или "года"."""
    file_stream = _build_invoice_xlsx(number="1333")
    # Подменяем дату на месяц с буквой "г" в середине названия.
    wb = openpyxl.load_workbook(file_stream)
    ws = wb.active
    ws["A1"] = "Приходная накладная № 1333 от 4 августа 2026 г."
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    invoice = parse_invoice(buf)

    assert invoice.invoice_date_text == "4 августа 2026"
