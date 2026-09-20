from datetime import date, datetime, time, timedelta
import io
import json

from openpyxl import Workbook

from wms.extensions import db
from wms.models import (
    AppSetting,
    Box,
    BoxItem,
    MovementDocument,
    MovementLine,
    MovementReceiptDiscrepancy,
    Nomenclature,
    ShipmentPlan,
    ShipmentPlanLine,
    Warehouse,
)

from wms.utils.google_sheets import (
    _current_plan_fact_totals,
    _ensure_output_row_capacity,
    _fact_ranges_for_sheet,
    build_wms_movement_rows,
    distribution_sheet_titles,
    movement_wms_totals,
)


class _ExecuteRecorder:
    def __init__(self):
        self.calls = []

    def spreadsheets(self):
        return self

    def batchUpdate(self, **kwargs):
        self.calls.append(kwargs)
        return self

    def execute(self):
        return {}


def test_distribution_sheet_titles_ignores_dates_and_non_distribution_sheets():
    titles = distribution_sheet_titles(
        [
            "Распределение FBO ВБ от 12.09",
            "Распределение ОЗОН ФБС от 27.08",
            " ВБ ФБС от 27.08",
            "Приоритет",
        ]
    )
    assert titles == [
        "Распределение FBO ВБ от 12.09",
        "Распределение ОЗОН ФБС от 27.08",
    ]


def test_wms_rows_separate_in_transit_and_received(db):
    sender = Warehouse(code="SYNC-FROM", name="Основной")
    target = Warehouse(
        code="SYNC-TO", name="ОЗОН: Москва", marketplace="ozon", marketplace_city="Москва"
    )
    item = Nomenclature(sku="ART-1", barcode="4600000000001", name="Товар", unit="шт")
    db.session.add_all([sender, target, item])
    db.session.commit()

    for number, received in (("MOVE-TRANSIT", False), ("MOVE-RECEIVED", True)):
        box = Box(box_number=f"BOX-{number}", warehouse_id=target.id, status="open")
        document = MovementDocument(
            number=number,
            from_warehouse_id=sender.id,
            to_warehouse_id=target.id,
            status="completed",
            received_at=datetime.utcnow() if received else None,
            marketplace_request_created_at=datetime.utcnow(),
        )
        db.session.add_all([box, document])
        db.session.flush()
        db.session.add_all(
            [
                BoxItem(box_id=box.id, nomenclature_id=item.id, qty=5),
                MovementLine(document_id=document.id, box_id=box.id, from_warehouse_id=sender.id),
            ]
        )
    db.session.commit()

    rows = build_wms_movement_rows()
    assert len(rows) == 1
    assert rows[0][6:] == [5.0, 5.0, 10.0]

    plan = ShipmentPlan(marketplace="ozon")
    db.session.add(plan)
    db.session.flush()
    db.session.add(
        ShipmentPlanLine(
            plan_id=plan.id,
            warehouse_id=target.id,
            nomenclature_id=item.id,
            barcode=item.barcode,
            planned_qty=20,
            # Даже если сохраненный счетчик строки плана устарел, факт
            # восстанавливается непосредственно из перемещений WMS.
            fulfilled_qty=0,
        )
    )
    db.session.commit()

    assert _current_plan_fact_totals()[("ozon", "москва", item.barcode)] == 10


def test_wms_rows_count_completed_movement_without_request_checkbox(db):
    """Для факта отгрузки достаточно завершенного перемещения."""
    sender = Warehouse(code="SYNC-NUM-FROM", name="Основной")
    target = Warehouse(
        code="SYNC-NUM-TO",
        name="ВБ: Казань",
        marketplace="wb",
        marketplace_city="Казань",
    )
    item = Nomenclature(sku="ART-NUM", barcode="4600000000099", name="Товар", unit="шт")
    box = Box(box_number="BOX-REQUEST-NUMBER", warehouse=sender, status="shipped")
    document = MovementDocument(
        number="MOVE-REQUEST-NUMBER",
        from_warehouse=sender,
        to_warehouse=target,
        status="completed",
    )
    db.session.add_all([sender, target, item, box, document])
    db.session.flush()
    db.session.add_all(
        [
            BoxItem(box_id=box.id, nomenclature_id=item.id, qty=7669),
            MovementLine(document_id=document.id, box_id=box.id, from_warehouse_id=sender.id),
        ]
    )
    db.session.commit()

    rows = build_wms_movement_rows()

    assert len(rows) == 1
    assert rows[0][6:] == [7669.0, 0.0, 7669.0]


def test_period_totals_sum_all_cities_from_0001_and_keep_sent_qty_on_shortage(db):
    sender = Warehouse(code="SYNC-DATE-FROM", name="Основной")
    kazan = Warehouse(
        code="SYNC-DATE-KZN", name="ВБ: Казань", marketplace="wb", marketplace_city="Казань"
    )
    samara = Warehouse(
        code="SYNC-DATE-SMR", name="ВБ: Самара", marketplace="wb", marketplace_city="Самара"
    )
    item = Nomenclature(sku="ART-DATE", barcode="4600000000088", name="Товар", unit="шт")
    db.session.add_all([sender, kazan, samara, item])
    db.session.flush()

    period_start = date(2026, 9, 12)

    def add_movement(number, warehouse, qty, completed_at, received_qty=None):
        box = Box(box_number=f"BOX-{number}", warehouse=sender, status="shipped")
        document = MovementDocument(
            number=number,
            from_warehouse=sender,
            to_warehouse=warehouse,
            status="completed",
            completed_at=completed_at,
            received_at=(completed_at + timedelta(days=1)) if received_qty is not None else None,
        )
        db.session.add_all([box, document])
        db.session.flush()
        db.session.add_all(
            [
                BoxItem(box_id=box.id, nomenclature_id=item.id, qty=qty),
                MovementLine(document_id=document.id, box_id=box.id, from_warehouse_id=sender.id),
            ]
        )
        if received_qty is not None:
            db.session.add(
                MovementReceiptDiscrepancy(
                    document_id=document.id,
                    nomenclature_id=item.id,
                    expected_qty=qty,
                    received_qty=received_qty,
                )
            )

    add_movement("MOVE-BEFORE", kazan, 100, datetime.combine(period_start, time(0, 0)))
    add_movement("MOVE-KZN-1", kazan, 20, datetime.combine(period_start, time(0, 1)))
    add_movement(
        "MOVE-KZN-2",
        kazan,
        30,
        datetime.combine(period_start, time(10, 0)),
        received_qty=25,
    )
    add_movement("MOVE-SMR", samara, 7, datetime.combine(period_start, time(12, 0)))
    db.session.commit()

    totals = movement_wms_totals(period_start)

    assert totals[(kazan.id, item.id)]["shipped"] == 50
    assert totals[(kazan.id, item.id)]["received"] == 25
    assert totals[(samara.id, item.id)]["shipped"] == 7


def test_fact_ranges_target_only_shipment_fact_column():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Распределение ВБ"
    sheet.append(["Артикул", "Баркод", "Москва", "отгружено"])
    sheet.append(["A1", "111", 10, 0])
    sheet.append(["ИТОГО", None, 10, "=SUM(D2:D2)"])

    ranges = _fact_ranges_for_sheet(sheet, "wb", {("wb", "москва", "111"): 7})

    assert ranges == [{"range": "'Распределение ВБ'!D2:D3", "values": [[7], [None]]}]


def test_fact_ranges_match_city_aliases_when_writing_back_to_google():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Распределение ВБ"
    sheet.append(["Артикул", "Баркод", "СПБ", "отгружено"])
    sheet.append(["A1", "111", 10, 0])

    ranges = _fact_ranges_for_sheet(
        sheet,
        "wb",
        {("wb", "санкт-петербург", "111"): 9},
    )

    assert ranges == [{"range": "'Распределение ВБ'!D2:D2", "values": [[9]]}]


def test_output_sheet_expands_only_when_required_rows_exceed_grid():
    service = _ExecuteRecorder()
    properties = {"sheetId": 42, "gridProperties": {"rowCount": 2139}}

    assert _ensure_output_row_capacity(service, "sheet-id", properties, 2139) == 2139
    assert service.calls == []

    assert _ensure_output_row_capacity(service, "sheet-id", properties, 2140) == 2140
    assert service.calls[0]["body"]["requests"][0]["updateSheetProperties"][
        "properties"
    ]["gridProperties"]["rowCount"] == 2140


def test_google_trigger_is_public_but_requires_its_own_token(client):
    response = client.post("/shipment-plan/google-trigger")

    assert response.status_code == 401
    assert response.get_json()["ok"] is False


def test_google_trigger_runs_sync_with_valid_token(client, db, monkeypatch):
    db.session.add(AppSetting(key="google_sheets_trigger_token", value="secret-token"))
    db.session.commit()

    monkeypatch.setattr(
        "wms.blueprints.shipment_plan.google_sheets_configured", lambda app: True
    )
    monkeypatch.setattr(
        "wms.blueprints.shipment_plan.sync_google_plans_and_movements",
        lambda: (["ВБ: 12 позиций"], ["Распределение ВБ"], 4, 8),
    )

    response = client.post(
        "/shipment-plan/google-trigger",
        headers={"X-WMS-Sync-Token": "secret-token"},
    )

    assert response.status_code == 200
    payload = response.get_json()
    assert payload["ok"] is True
    assert "ВБ: 12 позиций" in payload["message"]
    assert "обновлено ячеек факта: 8" in payload["message"]


def test_google_button_setup_uses_public_https_address(client_logged_in, app):
    app.config["WMS_PUBLIC_URL"] = "https://wms.wmsmeviar.ru"

    response = client_logged_in.get("/shipment-plan/google-button")

    assert response.status_code == 200
    assert b"https://wms.wmsmeviar.ru/shipment-plan/google-trigger" in response.data
    assert b"syncWms" in response.data


def test_dashboard_shows_google_button_setup_even_without_credentials(
    client_logged_in, monkeypatch
):
    monkeypatch.setattr(
        "wms.blueprints.shipment_plan.google_sheets_configured", lambda app: False
    )

    response = client_logged_in.get("/shipment-plan/")

    assert response.status_code == 200
    assert "/shipment-plan/google-button" in response.get_data(as_text=True)


def test_admin_can_upload_google_credentials_from_setup_page(
    client_logged_in, app, tmp_path
):
    target = tmp_path / "google-service-account.json"
    app.config["GOOGLE_SERVICE_ACCOUNT_FILE"] = str(target)
    credentials = {
        "type": "service_account",
        "project_id": "test-project",
        "private_key": "-----BEGIN PRIVATE KEY-----\ntest\n-----END PRIVATE KEY-----\n",
        "client_email": "wms@example.iam.gserviceaccount.com",
        "token_uri": "https://oauth2.googleapis.com/token",
    }

    response = client_logged_in.post(
        "/shipment-plan/google-button",
        data={
            "action": "upload_credentials",
            "credentials": (
                io.BytesIO(json.dumps(credentials).encode("utf-8")),
                "credentials.json",
            ),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert target.exists()
    assert json.loads(target.read_text(encoding="utf-8"))["project_id"] == "test-project"
    assert "Ключ Google сохранен на сервере" in response.get_data(as_text=True)
