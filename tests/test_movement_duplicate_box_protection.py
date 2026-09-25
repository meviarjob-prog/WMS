"""Один короб не должен попадать в один и тот же документ перемещения
дважды. add_box/route_box_add проверяют это перед вставкой, но два
одновременных запроса могут оба пройти проверку до того, как первый
закоммитится — именно так задвоился короб в документе 202 (см. чат:
"экспорт этого документа показывает 2220, в перемещении 2147" —
total_item_qty() считает короб один раз через SQL SUM с фильтром по
множеству box_id, где дубли естественно схлопываются, а построчный
экспорт по MovementLine удваивал его содержимое).

Защита теперь двухуровневая:
1. Уникальный индекс (document_id, box_id) на movement_lines — бэкстоп на
   уровне БД, ловит и гонку, и любой будущий баг в проверках выше.
2. export_movement_to_excel не показывает содержимое короба дважды, даже
   если в документе (например, в старых данных, заведенных до появления
   индекса) все-таки оказалось два MovementLine на один короб."""

from datetime import datetime

import pytest
from sqlalchemy.exc import IntegrityError

from wms.extensions import db
from wms.models import Box, BoxItem, MovementDocument, MovementLine, Nomenclature, Warehouse
from wms.utils.excel_io import export_movement_to_excel


def _make_doc_with_box(qty=10, suffix="1"):
    sender = Warehouse(code=f"WH-DUPBOX-{suffix}A", name="Отправитель")
    dest = Warehouse(code=f"WH-DUPBOX-{suffix}B", name="Получатель")
    db.session.add_all([sender, dest])
    db.session.commit()

    item = Nomenclature(sku=f"SKU-DUPBOX-{suffix}", barcode=f"999000111{suffix}", name="Товар", unit="шт")
    db.session.add(item)
    db.session.commit()

    box = Box(box_number=f"BOX-DUPBOX-{suffix}", warehouse_id=sender.id, status="open")
    db.session.add(box)
    db.session.commit()
    db.session.add(BoxItem(box_id=box.id, nomenclature_id=item.id, qty=qty))

    doc = MovementDocument(number=f"PER-DUPBOX-{suffix}", from_warehouse_id=sender.id, to_warehouse_id=dest.id)
    db.session.add(doc)
    db.session.commit()
    db.session.add(MovementLine(document_id=doc.id, box_id=box.id, from_warehouse_id=sender.id))
    db.session.commit()
    return doc, box


def test_unique_index_blocks_duplicate_movement_line(db):
    """Бэкстоп на уровне БД: вторая MovementLine на тот же (document, box)
    не должна закоммититься."""
    doc, box = _make_doc_with_box(suffix="A")

    db.session.add(MovementLine(document_id=doc.id, box_id=box.id, from_warehouse_id=box.warehouse_id))
    with pytest.raises(IntegrityError):
        db.session.commit()
    db.session.rollback()

    assert doc.lines.count() == 1


def test_export_does_not_double_count_box_listed_via_two_lines(db):
    """Даже если (например, в старых данных, заведенных до появления
    уникального индекса) короб все-таки оказался привязан к документу
    двумя строками, экспорт не должен удвоить его содержимое —
    total_item_qty() и построчный экспорт обязаны совпадать. Уникальный
    индекс на время теста снят — иначе такую строку было бы уже не
    создать, а именно эту "уже сломанную" ситуацию и нужно проверить."""
    from sqlalchemy import text

    doc, box = _make_doc_with_box(qty=100, suffix="B")

    db.session.execute(text('DROP INDEX IF EXISTS "uq_movement_lines_document_box"'))
    db.session.add(MovementLine(document_id=doc.id, box_id=box.id, from_warehouse_id=box.warehouse_id))
    db.session.commit()

    assert doc.lines.count() == 2
    assert doc.total_item_qty() == 100  # SQL SUM по множеству box_id дублей не считает

    data = export_movement_to_excel([doc])
    import io

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    total_from_export = sum(r[7] for r in rows if isinstance(r[7], (int, float)))

    assert total_from_export == 100  # не 200


def test_add_box_race_gives_friendly_message_not_500(db, client_logged_in, monkeypatch):
    """Симулирует гонку: коммит в add_box падает с IntegrityError (как это
    было бы, если бы уникальный индекс поймал дубль, вставленный вторым,
    параллельным запросом между предварительной проверкой и этим commit())
    — должно быть аккуратное сообщение, а не 500-я ошибка."""
    sender = Warehouse(code="WH-DUPBOX-CA", name="Отправитель")
    dest = Warehouse(code="WH-DUPBOX-CB", name="Получатель")
    db.session.add_all([sender, dest])
    db.session.commit()
    box = Box(box_number="BOX-DUPBOX-C", warehouse_id=sender.id, status="open")
    db.session.add(box)
    db.session.commit()
    doc = MovementDocument(number="PER-DUPBOX-C", from_warehouse_id=sender.id, to_warehouse_id=dest.id)
    db.session.add(doc)
    db.session.commit()

    import wms.blueprints.movement as movement_bp

    original_commit = db.session.commit
    calls = {"n": 0}

    def failing_once_commit():
        calls["n"] += 1
        if calls["n"] == 1:
            raise IntegrityError("simulated race", {}, Exception("UNIQUE constraint failed"))
        return original_commit()

    monkeypatch.setattr(movement_bp.db.session, "commit", failing_once_commit)

    resp = client_logged_in.post(
        f"/movement/{doc.id}/boxes/add",
        data={"box_number": box.box_number},
        follow_redirects=True,
    )

    assert resp.status_code == 200
    assert "кто-то другой" in resp.get_data(as_text=True)
    assert doc.lines.count() == 0  # неудавшаяся вставка ничего не оставила


def test_route_box_add_race_gives_friendly_message_not_500(db, client_logged_in, monkeypatch):
    """То же самое для второго места, откуда добавляют короб в перемещение —
    «Куда везти короб» → добавить (route_box_add)."""
    sender = Warehouse(code="WH-DUPBOX-DA", name="Отправитель")
    dest = Warehouse(code="WH-DUPBOX-DB", name="Получатель")
    db.session.add_all([sender, dest])
    db.session.commit()
    box = Box(box_number="BOX-DUPBOX-D", warehouse_id=sender.id, status="open")
    db.session.add(box)
    db.session.commit()

    import wms.blueprints.movement as movement_bp

    original_commit = db.session.commit
    calls = {"n": 0}

    def failing_once_commit():
        calls["n"] += 1
        if calls["n"] == 1:
            raise IntegrityError("simulated race", {}, Exception("UNIQUE constraint failed"))
        return original_commit()

    monkeypatch.setattr(movement_bp.db.session, "commit", failing_once_commit)

    resp = client_logged_in.post(
        "/movement/route-box/add",
        data={"box_id": box.id, "to_warehouse_id": dest.id},
        follow_redirects=True,
    )

    assert resp.status_code == 200
    assert "уже добавлен" in resp.get_data(as_text=True)
