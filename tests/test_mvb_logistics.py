"""«МВБ Логистика»: отдельный вход, заявки клиентов, индивидуальный штрихкод
на каждый короб и поштучные сканы (забор → склад МВБ → СЦ)."""

from flask import g

from wms.extensions import db
from datetime import datetime

from wms.models import (
    Box, MovementDocument, MovementLine, MvbBox, MvbClient, MvbOrder, MvbPallet, MvbPriceTier, MvbTrip,
    MvbVehicle,
    User, Warehouse,
)


def _user(username, role, client=None, password="password123"):
    user = User(username=username, role=role, is_admin=False, mvb_client_id=client.id if client else None)
    user.set_password(password)
    db.session.add(user)
    db.session.commit()
    return user


def _login(client, user):
    # Фикстура db держит app context на весь тест, а Flask-Login кеширует
    # текущего пользователя в g — при смене пользователя внутри теста кеш
    # надо сбросить, иначе следующий запрос увидит предыдущего.
    g.pop("_login_user", None)
    with client.session_transaction() as sess:
        sess["_user_id"] = str(user.id)
        sess["_fresh"] = True


def _mvb_client(name="ООО Ромашка"):
    c = MvbClient(name=name, address="Москва, ул. Ленина, 1", phone="+79990000000")
    db.session.add(c)
    db.session.commit()
    return c


def _create_order(http, **overrides):
    form = {
        "marketplace": "wb",
        "destination": "Коледино",
        "box_count": "3",
        "delivery_method": "pickup",
        "pickup_address": "Москва, ул. Ленина, 1",
        "planned_date": "2026-10-05",
        "slot_date": "2026-10-06",
        "time_from": "10:00",
        "time_to": "14:00",
    }
    form.update(overrides)
    return http.post("/mvb/orders/new", data=form)


def _vehicle(capacity=10, driver=None):
    v = MvbVehicle(plate="А123ВС77", capacity_boxes=capacity, driver_id=driver.id if driver else None)
    db.session.add(v)
    db.session.commit()
    return v


def _new_trip(http, *directions):
    """Рейс-маршрут по направлениям [(маркетплейс, СЦ), ...] — как кнопка
    «Сформировать рейс» на странице «К отправке»."""
    http.post("/mvb/trips/new", data={"dir": [f"{m}|{d}" for m, d in directions]})
    return MvbTrip.query.order_by(MvbTrip.id.desc()).first()


def _deliver_all(http, trip):
    for stop in list(trip.stops):
        http.post(f"/mvb/trips/{trip.id}/stops/{stop.id}/deliver")


def _trip(http, marketplace="wb", destination="Коледино", driver=None, capacity=10):
    """Рейс (вызывать под пользователем склада): создан и авто назначено."""
    _new_trip(http, (marketplace, destination))
    trip = MvbTrip.query.order_by(MvbTrip.id.desc()).first()
    vehicle = _vehicle(capacity, driver)
    http.post(f"/mvb/trips/{trip.id}/plan", data={
        "vehicle_id": str(vehicle.id), "planned_arrival_at": "2026-10-06T09:00",
    })
    db.session.refresh(trip)
    return trip


def _confirmed_order(http, **overrides):
    _create_order(http, **overrides)
    order = MvbOrder.query.order_by(MvbOrder.id.desc()).first()
    http.post(f"/mvb/orders/{order.id}/confirm")
    return order


def test_mvb_login_page_is_public_and_separate(db, client):
    response = client.get("/mvb/login")
    assert response.status_code == 200
    assert "МВБ Логистика" in response.get_data(as_text=True)

    response = client.get("/mvb/orders")
    assert response.status_code == 302
    assert "/mvb/login" in response.headers["Location"]


def test_mvb_user_logs_in_via_mvb_page_but_not_wms(db, client):
    _user("driver1", "mvb_driver")

    response = client.post("/login", data={"username": "driver1", "password": "password123"})
    assert response.status_code == 302
    assert "/mvb/login" in response.headers["Location"]

    response = client.post("/mvb/login", data={"username": "driver1", "password": "password123"})
    assert response.status_code == 302
    assert response.headers["Location"].endswith("/mvb/")


def test_wms_user_cannot_login_to_mvb_or_open_it(db, client):
    wms_user = _user("storekeeper", "warehouse")

    response = client.post("/mvb/login", data={"username": "storekeeper", "password": "password123"})
    assert "Неверный логин или пароль" in response.get_data(as_text=True)

    _login(client, wms_user)
    response = client.get("/mvb/orders")
    assert response.status_code == 302
    assert "/mvb" not in response.headers["Location"]


def test_mvb_user_cannot_open_wms_sections(db, client):
    _login(client, _user("staff1", "mvb_admin"))
    for url in ("/", "/nomenclature/", "/movement/", "/users"):
        response = client.get(url)
        assert response.status_code == 302, url
        assert "/mvb/" in response.headers["Location"], url


def test_client_creates_and_confirms_order_with_unique_barcode_per_box(db, client):
    rom = _mvb_client()
    _login(client, _user("client1", "mvb_client", rom))

    response = _create_order(client)
    assert response.status_code == 302
    order = MvbOrder.query.one()
    assert order.client_id == rom.id
    assert order.status == "draft"
    assert order.boxes == []

    client.post(f"/mvb/orders/{order.id}/confirm")
    db.session.refresh(order)
    assert order.status == "confirmed"
    barcodes = [b.barcode for b in order.boxes]
    assert barcodes == [f"{order.number}-001", f"{order.number}-002", f"{order.number}-003"]
    assert len(set(barcodes)) == 3

    pdf = client.get(f"/mvb/orders/{order.id}/labels.pdf")
    assert pdf.status_code == 200
    assert pdf.mimetype == "application/pdf"


def test_pickup_requires_address(db, client):
    _login(client, _user("client1", "mvb_client", _mvb_client()))
    response = _create_order(client, pickup_address="")
    assert response.status_code == 200
    assert "укажите адрес" in response.get_data(as_text=True)
    assert MvbOrder.query.count() == 0


def test_client_sees_only_own_orders(db, client):
    a, b = _mvb_client("A"), _mvb_client("B")
    client_a = _user("ca", "mvb_client", a)
    client_b = _user("cb", "mvb_client", b)

    _login(client, client_a)
    order_a = _confirmed_order(client)

    _login(client, client_b)
    assert client.get(f"/mvb/orders/{order_a.id}").status_code == 404
    assert client.get(f"/mvb/orders/{order_a.id}/labels.pdf").status_code == 404
    assert f">{order_a.number}</a>" not in client.get("/mvb/orders").get_data(as_text=True)


def test_box_scanned_through_all_stages(db, client):
    rom = _mvb_client()
    _login(client, _user("client1", "mvb_client", rom))
    order = _confirmed_order(client)
    box = order.boxes[0]

    driver = _user("driver1", "mvb_driver")
    staff = _user("staff1", "mvb_admin")

    _login(client, driver)
    data = client.post("/mvb/scan/pickup", data={"barcode": box.barcode}).get_json()
    assert data["ok"] and not data["already"] and data["done"] == 1
    # повторный скан того же короба — не ошибка и не двойной учет
    data = client.post("/mvb/scan/pickup", data={"barcode": box.barcode.lower()}).get_json()
    assert data["ok"] and data["already"]
    # водителю приемка на складе недоступна
    assert client.post("/mvb/scan/receive", data={"barcode": box.barcode}).status_code == 403

    _login(client, staff)
    assert client.post("/mvb/scan/receive", data={"barcode": box.barcode}).get_json()["ok"]
    trip = _trip(client, driver=driver)
    assert client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": box.barcode}).get_json()["ok"]
    client.post(f"/mvb/trips/{trip.id}/depart")

    _login(client, driver)
    _deliver_all(client, trip)

    db.session.refresh(box)
    assert box.status == "delivered"
    assert box.picked_up_at and box.received_at and box.loaded_at and box.shipped_at and box.delivered_at
    assert [e.status for e in box.events] == ["picked_up", "received", "loaded", "shipped", "delivered"]
    # остальные короба заявки не тронуты — видно, какие именно забрали
    assert {b.status for b in order.boxes[1:]} == {"created"}


def test_scan_rejects_wrong_order_of_stages_and_unknown_boxes(db, client):
    _login(client, _user("client1", "mvb_client", _mvb_client()))
    order = _confirmed_order(client)
    _login(client, _user("staff1", "mvb_admin"))

    # короб еще не принят на складе — в рейс его не погрузить
    trip = _trip(client)
    response = client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": order.boxes[0].barcode})
    assert response.status_code == 409
    assert client.post("/mvb/scan/ship", data={"barcode": order.boxes[0].barcode}).status_code == 404
    assert client.post("/mvb/scan/receive", data={"barcode": "NOPE-1"}).status_code == 404


def test_self_delivery_boxes_cannot_be_picked_up_but_are_received(db, client):
    _login(client, _user("client1", "mvb_client", _mvb_client()))
    order = _confirmed_order(client, delivery_method="self", pickup_address="")
    barcode = order.boxes[0].barcode

    _login(client, _user("driver1", "mvb_driver"))
    assert client.post("/mvb/scan/pickup", data={"barcode": barcode}).status_code == 409

    _login(client, _user("staff1", "mvb_admin"))
    assert client.post("/mvb/scan/receive", data={"barcode": barcode}).get_json()["ok"]


def test_draft_and_cancelled_orders_cannot_be_scanned_or_cancelled_after_scan(db, client):
    _login(client, _user("client1", "mvb_client", _mvb_client()))
    order = _confirmed_order(client)
    barcode = order.boxes[0].barcode

    _login(client, _user("staff1", "mvb_admin"))
    client.post("/mvb/scan/receive", data={"barcode": barcode})
    client.post(f"/mvb/orders/{order.id}/cancel")
    db.session.refresh(order)
    assert order.status == "confirmed"


def test_progress_label_shows_how_many_boxes_moved(db, client):
    _login(client, _user("client1", "mvb_client", _mvb_client()))
    order = _confirmed_order(client)
    assert order.progress_label() == "Ожидает передачи"

    _login(client, _user("staff1", "mvb_admin"))
    client.post("/mvb/scan/receive", data={"barcode": order.boxes[0].barcode})
    db.session.refresh(order)
    assert order.progress_label() == "На складе МВБ: 1 из 3"


def test_mvb_admin_creates_client_and_users(db, client):
    _login(client, _user("boss", "mvb_admin"))
    client.post("/mvb/admin/clients", data={"name": "ИП Иванов", "address": "Казань"})
    ivanov = MvbClient.query.filter_by(name="ИП Иванов").one()

    client.post("/mvb/admin/users", data={
        "username": "ivanov", "password": "secret1", "role": "mvb_client", "client_id": str(ivanov.id),
    })
    user = User.query.filter_by(username="ivanov").one()
    assert user.role == "mvb_client" and user.mvb_client_id == ivanov.id and not user.is_admin

    # без клиента роль «Клиент» не создается
    client.post("/mvb/admin/users", data={"username": "nobody", "password": "secret1", "role": "mvb_client"})
    assert User.query.filter_by(username="nobody").first() is None


def test_non_admin_mvb_users_cannot_manage(db, client):
    _login(client, _user("staff1", "mvb_staff"))
    response = client.post("/mvb/admin/users", data={"username": "x", "password": "secret1", "role": "mvb_admin"})
    assert response.status_code == 302
    assert User.query.filter_by(username="x").first() is None


def test_wms_admin_has_access_and_mvb_users_hidden_from_wms_settings(db, client_logged_in):
    _user("driver1", "mvb_driver")
    assert client_logged_in.get("/mvb/orders").status_code == 200
    assert "driver1" not in client_logged_in.get("/users").get_data(as_text=True)


def test_pages_render(db, client_logged_in):
    rom = _mvb_client()
    client_logged_in.post("/mvb/orders/new", data={
        "client_id": str(rom.id), "marketplace": "ozon", "destination": "Хоругвино", "box_count": "2",
        "delivery_method": "pickup", "pickup_address": "Москва", "slot_date": "2026-10-07",
    })
    order = MvbOrder.query.one()
    for url in (
        "/mvb/orders", "/mvb/orders/new", f"/mvb/orders/{order.id}", f"/mvb/orders/{order.id}/edit",
        "/mvb/driver", "/mvb/scan/pickup", "/mvb/scan/receive", "/mvb/admin/clients", "/mvb/admin/users",
    ):
        assert client_logged_in.get(url).status_code == 200, url
    client_logged_in.post(f"/mvb/orders/{order.id}/confirm")
    assert client_logged_in.get(f"/mvb/orders/{order.id}").status_code == 200
    assert client_logged_in.get("/mvb/driver").status_code == 200
    assert MvbBox.query.count() == 2


# ---------- этап 2: транспорт, паллеты, рейсы, пропуск ----------


def _received_order(http, client_user, staff, **overrides):
    _login(http, client_user)
    order = _confirmed_order(http, **overrides)
    _login(http, staff)
    for box in order.boxes:
        http.post("/mvb/scan/receive", data={"barcode": box.barcode})
    return order


def test_trip_lifecycle_with_plan_and_fact(db, client):
    rom = _mvb_client()
    client_user = _user("client1", "mvb_client", rom)
    staff = _user("staff1", "mvb_admin")
    driver = _user("driver1", "mvb_driver")
    order = _received_order(client, client_user, staff)

    trip = _new_trip(client, ("wb", "Коледино"))
    assert trip.status == "searching" and trip.planned_boxes == 3
    # без авто погрузка закрыта
    assert client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": order.boxes[0].barcode}).status_code == 409

    vehicle = _vehicle(capacity=2, driver=driver)
    client.post(f"/mvb/trips/{trip.id}/plan", data={
        "vehicle_id": str(vehicle.id), "planned_arrival_at": "2026-10-06T09:00",
        "planned_load_start_at": "2026-10-06T09:15", "planned_load_end_at": "2026-10-06T10:00",
    })
    db.session.refresh(trip)
    assert trip.status == "assigned"
    assert trip.driver_id == driver.id  # водитель подставлен из авто
    assert trip.planned_arrival_at.hour == 6  # 09:00 МСК хранится как 06:00 UTC

    client.post(f"/mvb/trips/{trip.id}/arrive")
    db.session.refresh(trip)
    assert trip.status == "arrived" and trip.arrived_at

    r1 = client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": order.boxes[0].barcode}).get_json()
    assert r1["ok"] and r1["count"] == 1 and r1["warning"] is None
    db.session.refresh(trip)
    assert trip.status == "loading" and trip.load_started_at
    client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": order.boxes[1].barcode})
    r3 = client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": order.boxes[2].barcode}).get_json()
    assert r3["ok"] and "вместимости" in r3["warning"]

    client.post(f"/mvb/trips/{trip.id}/depart")
    db.session.refresh(trip)
    assert trip.status == "departed" and trip.load_finished_at and trip.departed_at
    assert {b.status for b in order.boxes} == {"shipped"}

    _login(client, driver)
    assert client.get("/mvb/driver").status_code == 200
    _deliver_all(client, trip)
    db.session.refresh(trip)
    assert trip.status == "delivered"
    assert {b.status for b in order.boxes} == {"delivered"}


def test_driver_cannot_operate_other_trips_or_staff_actions(db, client):
    staff = _user("staff1", "mvb_admin")
    driver = _user("driver1", "mvb_driver")
    other = _user("driver2", "mvb_driver")
    _login(client, staff)
    trip = _trip(client, driver=driver)

    _login(client, other)
    assert client.get(f"/mvb/trips/{trip.id}").status_code == 404
    _login(client, driver)
    assert client.post(f"/mvb/trips/{trip.id}/start_loading").status_code == 403
    assert client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": "x"}).status_code == 403


def test_trip_rejects_other_direction_and_cancel_returns_boxes(db, client):
    client_user = _user("client1", "mvb_client", _mvb_client())
    staff = _user("staff1", "mvb_admin")
    order = _received_order(client, client_user, staff, marketplace="ozon", destination="Хоругвино")
    trip = _trip(client, marketplace="wb", destination="Коледино")
    assert client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": order.boxes[0].barcode}).status_code == 409

    ozon_trip = _trip(client, marketplace="ozon", destination="хоругвино")
    assert client.post(f"/mvb/trips/{ozon_trip.id}/scan", data={"barcode": order.boxes[0].barcode}).get_json()["ok"]
    client.post(f"/mvb/trips/{ozon_trip.id}/cancel")
    box = db.session.get(MvbBox, order.boxes[0].id)
    assert box.status == "received" and box.trip_id is None and box.loaded_at is None


def test_pallet_scan_and_load_whole_pallet(db, client):
    client_user = _user("client1", "mvb_client", _mvb_client())
    staff = _user("staff1", "mvb_admin")
    order = _received_order(client, client_user, staff)
    other = _received_order(client, client_user, staff, marketplace="ozon", box_count="1")

    client.post("/mvb/pallets", data={"marketplace": "wb", "destination": "Коледино"})
    pallet = MvbPallet.query.one()
    for box in order.boxes[:2]:
        assert client.post(f"/mvb/pallets/{pallet.id}/scan", data={"barcode": box.barcode}).get_json()["ok"]
    # другое направление на паллету не встает
    assert client.post(f"/mvb/pallets/{pallet.id}/scan", data={"barcode": other.boxes[0].barcode}).status_code == 409
    assert client.get(f"/mvb/pallets/{pallet.id}/label.pdf").mimetype == "application/pdf"
    assert client.get("/mvb/pallets").status_code == 200
    assert client.get(f"/mvb/pallets/{pallet.id}").status_code == 200

    trip = _trip(client)
    data = client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": pallet.number}).get_json()
    assert data["ok"] and data["count"] == 2
    assert {b.status for b in order.boxes[:2]} == {"loaded"}
    assert order.boxes[2].status == "received"


def test_dispatch_groups_ready_boxes_by_direction_fifo(db, client):
    client_user = _user("client1", "mvb_client", _mvb_client())
    staff = _user("staff1", "mvb_admin")
    _received_order(client, client_user, staff)
    _received_order(client, client_user, staff, marketplace="ozon", destination="Хоругвино", box_count="2")
    html = client.get("/mvb/dispatch").get_data(as_text=True)
    assert "Коледино" in html and "Хоругвино" in html
    assert html.index("Коледино") < html.index("Хоругвино")  # раньше принятые — выше


def test_driver_sees_assigned_and_unassigned_pickups_only(db, client):
    client_user = _user("client1", "mvb_client", _mvb_client())
    staff = _user("staff1", "mvb_admin")
    d1, d2 = _user("driver1", "mvb_driver"), _user("driver2", "mvb_driver")
    _login(client, client_user)
    mine = _confirmed_order(client)
    theirs = _confirmed_order(client)
    free = _confirmed_order(client)

    _login(client, staff)
    client.post(f"/mvb/orders/{mine.id}/driver", data={"driver_id": str(d1.id)})
    client.post(f"/mvb/orders/{theirs.id}/driver", data={"driver_id": str(d2.id)})

    _login(client, d1)
    client.get("/mvb/driver")  # первый показ забирает накопившиеся flash-сообщения
    html = client.get("/mvb/driver").get_data(as_text=True)
    assert mine.number in html and free.number in html and theirs.number not in html
    # водитель не может назначать
    client.post(f"/mvb/orders/{free.id}/driver", data={"driver_id": str(d1.id)})
    assert db.session.get(MvbOrder, free.id).driver_id is None


def test_vehicles_page_staff_only(db, client):
    _login(client, _user("staff1", "mvb_admin"))
    client.post("/mvb/vehicles", data={"plate": "а001аа77", "capacity_boxes": "40"})
    assert MvbVehicle.query.one().plate == "А001АА77"
    _login(client, _user("client1", "mvb_client", _mvb_client()))
    assert client.get("/mvb/vehicles").status_code == 302


def test_stage2_pages_render(db, client_logged_in):
    for url in ("/mvb/dispatch", "/mvb/trips", "/mvb/trips?status=all", "/mvb/pallets", "/mvb/vehicles"):
        assert client_logged_in.get(url).status_code == 200, url
    trip = _new_trip(client_logged_in, ("wb", "Коледино"), ("ozon", "Хоругвино"))
    assert len(trip.stops) == 2
    assert client_logged_in.get(f"/mvb/trips/{trip.id}").status_code == 200


# ---------- этап 3: маршрут по точкам, лента водителя, короба из WMS ----------


def test_multi_stop_route_driver_delivers_each_point(db, client):
    client_user = _user("client1", "mvb_client", _mvb_client())
    staff = _user("staff1", "mvb_admin")
    driver = _user("driver1", "mvb_driver")
    wb = _received_order(client, client_user, staff, box_count="2")
    oz = _received_order(client, client_user, staff, marketplace="ozon", destination="Хоругвино", box_count="1")
    other = _received_order(client, client_user, staff, marketplace="wb", destination="Казань", box_count="1")

    trip = _new_trip(client, ("wb", "Коледино"), ("ozon", "Хоругвино"))
    assert [s.label() for s in trip.stops] == ["Wildberries · Коледино", "Ozon · Хоругвино"]
    vehicle = _vehicle(10, driver)
    client.post(f"/mvb/trips/{trip.id}/plan", data={"vehicle_id": str(vehicle.id)})
    for box in wb.boxes + oz.boxes:
        assert client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": box.barcode}).get_json()["ok"]
    # направления нет в маршруте — отказ с подсказкой
    r = client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": other.boxes[0].barcode})
    assert r.status_code == 409 and "добавьте точку" in r.get_json()["message"]
    # точку можно добавить и тогда короб грузится; пустая точка убирается при отправке
    client.post(f"/mvb/trips/{trip.id}/stops", data={"marketplace": "wb", "destination": "Электросталь"})
    db.session.refresh(trip)
    assert len(trip.stops) == 3
    client.post(f"/mvb/trips/{trip.id}/depart")
    db.session.refresh(trip)
    assert len(trip.stops) == 2 and trip.status == "departed"

    first, second = trip.stops
    _login(client, driver)
    client.get("/mvb/driver")
    client.post(f"/mvb/trips/{trip.id}/stops/{first.id}/deliver")
    db.session.refresh(trip)
    assert trip.status == "departed"
    assert {b.status for b in wb.boxes} == {"delivered"} and oz.boxes[0].status == "shipped"
    client.post(f"/mvb/trips/{trip.id}/stops/{second.id}/deliver")
    db.session.refresh(trip)
    assert trip.status == "delivered" and oz.boxes[0].status == "delivered"


def test_stop_order_can_be_changed(db, client):
    _login(client, _user("staff1", "mvb_admin"))
    trip = _new_trip(client, ("wb", "Коледино"), ("ozon", "Хоругвино"))
    second = trip.stops[1]
    client.post(f"/mvb/trips/{trip.id}/stops/{second.id}/up")
    db.session.refresh(trip)
    assert trip.stops[0].id == second.id


def test_driver_feed_shows_free_space_and_take(db, client):
    client_user = _user("client1", "mvb_client", _mvb_client())
    driver = _user("driver1", "mvb_driver")
    _vehicle(capacity=5, driver=driver)
    _login(client, client_user)
    small = _confirmed_order(client, box_count="2")
    big = _confirmed_order(client, box_count="8")

    _login(client, driver)
    client.get("/mvb/driver")
    html = client.get("/mvb/driver").get_data(as_text=True)
    assert "свободно <b>5</b>" in html
    assert "✓ помещается" in html and "✗ не помещается" in html
    assert "новая" in html

    client.post(f"/mvb/driver/orders/{small.id}/take")
    assert db.session.get(MvbOrder, small.id).driver_id == driver.id
    for box in small.boxes:
        data = client.post("/mvb/scan/pickup", data={"barcode": box.barcode}).get_json()
    assert data["warning"] == "В машине 2 из 5 кор."
    client.get("/mvb/driver")
    assert "свободно <b>3</b>" in client.get("/mvb/driver").get_data(as_text=True)
    assert big.number in html


def test_driver_scan_of_free_order_assigns_it(db, client):
    _login(client, _user("client1", "mvb_client", _mvb_client()))
    order = _confirmed_order(client)
    driver = _user("driver1", "mvb_driver")
    _login(client, driver)
    client.post("/mvb/scan/pickup", data={"barcode": order.boxes[0].barcode})
    assert db.session.get(MvbOrder, order.id).driver_id == driver.id


def _wms_movement(number="PER-000777", boxes=2, request_number="WB-1"):
    sender = Warehouse(code=f"S-{number}", name="Основной склад", address="Москва, Складская 5")
    dest = Warehouse(code=f"D-{number}", name="WB Коледино", marketplace="wb", marketplace_city="Коледино")
    db.session.add_all([sender, dest])
    db.session.flush()
    doc = MovementDocument(
        number=number, from_warehouse_id=sender.id, to_warehouse_id=dest.id, status="completed",
        completed_at=datetime.utcnow(),
        marketplace_request_number=request_number,
        marketplace_request_created_at=datetime.utcnow() if request_number else None,
    )
    db.session.add(doc)
    db.session.flush()
    wms_boxes = []
    for i in range(boxes):
        box = Box(box_number=f"BOX-{number[-3:]}{i:03d}", warehouse_id=sender.id)
        db.session.add(box)
        db.session.flush()
        db.session.add(MovementLine(document_id=doc.id, box_id=box.id))
        wms_boxes.append(box)
    db.session.commit()
    return doc, wms_boxes


def test_wms_movement_import_keeps_wms_barcodes(db, client):
    _login(client, _user("staff1", "mvb_admin"))
    doc, wms_boxes = _wms_movement()
    assert doc.number in client.get("/mvb/wms").get_data(as_text=True)

    client.post("/mvb/wms/import", data={"doc_id": str(doc.id), "delivery_method": "self"})
    order = MvbOrder.query.filter_by(wms_movement_id=doc.id).one()
    assert order.client.is_internal and order.marketplace == "wb" and order.destination == "Коледино"
    assert [b.barcode for b in order.boxes] == [b.barcode_value for b in wms_boxes]
    # повторно не импортируется
    client.post("/mvb/wms/import", data={"doc_id": str(doc.id)})
    assert MvbOrder.query.filter_by(wms_movement_id=doc.id).count() == 1
    # скан этикетки WMS (цифры) и ручной ввод номера BOX- находят короб МВБ
    assert client.post("/mvb/scan/receive", data={"barcode": wms_boxes[0].barcode_value}).get_json()["ok"]
    assert client.post("/mvb/scan/receive", data={"barcode": wms_boxes[1].box_number}).get_json()["ok"]


def test_receive_scan_of_wms_box_auto_imports_movement_and_trip_marks_wms_shipped(db, client):
    staff = _user("staff1", "mvb_admin")
    _login(client, staff)
    doc, wms_boxes = _wms_movement()

    data = client.post("/mvb/scan/receive", data={"barcode": wms_boxes[0].barcode_value}).get_json()
    assert data["ok"] and doc.number in data["message"]
    order = MvbOrder.query.filter_by(wms_movement_id=doc.id).one()
    client.post("/mvb/scan/receive", data={"barcode": wms_boxes[1].barcode_value})

    trip = _trip(client)
    for box in wms_boxes:
        assert client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": box.barcode_value}).get_json()["ok"]
    client.post(f"/mvb/trips/{trip.id}/depart")
    assert db.session.get(MovementDocument, doc.id).shipped_at is not None
    assert {b.status for b in order.boxes} == {"shipped"}


def test_wms_shipped_not_set_without_marketplace_request_or_partial(db, client):
    _login(client, _user("staff1", "mvb_admin"))
    doc, wms_boxes = _wms_movement(request_number=None)
    client.post("/mvb/wms/import", data={"doc_id": str(doc.id)})
    for box in wms_boxes:
        client.post("/mvb/scan/receive", data={"barcode": box.barcode_value})
    trip = _trip(client)
    for box in wms_boxes:
        client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": box.barcode_value})
    client.post(f"/mvb/trips/{trip.id}/depart")
    assert db.session.get(MovementDocument, doc.id).shipped_at is None

    doc2, boxes2 = _wms_movement(number="PER-000888")
    client.post(f"/mvb/wms/{doc2.id}/import")
    for box in boxes2:
        client.post("/mvb/scan/receive", data={"barcode": box.barcode_value})
    trip2 = _trip(client)
    client.post(f"/mvb/trips/{trip2.id}/scan", data={"barcode": boxes2[0].barcode_value})
    client.post(f"/mvb/trips/{trip2.id}/depart")
    assert db.session.get(MovementDocument, doc2.id).shipped_at is None  # уехал не весь


# ---------- сквозная цепочка (как описал владелец) ----------


def test_full_chain_seller_to_sc(db, client):
    """Селлер создает заявку → оператор видит её и назначает водителя из
    списка → водитель по дороге видит другие заявки и берет ту, что влезает
    → короба сканируются при заборе → приемка на складе → программа считает
    машины и формирует рейсы по наполненности (статус «Поиск авто») →
    погрузка сканом каждого короба → водитель на точках отмечает
    «Сдано» / «Не сдано»."""
    seller = _user("seller", "mvb_client", _mvb_client("ИП Селлер"))
    seller2 = _user("seller2", "mvb_client", _mvb_client("ООО Второй"))
    operator = _user("operator", "mvb_staff")
    keeper = _user("keeper", "mvb_storekeeper")
    driver = _user("driver", "mvb_driver")
    _vehicle(capacity=6, driver=driver)

    # 1. селлеры создают заявки на забор
    _login(client, seller)
    first = _confirmed_order(client, box_count="3", destination="Коледино")
    _login(client, seller2)
    second = _confirmed_order(client, box_count="2", marketplace="ozon", destination="Хоругвино")
    too_big = _confirmed_order(client, box_count="9", destination="Коледино")

    # 2. оператор видит заявки без водителя и назначает водителя из списка
    _login(client, operator)
    client.get("/mvb/orders")
    html = client.get("/mvb/orders").get_data(as_text=True)
    assert "без водителя: <b>3</b>" in html and driver.display_name() in html
    client.post(f"/mvb/orders/{first.id}/driver", data={"driver_id": str(driver.id), "back": "orders"})
    assert db.session.get(MvbOrder, first.id).driver_id == driver.id

    # 3. водитель в пути: видит свою и чужие свободные заявки, что влезает
    _login(client, driver)
    client.get("/mvb/driver")
    feed = client.get("/mvb/driver").get_data(as_text=True)
    assert first.number in feed and second.number in feed and too_big.number in feed
    for box in first.boxes:
        client.post("/mvb/scan/pickup", data={"barcode": box.barcode})
    client.get("/mvb/driver")
    feed = client.get("/mvb/driver").get_data(as_text=True)
    assert "свободно <b>3</b>" in feed  # 3 из 6 заняты
    # вторая заявка (2 кор.) помещается — берет её по дороге
    client.post(f"/mvb/driver/orders/{second.id}/take")
    for box in second.boxes:
        data = client.post("/mvb/scan/pickup", data={"barcode": box.barcode}).get_json()
    assert data["warning"] == "В машине 5 из 6 кор."

    # 4. водитель привез короба на склад — одна кнопка «Короба сданы в МВБ»
    assert "Короба сданы в МВБ" in client.get("/mvb/driver").get_data(as_text=True)
    client.post("/mvb/driver/handover")
    assert {b.status for b in first.boxes + second.boxes} == {"received"}
    # приемка — у кладовщика, оператору и водителю она не нужна
    assert client.post("/mvb/scan/receive", data={"barcode": first.boxes[0].barcode}).status_code == 403
    _login(client, operator)
    assert client.post("/mvb/scan/receive", data={"barcode": first.boxes[0].barcode}).status_code == 403

    # 5. программа считает машины и компонует рейсы по наполненности,
    # не разбивая заявки: 3 + 2 кор. в машины по 4 — две машины
    html = client.get("/mvb/dispatch?capacity=4").get_data(as_text=True)
    assert "нужно машин по 4 кор.: <b>2</b>" in html
    client.post("/mvb/trips/new", data={
        "dir": ["wb|Коледино", "ozon|Хоругвино"], "mode": "fill", "capacity": "4",
    })
    trips = MvbTrip.query.order_by(MvbTrip.id).all()
    assert len(trips) == 2 and {t.status for t in trips} == {"searching"}
    assert [t.planned_boxes for t in trips] == [3, 2]
    assert [[s.label() for s in t.stops] for t in trips] == [["Wildberries · Коледино"], ["Ozon · Хоругвино"]]
    assert db.session.get(MvbOrder, first.id).lines[0].planned_trip_id == trips[0].id
    assert db.session.get(MvbOrder, second.id).lines[0].planned_trip_id == trips[1].id
    assert "Поиск авто: 2" in client.get("/mvb/trips").get_data(as_text=True)

    # 6. авто найдено — наемный водитель без регистрации, погрузка сканом
    # каждого короба; водителю уходит ссылка
    links = []
    for trip, order in zip(trips, (first, second)):
        client.post(f"/mvb/trips/{trip.id}/plan", data={
            "car_plate": "в777ор77", "driver_name": "Случайный Водитель", "driver_phone": "+79990000000",
        })
        db.session.refresh(trip)
        assert trip.status == "assigned" and trip.car_plate == "В777ОР77" and trip.driver_id is None
        links.append(f"/mvb/t/{trip.access_token}")
    # погрузка и отправка — кладовщик
    _login(client, keeper)
    for trip, order in zip(trips, (first, second)):
        for box in order.boxes:
            assert client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": box.barcode}).get_json()["ok"]
        client.post(f"/mvb/trips/{trip.id}/depart")

    # 7. водитель по ссылке (без входа): Коледино сдано, Хоругвино не сдано
    client.post("/mvb/logout")
    wb_trip, oz_trip = trips
    page = client.get(links[0]).get_data(as_text=True)
    assert wb_trip.number in page and "Сдан на СЦ" in page
    client.post(f"{links[0]}/stops/{wb_trip.stops[0].id}/deliver")
    client.post(f"{links[1]}/stops/{oz_trip.stops[0].id}/reject", data={"comment": "СЦ не принял: нет слота"})
    db.session.refresh(wb_trip)
    db.session.refresh(oz_trip)
    assert wb_trip.status == "delivered" and oz_trip.status == "delivered"
    assert {b.status for b in first.boxes} == {"delivered"}
    assert {b.status for b in second.boxes} == {"not_delivered"}
    oz_stop = oz_trip.stops[0]
    assert oz_stop.result == "rejected" and oz_stop.delivery_comment == "СЦ не принял: нет слота"

    # не сданный короб возвращается на склад (приемка кладовщиком) и снова
    # готов к отправке
    _login(client, keeper)
    assert client.post("/mvb/scan/receive", data={"barcode": second.boxes[0].barcode}).get_json()["ok"]
    box = db.session.get(MvbBox, second.boxes[0].id)
    assert box.status == "received" and box.trip_id is None
    _login(client, operator)
    assert "Хоругвино" in client.get("/mvb/dispatch").get_data(as_text=True)


def test_driver_on_sc_trip_does_not_see_free_pickups(db, client):
    """Водителя можно назначить и на забор, и на рейс на СЦ; в рейсе на СЦ
    функция «забрать по дороге» ему не нужна."""
    staff = _user("staff1", "mvb_admin")
    driver = _user("driver1", "mvb_driver")
    client_user = _user("client1", "mvb_client", _mvb_client())
    received = _received_order(client, client_user, staff, box_count="1")
    _login(client, client_user)
    free = _confirmed_order(client, box_count="2")
    mine = _confirmed_order(client, box_count="1")

    _login(client, staff)
    client.post(f"/mvb/orders/{mine.id}/driver", data={"driver_id": str(driver.id)})
    trip = _trip(client, driver=driver)
    assert trip.driver_id == driver.id
    client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": received.boxes[0].barcode})
    client.post(f"/mvb/trips/{trip.id}/depart")

    _login(client, driver)
    client.get("/mvb/driver")
    html = client.get("/mvb/driver").get_data(as_text=True)
    assert trip.number in html and "Сдан на СЦ" in html
    assert mine.number in html and free.number not in html and "Заберу" not in html
    client.post(f"/mvb/driver/orders/{free.id}/take")
    assert db.session.get(MvbOrder, free.id).driver_id is None

    # рейс закрыт — снова видит свободные заявки и может взять
    _deliver_all(client, trip)
    client.get("/mvb/driver")
    html = client.get("/mvb/driver").get_data(as_text=True)
    assert free.number in html and "Заберу" in html
    client.post(f"/mvb/driver/orders/{free.id}/take")
    assert db.session.get(MvbOrder, free.id).driver_id == driver.id


def test_reject_requires_reason(db, client):
    client_user = _user("client1", "mvb_client", _mvb_client())
    staff = _user("staff1", "mvb_admin")
    order = _received_order(client, client_user, staff, box_count="1")
    trip = _trip(client)
    client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": order.boxes[0].barcode})
    client.post(f"/mvb/trips/{trip.id}/depart")
    db.session.refresh(trip)
    client.post(f"/mvb/trips/{trip.id}/stops/{trip.stops[0].id}/reject", data={"comment": ""})
    db.session.refresh(trip)
    assert trip.stops[0].result is None and order.boxes[0].status == "shipped"


def test_fill_mode_packs_whole_orders(db, client):
    """Компоновка рейсов не разбивает заявки: 3 + 3 + 2 кор. в машины по 5 —
    [3 + 2] и [3]; заявка больше машины едет отдельно целиком."""
    client_user = _user("client1", "mvb_client", _mvb_client())
    staff = _user("staff1", "mvb_admin")
    a = _received_order(client, client_user, staff, box_count="3")
    b = _received_order(client, client_user, staff, box_count="3")
    c = _received_order(client, client_user, staff, marketplace="ozon", destination="Хоругвино", box_count="2")
    html = client.get("/mvb/dispatch?capacity=5").get_data(as_text=True)
    assert "нужно машин по 5 кор.: <b>2</b>" in html and "Предложение по рейсам" in html
    client.post("/mvb/trips/new", data={"dir": ["wb|Коледино", "ozon|Хоругвино"], "mode": "fill", "capacity": "5"})
    trips = MvbTrip.query.order_by(MvbTrip.id).all()
    assert [t.planned_boxes for t in trips] == [5, 3]
    assert [l.order.number for l in trips[0].planned_lines] == [a.number, c.number]
    assert [l.order.number for l in trips[1].planned_lines] == [b.number]

    # погрузка короба заявки из другого рейса — предупреждение
    client.post(f"/mvb/trips/{trips[1].id}/plan", data={"car_plate": "А1"})
    data = client.post(f"/mvb/trips/{trips[1].id}/scan", data={"barcode": a.boxes[0].barcode}).get_json()
    assert data["ok"] and trips[0].number in data["warning"]

    # отмена рейса освобождает заявки
    client.post(f"/mvb/trips/{trips[0].id}/cancel")
    assert db.session.get(MvbOrder, c.id).lines[0].planned_trip_id is None


def test_slot_date_required_and_separates_trips(db, client):
    """Дата слота на СЦ обязательна; заявки с разными слотами в одну машину
    не компонуются, ближайший слот — первым."""
    client_user = _user("client1", "mvb_client", _mvb_client())
    _login(client, client_user)
    html = _create_order(client, slot_date="").get_data(as_text=True)
    assert "Укажите дату слота" in html and MvbOrder.query.count() == 0

    staff = _user("staff1", "mvb_admin")
    late = _received_order(client, client_user, staff, box_count="2", slot_date="2026-10-09")
    early = _received_order(client, client_user, staff, box_count="1", slot_date="2026-10-07")
    assert early.slot_date.isoformat() == "2026-10-07"
    html = client.get("/mvb/dispatch?capacity=10").get_data(as_text=True)
    assert html.index("07.10.2026") < html.index("09.10.2026")
    assert "нужно машин по 10 кор.: <b>2</b>" in html  # 3 кор. влезли бы в одну, но слоты разные
    client.post("/mvb/trips/new", data={"dir": ["wb|Коледино"], "mode": "fill", "capacity": "10"})
    trips = MvbTrip.query.order_by(MvbTrip.id).all()
    assert [(t.slot_date.isoformat(), t.planned_boxes) for t in trips] == [("2026-10-07", 1), ("2026-10-09", 2)]

    # короб с другим слотом грузится с предупреждением
    client.post(f"/mvb/trips/{trips[0].id}/plan", data={"car_plate": "А1"})
    data = client.post(f"/mvb/trips/{trips[0].id}/scan", data={"barcode": late.boxes[0].barcode}).get_json()
    assert data["ok"] and "слот заявки 09.10" in data["warning"]
    assert "07.10.2026" in client.get(f"/mvb/t/{trips[0].access_token}").get_data(as_text=True)

    # слот оформленной заявки можно перенести, пока короба не уехали на СЦ
    client.post(f"/mvb/orders/{late.id}/slot", data={"slot_date": "2026-10-12"})
    assert db.session.get(MvbOrder, late.id).slot_date.isoformat() == "2026-10-12"


def test_order_bigger_than_truck_goes_whole(db, client):
    client_user = _user("client1", "mvb_client", _mvb_client())
    staff = _user("staff1", "mvb_admin")
    _received_order(client, client_user, staff, box_count="7")
    client.post("/mvb/trips/new", data={"dir": ["wb|Коледино"], "mode": "fill", "capacity": "3"})
    trips = MvbTrip.query.order_by(MvbTrip.id).all()
    assert [t.planned_boxes for t in trips] == [7]


def test_sc_driver_link_without_login(db, client):
    """Наемный водитель на СЦ не регистрируется: по ссылке он отмечает подачу
    и итог на точках; по чужому/неверному токену — 404, служебное закрыто."""
    client_user = _user("client1", "mvb_client", _mvb_client())
    staff = _user("staff1", "mvb_admin")
    order = _received_order(client, client_user, staff, box_count="2")
    trip = _new_trip(client, ("wb", "Коледино"))
    client.post(f"/mvb/trips/{trip.id}/plan", data={"car_plate": "е555кх77", "capacity_boxes": "10"})
    db.session.refresh(trip)
    assert trip.status == "assigned" and trip.capacity() == 10
    detail = client.get(f"/mvb/trips/{trip.id}").get_data(as_text=True)
    link = f"/mvb/t/{trip.access_token}"
    assert link in detail and "wa.me" not in detail

    client.post("/mvb/logout")
    assert client.get("/mvb/t/wrong-token").status_code == 404
    assert client.get(link).status_code == 200
    assert client.get(f"/mvb/trips/{trip.id}").status_code == 302  # служебное — только со входом
    client.post(f"{link}/arrive")
    assert db.session.get(MvbTrip, trip.id).status == "arrived"
    # до отправки «Сдано» не принимается
    client.post(f"{link}/stops/{trip.stops[0].id}/deliver")
    assert db.session.get(MvbTrip, trip.id).stops[0].result is None

    _login(client, staff)
    for box in order.boxes:
        client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": box.barcode})
    client.post(f"/mvb/trips/{trip.id}/depart")
    client.post("/mvb/logout")
    client.post(f"{link}/stops/{trip.stops[0].id}/reject", data={"comment": ""})
    assert db.session.get(MvbTrip, trip.id).stops[0].result is None  # без причины нельзя
    client.post(f"{link}/stops/{trip.stops[0].id}/deliver")
    trip = db.session.get(MvbTrip, trip.id)
    assert trip.status == "delivered" and trip.stops[0].delivered_by_id is None
    assert {b.status for b in order.boxes} == {"delivered"}
    assert "Рейс завершен" in client.get(link).get_data(as_text=True)


# ---------- прайс, стоимость, отчет ----------


def _set_prices(http):
    for kind, min_boxes, price in [("pickup", 1, "50"), ("sc", 1, "120"), ("sc", 10, "100"), ("sc", 50, "80,5")]:
        http.post("/mvb/prices", data={"kind": kind, "min_boxes": str(min_boxes), "price_per_box": price})


def test_price_tiers_and_order_cost(db, client):
    staff = _user("staff1", "mvb_admin")
    client_user = _user("client1", "mvb_client", _mvb_client())
    _login(client, staff)
    _set_prices(client)
    assert MvbPriceTier.price_for("sc", 9) == 120
    assert MvbPriceTier.price_for("sc", 10) == 100
    assert MvbPriceTier.price_for("sc", 60) == 80.5
    # та же ступень второй раз не дублируется, а обновляется
    client.post("/mvb/prices", data={"kind": "sc", "min_boxes": "10", "price_per_box": "95"})
    assert MvbPriceTier.query.filter_by(kind="sc", min_boxes=10).one().price_per_box == 95
    assert "Отправка на СЦ" in client.get("/mvb/prices").get_data(as_text=True)

    _login(client, client_user)
    pickup = _confirmed_order(client, box_count="12")
    self_order = _confirmed_order(client, box_count="3", delivery_method="self", pickup_address="")
    db.session.refresh(pickup)
    assert pickup.pickup_cost == 600 and pickup.sc_cost == 1140 and pickup.total_cost == 1740
    db.session.refresh(self_order)
    assert self_order.pickup_cost is None and self_order.sc_cost == 360

    # клиент видит стоимость, но поправить не может
    assert "1740.00" in client.get(f"/mvb/orders/{pickup.id}").get_data(as_text=True)
    client.post(f"/mvb/orders/{pickup.id}/costs", data={"pickup_cost": "0", "sc_cost": "0"})
    assert db.session.get(MvbOrder, pickup.id).total_cost == 1740

    _login(client, staff)
    client.post(f"/mvb/orders/{pickup.id}/costs", data={"pickup_cost": "500", "sc_cost": "1 000,50"})
    order = db.session.get(MvbOrder, pickup.id)
    assert order.pickup_cost == 500 and order.sc_cost == 1000.5
    client.post(f"/mvb/orders/{pickup.id}/costs", data={"action": "recalc"})
    assert db.session.get(MvbOrder, pickup.id).sc_cost == 1140
    client.post(f"/mvb/orders/{pickup.id}/costs", data={"pickup_cost": "-5"})
    assert db.session.get(MvbOrder, pickup.id).pickup_cost == 600


def test_prices_page_staff_only(db, client):
    _login(client, _user("client1", "mvb_client", _mvb_client()))
    client.post("/mvb/prices", data={"kind": "sc", "min_boxes": "1", "price_per_box": "1"})
    assert MvbPriceTier.query.count() == 0
    assert client.get("/mvb/reports").status_code == 302


def test_client_report_for_period(db, client):
    staff = _user("staff1", "mvb_admin")
    a = _user("ca", "mvb_client", _mvb_client("Альфа"))
    b = _user("cb", "mvb_client", _mvb_client("Бета"))
    _login(client, staff)
    _set_prices(client)
    order_a = _received_order(client, a, staff, box_count="3")
    _received_order(client, b, staff, box_count="2", delivery_method="self", pickup_address="")

    trip = _trip(client)
    for box in order_a.boxes[:2]:
        client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": box.barcode})
    client.post(f"/mvb/trips/{trip.id}/depart")
    db.session.refresh(trip)
    client.post(f"/mvb/trips/{trip.id}/stops/{trip.stops[0].id}/deliver")

    today = datetime.utcnow().date().isoformat()
    html = client.get(f"/mvb/reports?date_from={today}&date_to={today}").get_data(as_text=True)
    assert "Альфа" in html and "Бета" in html
    assert "Дата заявки" in html and order_a.number in html
    from wms.blueprints.mvb import _report_rows

    rows, totals, order_rows = _report_rows(datetime.utcnow().date(), datetime.utcnow().date())
    by_number = {r["order"].number: r for r in order_rows}
    assert by_number[order_a.number]["date"] is not None and by_number[order_a.number]["shipped"] == 2
    by_name = {r["client"].name: r for r in rows}
    assert by_name["Альфа"]["boxes"] == 3 and by_name["Альфа"]["shipped"] == 2 and by_name["Альфа"]["delivered"] == 2
    assert by_name["Альфа"]["total"] == 3 * 50 + 3 * 120
    assert by_name["Бета"]["shipped"] == 0 and by_name["Бета"]["pickup_cost"] == 0
    assert totals["boxes"] == 5

    xlsx = client.get(f"/mvb/reports.xlsx?date_from={today}&date_to={today}")
    assert xlsx.status_code == 200 and xlsx.data[:2] == b"PK"
    import io

    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(xlsx.data))["По заявкам"]
    assert sheet["A1"].value == "Дата заявки" and sheet["B2"].value == order_a.number
    # другой период — пусто
    assert _report_rows(datetime(2020, 1, 1).date(), datetime(2020, 1, 31).date())[0] == []


# ---------- регистрация клиента ----------


def _register(http, **overrides):
    data = {
        "name": "ИП Новый", "address": "Москва, ул. Ленина 1", "phone": "+7 900 123-45-67",
        "username": "newclient", "password": "secret1",
    }
    data.update(overrides)
    return http.post("/mvb/register", data=data)


def test_client_self_registration_needs_operator_approval(db, client):
    assert client.get("/mvb/register").status_code == 200
    assert "Зарегистрироваться" in client.get("/mvb/login").get_data(as_text=True)
    response = _register(client)
    assert response.status_code == 302 and response.headers["Location"].endswith("/mvb/login")
    new = MvbClient.query.filter_by(name="ИП Новый").one()
    assert new.approval == "pending" and new.users[0].username == "newclient"

    # до подтверждения войти нельзя
    html = client.post("/mvb/login", data={"username": "newclient", "password": "secret1"}).get_data(as_text=True)
    assert "на проверке" in html
    assert client.get("/mvb/orders").status_code == 302

    # оператор видит новую регистрацию и подтверждает
    staff = _user("staff1", "mvb_admin")
    _login(client, staff)
    client.get("/mvb/admin/clients")
    page = client.get("/mvb/admin/clients").get_data(as_text=True)
    assert "ИП Новый" in page and "+7 900 123-45-67" in page
    assert "Клиенты <span" in page
    client.post(f"/mvb/registrations/{new.id}/approve")
    assert db.session.get(MvbClient, new.id).approval == "approved"

    client.post("/mvb/logout")
    response = client.post("/mvb/login", data={"username": "newclient", "password": "secret1"})
    assert response.status_code == 302
    form = client.get("/mvb/orders/new").get_data(as_text=True)
    assert "Москва, ул. Ленина 1" in form  # адрес из регистрации подставляется в заявку


def test_registration_validation_and_reject(db, client):
    _user("taken", "mvb_client", _mvb_client())
    for overrides, message in [
        ({"address": ""}, "адрес забора"),
        ({"username": "Taken"}, "логин уже занят"),
        ({"password": "123"}, "не короче 6"),
        ({"phone": "12"}, "телефон"),
    ]:
        html = _register(client, **overrides).get_data(as_text=True)
        assert message in html
    assert MvbClient.query.filter_by(name="ИП Новый").count() == 0

    _register(client)
    new = MvbClient.query.filter_by(name="ИП Новый").one()
    staff = _user("staff1", "mvb_admin")
    _login(client, staff)
    client.post(f"/mvb/registrations/{new.id}/reject")
    client.post("/mvb/logout")
    html = client.post("/mvb/login", data={"username": "newclient", "password": "secret1"}).get_data(as_text=True)
    assert "отклонена" in html

    # клиент не может подтверждать регистрации
    other = _user("client2", "mvb_client", _mvb_client("Другой"))
    _login(client, other)
    client.post(f"/mvb/registrations/{new.id}/approve")
    assert db.session.get(MvbClient, new.id).approval == "rejected"


def test_wms_login_page_links_to_mvb(db, client):
    html = client.get("/login").get_data(as_text=True)
    assert "МВБ Логистика" in html and 'href="/mvb/login"' in html


def _multi_order_form(**overrides):
    form = {
        "direction_count": "2",
        "line_marketplace": ["wb", "ozon"],
        "line_destination": ["Коледино", "Хоругвино"],
        "line_slot": ["2026-10-06", "2026-10-08"],
        "line_boxes": ["3", "2"],
        "delivery_method": "pickup",
        "pickup_address": "Москва, ул. Ленина, 1",
    }
    form.update(overrides)
    return form


def test_one_order_many_directions_with_own_box_counters(db, client):
    """Одна заявка на несколько направлений: у каждого направления свой
    СЦ, слот, количество, свой список коробов и счетчик 1…N."""
    seller = _user("seller", "mvb_client", _mvb_client())
    _login(client, seller)
    form = client.get("/mvb/orders/new").get_data(as_text=True)
    assert "Сколько направлений" in form
    client.post("/mvb/orders/new", data=_multi_order_form())
    order = MvbOrder.query.one()
    assert [l.short_label() for l in order.lines] == ["WB Коледино", "OZON Хоругвино"]
    assert order.box_count == 5 and order.slot_date.isoformat() == "2026-10-06"
    client.post(f"/mvb/orders/{order.id}/confirm")
    wb, oz = order.lines
    assert [b.barcode for b in wb.boxes] == [f"{order.number}-1-001", f"{order.number}-1-002", f"{order.number}-1-003"]
    assert [b.barcode for b in oz.boxes] == [f"{order.number}-2-001", f"{order.number}-2-002"]
    assert [(b.line_position, b.line_total) for b in oz.boxes] == [(1, 2), (2, 2)]
    client.get(f"/mvb/orders/{order.id}")
    page = client.get(f"/mvb/orders/{order.id}").get_data(as_text=True)
    assert "OZON Хоругвино" in page and "сдано на СЦ 0 / 2" in page and "Этикетки направления" in page
    # этикетки одного направления
    response = client.get(f"/mvb/orders/{order.id}/labels.pdf?line={oz.id}")
    assert response.status_code == 200 and response.mimetype == "application/pdf"
    # слот меняется по направлению
    client.post(f"/mvb/orders/{order.id}/slot", data={"line_id": str(oz.id), "slot_date": "2026-10-09"})
    assert db.session.get(MvbOrder, order.id).lines[1].slot_date.isoformat() == "2026-10-09"
    # водитель видит одну заявку с общим количеством коробов
    _login(client, _user("driver", "mvb_driver"))
    client.get("/mvb/driver")
    feed = client.get("/mvb/driver").get_data(as_text=True)
    assert order.number in feed and "<b>Коробов:</b> 5" in feed and "<b>Телефон:</b> +79990000000" in feed
    # скан короба: номер внутри своего направления
    data = client.post("/mvb/scan/pickup", data={"barcode": oz.boxes[1].barcode}).get_json()
    assert data["seq"] == 2 and data["total"] == 2 and data["direction"] == "OZON Хоругвино"


def test_multi_direction_form_validates_each_row(db, client):
    _login(client, _user("seller", "mvb_client", _mvb_client()))
    html = client.post("/mvb/orders/new", data=_multi_order_form(
        line_destination=["Коледино", ""],
    )).get_data(as_text=True)
    assert "Направление 2: выберите город" in html and MvbOrder.query.count() == 0
    # лишние строки сверх выбранного количества не учитываются
    client.post("/mvb/orders/new", data=_multi_order_form(direction_count="1"))
    assert len(MvbOrder.query.one().lines) == 1


def test_wms_movements_from_one_warehouse_become_one_order_for_drivers(db, client):
    _login(client, _user("operator", "mvb_staff"))
    doc1, boxes1 = _wms_movement("PER-000701", boxes=2)
    doc2, boxes2 = _wms_movement("PER-000702", boxes=1)
    from datetime import date as _date

    doc1.delivery_slot_date = _date(2026, 10, 9)
    doc2.from_warehouse_id = doc1.from_warehouse_id
    doc2.to_warehouse.marketplace = "ozon"
    doc2.to_warehouse.marketplace_city = "Хоругвино"
    db.session.commit()
    client.post("/mvb/wms/import", data={"doc_id": [str(doc1.id), str(doc2.id)]})
    order = MvbOrder.query.one()
    assert order.delivery_method == "pickup" and order.pickup_address == "Москва, Складская 5"
    assert [(l.short_label(), l.box_count) for l in order.lines] == [("WB Коледино", 2), ("OZON Хоругвино", 1)]
    assert [b.barcode for b in order.lines[1].boxes] == [boxes2[0].barcode_value]
    assert order.lines[0].slot_date == doc1.delivery_slot_date
    # заявка из WMS падает в заявки водителям
    _login(client, _user("driver", "mvb_driver"))
    assert order.number in client.get("/mvb/driver").get_data(as_text=True)


def test_wms_movements_from_different_warehouses_are_not_merged(db, client):
    _login(client, _user("operator", "mvb_staff"))
    doc1, _ = _wms_movement("PER-000711")
    doc2, _ = _wms_movement("PER-000712")
    client.post("/mvb/wms/import", data={"doc_id": [str(doc1.id), str(doc2.id)]})
    assert MvbOrder.query.count() == 0


def test_menu_by_role(db, client):
    seller = _user("seller", "mvb_client", _mvb_client())
    roles = {
        "seller": (seller, ["Мои заявки"], ["Заявки в работе", "Приемка", "Скан забора", "Настройки"]),
        "driver": (_user("driver", "mvb_driver"), ["Заявки на забор", "Скан забора"], ["Заявки в работе", "Приемка", "Паллеты"]),
        "operator": (_user("operator", "mvb_staff"), ["Заявки в работе", "Клиенты", "Настройки", "Прайс", "Транспорт", "Пользователи", "Рейсы"], ["Приемка", "Скан забора", "Заявки на забор", "Паллеты"]),
        "keeper": (_user("keeper", "mvb_storekeeper"), ["Приемка", "Паллеты", "Погрузка"], ["Заявки в работе", "Скан забора", "Настройки", "Клиенты", "Рейсы"]),
    }
    for name, (user, visible, hidden) in roles.items():
        _login(client, user)
        page = client.get({"operator": "/mvb/trips", "keeper": "/mvb/pallets"}.get(name, "/mvb/")).get_data(as_text=True)
        if name in ("seller", "driver"):
            page = client.get(client.get("/mvb/").headers["Location"]).get_data(as_text=True)
        nav = page.split("</nav>")[0]
        for item in visible:
            assert item in nav, (name, item)
        for item in hidden:
            assert item not in nav, (name, item)
    # кладовщик попадает сразу на приемку, оператору приемка недоступна
    _login(client, roles["keeper"][0])
    assert client.get("/mvb/").headers["Location"].endswith("/mvb/scan/receive")
    _login(client, roles["operator"][0])
    assert client.get("/mvb/scan/receive").status_code == 302
    assert client.get("/mvb/scan/pickup").status_code == 302


def test_storekeeper_cannot_do_operator_work(db, client):
    _login(client, _user("keeper", "mvb_storekeeper"))
    assert client.get("/mvb/prices").status_code == 302
    assert client.get("/mvb/admin/clients").status_code == 302
    client.post("/mvb/trips/new", data={"dir": ["wb|Коледино"]})
    assert MvbTrip.query.count() == 0


def test_clients_and_registrations_in_one_section_for_operator(db, client):
    pending = MvbClient(name="ИП Ждет", approval="pending")
    db.session.add(pending)
    db.session.commit()
    _login(client, _user("operator", "mvb_staff"))
    page = client.get("/mvb/admin/clients").get_data(as_text=True)
    assert "Новые регистрации" in page and "ИП Ждет" in page
    client.post(f"/mvb/registrations/{pending.id}/approve")
    assert db.session.get(MvbClient, pending.id).approval == "approved"
    assert client.get("/mvb/registrations").headers["Location"].endswith("/mvb/admin/clients")


def test_operator_manages_users_except_mvb_admins(db, client):
    admin = _user("boss", "mvb_admin")
    _login(client, _user("operator", "mvb_staff"))
    page = client.get("/mvb/admin/users").get_data(as_text=True)
    assert "Кладовщик МВБ" in page and 'value="mvb_admin"' not in page
    client.post("/mvb/admin/users", data={"username": "k1", "password": "secret1", "role": "mvb_storekeeper"})
    assert User.query.filter_by(username="k1").one().role == "mvb_storekeeper"
    assert client.post(f"/mvb/admin/users/{admin.id}/password", data={"password": "hacked1"}).status_code == 403
    assert client.post(f"/mvb/admin/users/{admin.id}/toggle").status_code == 403
    assert "Создать тестовых пользователей" not in page


def test_test_accounts_one_per_role(db, client):
    _login(client, _user("boss", "mvb_admin"))
    client.post("/mvb/admin/users/test-accounts")
    page = client.get("/mvb/admin/users").get_data(as_text=True)
    users = {u.username: u for u in User.query.filter(User.username.like("test_%")).all()}
    assert {u.role for u in users.values()} == {"mvb_client", "mvb_driver", "mvb_staff", "mvb_storekeeper", "mvb_admin"}
    assert users["test_client"].mvb_client.is_approved()
    assert "сохраните пароли" in page
    # пароль показан один раз и подходит для входа
    import re
    password = re.search(r"test_driver</td><td class=\"font-monospace\">([^<]+)<", page).group(1)
    assert "сохраните пароли" not in client.get("/mvb/admin/users").get_data(as_text=True)
    client.post("/mvb/logout")
    g.pop("_login_user", None)
    response = client.post("/mvb/login", data={"username": "test_driver", "password": password})
    assert response.status_code == 302 and response.headers["Location"].endswith("/mvb/")
    # повторное нажатие — новые пароли, без дублей
    _login(client, User.query.filter_by(username="boss").one())
    client.post("/mvb/admin/users/test-accounts")
    assert User.query.filter(User.username.like("test_%")).count() == 5
    assert MvbClient.query.filter_by(name="Тестовый клиент").count() == 1


def test_old_orders_get_one_direction_on_startup(db, client):
    from wms import _ensure_mvb_lines

    rom = _mvb_client()
    order = MvbOrder(number="MVB-OLD", client_id=rom.id, marketplace="wb", destination="Коледино", box_count=2,
                     delivery_method="self", status="confirmed")
    order.boxes = [MvbBox(seq=1, barcode="MVB-OLD-001"), MvbBox(seq=2, barcode="MVB-OLD-002")]
    db.session.add(order)
    db.session.commit()
    assert not order.lines
    _ensure_mvb_lines()
    _ensure_mvb_lines()
    order = db.session.get(MvbOrder, order.id)
    assert len(order.lines) == 1 and order.lines[0].short_label() == "WB Коледино"
    assert all(b.line_id == order.lines[0].id for b in order.boxes)


def test_fulfillment_as_destination_for_orders_pallets_and_trips(db, client):
    """Кроме складов маркетплейсов, короба и паллеты едут на фулфилменты."""
    client_user = _user("client1", "mvb_client", _mvb_client())
    staff = _user("staff1", "mvb_admin")
    order = _received_order(client, client_user, staff, marketplace="ff", destination="Север", box_count="2")
    assert order.lines[0].short_label() == "ФФ Север"
    client.post("/mvb/pallets", data={"marketplace": "ff", "destination": "Север"})
    pallet = MvbPallet.query.one()
    for box in order.boxes:
        assert client.post(f"/mvb/pallets/{pallet.id}/scan", data={"barcode": box.barcode}).get_json()["ok"]
    assert "Фулфилмент" in client.get("/mvb/dispatch").get_data(as_text=True)
    trip = _trip(client, marketplace="ff", destination="Север")
    data = client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": pallet.number}).get_json()
    assert data["ok"] and data["count"] == 2
    assert [s.label() for s in trip.stops] == ["Фулфилмент · Север"]


def test_cities_with_prices_and_destinations(db, client):
    """Справочник городов в прайсе: у города свой прайс отправки на СЦ —
    одинаковый для WB, Ozon и ФФ этого города. Пункты назначения: WB / Ozon /
    ФФ → город из справочника → (для ФФ) название; в заявке выбирают из
    списка."""
    from wms.models import MvbCity, MvbDestination

    _login(client, _user("operator", "mvb_staff"))
    for name in ("Коледино", "Казань", "Москва"):
        client.post("/mvb/prices", data={"action": "city_add", "name": name})
    client.post("/mvb/prices", data={"action": "city_add", "name": "казань"})  # дубль
    assert sorted(c.name for c in MvbCity.active()) == ["Казань", "Коледино", "Москва"]
    for mp, city, ff in [("wb", "Коледино", ""), ("wb", "Казань", ""), ("ozon", "Казань", ""), ("ff", "Москва", "ФФ Север")]:
        client.post("/mvb/prices", data={"action": "dest_add", "marketplace": mp, "city": city, "ff_name": ff})
    client.post("/mvb/prices", data={"action": "dest_add", "marketplace": "wb", "city": "коледино"})  # дубль
    client.post("/mvb/prices", data={"action": "dest_add", "marketplace": "ff", "city": "Москва"})  # без названия ФФ
    client.post("/mvb/prices", data={"action": "dest_add", "marketplace": "wb", "city": "Тула"})  # нет в справочнике
    assert sorted(d.label() for d in MvbDestination.active()) == [
        "OZON Казань", "WB Казань", "WB Коледино", "ФФ ФФ Север (Москва)",
    ]
    kazan = MvbCity.find("казань")
    client.post("/mvb/prices", data={"kind": "pickup", "min_boxes": "1", "price_per_box": "50"})
    client.post("/mvb/prices", data={"kind": "sc", "min_boxes": "1", "price_per_box": "100"})
    client.post("/mvb/prices", data={"kind": "sc", "min_boxes": "1", "price_per_box": "300", "city_id": str(kazan.id)})
    page = client.get("/mvb/prices").get_data(as_text=True)
    assert "Города и прайс отправки" in page and "Пункты назначения" in page
    assert MvbPriceTier.price_for("sc", 5, kazan.id) == 300
    assert MvbPriceTier.price_for("sc", 5, MvbCity.find("Москва").id) == 100

    _login(client, _user("seller", "mvb_client", _mvb_client()))
    form = client.get("/mvb/orders/new").get_data(as_text=True)
    import json as _json
    dests = _json.loads(form.split("const DESTS = ", 1)[1].split(";", 1)[0])
    assert {"mp": "ff", "city": "Москва", "value": "ФФ Север"} in dests
    # Ozon Коледино в списке нет — нельзя
    html = client.post("/mvb/orders/new", data=_multi_order_form(
        line_marketplace=["wb", "ozon"], line_destination=["Коледино", "Коледино"])).get_data(as_text=True)
    assert "нет в списке пунктов назначения" in html and MvbOrder.query.count() == 0
    client.post("/mvb/orders/new", data=_multi_order_form(
        line_marketplace=["ozon", "ff"], line_destination=["казань", "ФФ Север"]))
    order = MvbOrder.query.one()
    assert [l.short_label() for l in order.lines] == ["OZON Казань", "ФФ ФФ Север"]
    client.post(f"/mvb/orders/{order.id}/confirm")
    order = db.session.get(MvbOrder, order.id)
    # Ozon Казань — по прайсу Казани, ФФ в Москве — по общему
    assert order.sc_cost == 3 * 300 + 2 * 100 and order.pickup_cost == 5 * 50

    _login(client, User.query.filter_by(username="operator").one())
    # переименование города переносится в его пункты
    client.post("/mvb/prices", data={"action": "city_save", "city_id": str(kazan.id), "name": "Казань-2"})
    assert {d.city for d in MvbDestination.query.filter_by(marketplace="ozon")} == {"Казань-2"}
    # скрытый город и его пункты из списка пропадают
    client.post("/mvb/prices", data={"action": "city_toggle", "city_id": str(kazan.id)})
    assert not any(d.city == "Казань-2" for d in MvbDestination.active())


def test_old_cities_and_destination_prices_migrate(db, client):
    from wms import _ensure_mvb_destinations
    from wms.models import MvbCity, MvbDestination

    city = MvbCity(name="Казань", address_wb="Казань, ул. WB 1", address_ozon="Казань, ул. Ozon 2")
    db.session.add(city)
    db.session.flush()
    db.session.add(MvbPriceTier(kind="sc", min_boxes=1, price_per_box=250, city_id=city.id))
    # прайс, заведенный на пункт назначения (предыдущая версия), — на его город
    ff = MvbDestination(marketplace="ff", city="Москва", ff_name="ФФ Север")
    db.session.add(ff)
    db.session.flush()
    db.session.add(MvbPriceTier(kind="sc", min_boxes=1, price_per_box=180, destination_id=ff.id))
    db.session.commit()
    _ensure_mvb_destinations()
    _ensure_mvb_destinations()
    wb, oz = MvbDestination.find("wb", "Казань"), MvbDestination.find("ozon", "Казань")
    assert wb.address == "Казань, ул. WB 1" and oz.address == "Казань, ул. Ozon 2"
    assert MvbDestination.query.count() == 3
    kazan, moscow = MvbCity.find("Казань"), MvbCity.find("Москва")
    assert kazan.address_wb is None
    assert MvbPriceTier.price_for("sc", 3, kazan.id) == 250 and MvbPriceTier.price_for("sc", 3, moscow.id) == 180


def test_pickup_done_button_requires_all_boxes(db, client):
    """После скана коробов на заборе — «Готово»; пока не все короба
    отсканированы, завершить нельзя."""
    _login(client, _user("seller", "mvb_client", _mvb_client()))
    order = _confirmed_order(client, box_count="3")
    driver = _user("driver", "mvb_driver")
    _login(client, driver)
    client.post(f"/mvb/driver/orders/{order.id}/take")
    assert f"/mvb/scan/pickup?order={order.id}" in client.get("/mvb/driver").get_data(as_text=True)
    data = client.post("/mvb/scan/pickup", data={"barcode": order.boxes[0].barcode}).get_json()
    assert data["order_id"] == order.id and data["order_done"] == 1 and data["order_total"] == 3
    response = client.post(f"/mvb/driver/orders/{order.id}/done")
    assert response.headers["Location"].endswith(f"/mvb/scan/pickup?order={order.id}")
    page = client.get(response.headers["Location"]).get_data(as_text=True)
    assert "Не все короба отсканированы: 1 из 3" in page and "Готово" in page
    assert db.session.get(MvbOrder, order.id).pickup_done_at is None
    for box in order.boxes[1:]:
        client.post("/mvb/scan/pickup", data={"barcode": box.barcode})
    response = client.post(f"/mvb/driver/orders/{order.id}/done")
    assert response.headers["Location"].endswith("/mvb/driver")
    assert db.session.get(MvbOrder, order.id).pickup_done_at is not None


def test_sc_addresses_route_and_driver_sees_only_address_and_count(db, client):
    """Адреса СЦ в списке городов — маршрут водителю; на точке — адрес,
    сколько сдать и «Сдан на СЦ»."""
    from wms.models import MvbDestination

    staff = _user("staff1", "mvb_admin")
    _login(client, staff)
    client.post("/mvb/prices", data={"action": "city_add", "name": "Коледино"})
    client.post("/mvb/prices", data={
        "action": "dest_add", "marketplace": "wb", "city": "Коледино", "address": "Подольск, Коледино, ул. Троицкая 20",
    })
    assert MvbDestination.find("wb", "Коледино").address == "Подольск, Коледино, ул. Троицкая 20"
    a = _received_order(client, _user("c1", "mvb_client", _mvb_client("ИП А")), staff, box_count="2")
    b = _received_order(client, _user("c2", "mvb_client", _mvb_client("ИП Б")), staff, box_count="1")
    driver = _user("driver", "mvb_driver")
    trip = _trip(client, driver=driver)
    for box in a.boxes + b.boxes:
        client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": box.barcode})
    client.post(f"/mvb/trips/{trip.id}/depart")
    db.session.refresh(trip)
    stop = trip.stops[0]
    assert stop.address == "Подольск, Коледино, ул. Троицкая 20"
    assert "rtext=~" in trip.route_url() and "%D0%9F%D0%BE%D0%B4%D0%BE%D0%BB%D1%8C%D1%81%D0%BA" in trip.route_url()

    _login(client, driver)
    client.get("/mvb/driver")
    page = client.get("/mvb/driver").get_data(as_text=True)
    # водителю — адрес, сколько паллет/коробов сдать и «Сдан на СЦ», без клиентов
    assert "Маршрут по точкам" in page and "Подольск, Коледино" in page and "<b>3</b> кор." in page
    assert "ИП А" not in page and page.count("Сдан на СЦ</button>") == 1
    client.post(f"/mvb/trips/{trip.id}/stops/{stop.id}/deliver")
    db.session.refresh(trip)
    assert trip.status == "delivered" and {x.status for x in a.boxes + b.boxes} == {"delivered"}


def test_pallet_page_lists_direction_boxes_and_highlights_scanned(db, client):
    staff = _user("staff1", "mvb_admin")
    order = _received_order(client, _user("c1", "mvb_client", _mvb_client()), staff, box_count="3")
    other = _received_order(client, User.query.filter_by(username="c1").one(), staff,
                            marketplace="ozon", destination="Хоругвино", box_count="1")
    _login(client, staff)
    client.post("/mvb/pallets", data={"marketplace": "wb", "destination": "Коледино"})
    pallet = MvbPallet.query.one()
    data = client.post(f"/mvb/pallets/{pallet.id}/scan", data={"barcode": order.boxes[0].barcode}).get_json()
    assert data["barcode"] == order.boxes[0].barcode
    page = client.get(f"/mvb/pallets/{pallet.id}").get_data(as_text=True)
    assert "на паллете <span id=\"palletCount\">1</span> из 3" in page
    assert f'data-barcode="{order.boxes[0].barcode}" class="table-success"' in page
    assert f'data-barcode="{order.boxes[1].barcode}" class=""' in page
    assert other.boxes[0].barcode not in page  # другое направление не показываем


def test_operator_assigns_car_and_storekeeper_loads_by_plate(db, client):
    """Оператор вносит авто (госномер, модель, ФИО, телефон, план) и жмет
    «Назначить на рейс»; кладовщик выбирает авто по госномеру, «Начать
    погрузку», сканирует паллету, «Завершить погрузку»."""
    admin = _user("staff1", "mvb_admin")
    order = _received_order(client, _user("c1", "mvb_client", _mvb_client()), admin, box_count="2")
    client.post("/mvb/pallets", data={"marketplace": "wb", "destination": "Коледино"})
    pallet = MvbPallet.query.one()
    for box in order.boxes:
        client.post(f"/mvb/pallets/{pallet.id}/scan", data={"barcode": box.barcode})

    _login(client, _user("operator", "mvb_staff"))
    trip = _new_trip(client, ("wb", "Коледино"))
    page = client.get(f"/mvb/trips/{trip.id}").get_data(as_text=True)
    assert "Назначить на рейс" in page and "WhatsApp" not in page and "Telegram" not in page
    client.post(f"/mvb/trips/{trip.id}/plan", data={
        "car_plate": "в123ор77", "car_model": "ГАЗель Next", "driver_name": "Иванов И.", "driver_phone": "+79990001122",
        "planned_arrival_at": "2026-10-06T08:00", "planned_load_start_at": "2026-10-06T08:30",
        "planned_load_end_at": "2026-10-06T09:00",
    })
    db.session.refresh(trip)
    assert trip.status == "assigned" and trip.transport_label() == "В123ОР77 · ГАЗель Next"

    _login(client, _user("keeper", "mvb_storekeeper"))
    page = client.get("/mvb/loading").get_data(as_text=True)
    assert "В123ОР77" in page and "Начать погрузку" in page
    response = client.post("/mvb/loading", data={"trip_id": str(trip.id)})
    assert response.headers["Location"].endswith(f"/mvb/loading/{trip.id}")
    db.session.refresh(trip)
    assert trip.status == "loading" and trip.load_started_at is not None
    data = client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": pallet.number}).get_json()
    assert data["ok"] and data["loaded"] == 2
    assert "Завершить погрузку" in client.get(f"/mvb/loading/{trip.id}").get_data(as_text=True)
    response = client.post(f"/mvb/trips/{trip.id}/depart", data={"back": "loading"})
    assert response.headers["Location"].endswith("/mvb/loading")
    db.session.refresh(trip)
    assert trip.status == "departed" and {b.status for b in order.boxes} == {"shipped"}
    # водитель по ссылке видит, сколько паллет сдать на точке
    client.post("/mvb/logout")
    page = client.get(f"/mvb/t/{trip.access_token}").get_data(as_text=True)
    assert "Сдать: <b>1</b> пал." in page and "Сдан на СЦ" in page


def test_admin_deletes_orders_trips_and_pallets(db, client):
    client_user = _user("client1", "mvb_client", _mvb_client())
    operator = _user("oper1", "mvb_staff")
    admin = _user("boss1", "mvb_admin")
    order = _received_order(client, client_user, admin)
    second = _received_order(client, client_user, admin, box_count="2")
    client.post("/mvb/pallets", data={"marketplace": "wb", "destination": "Коледино"})
    pallet = MvbPallet.query.one()
    for box in second.boxes:
        client.post(f"/mvb/pallets/{pallet.id}/scan", data={"barcode": box.barcode})
    trip = _trip(client)
    client.post(f"/mvb/trips/{trip.id}/scan", data={"barcode": order.boxes[0].barcode})
    client.post(f"/mvb/trips/{trip.id}/depart")
    assert order.boxes[0].status == "shipped"
    assert "🗑 Удалить" in client.get(f"/mvb/trips/{trip.id}").get_data(as_text=True)

    # оператор и клиент удалять не могут
    for user in (operator, client_user):
        _login(client, user)
        client.post(f"/mvb/orders/{order.id}/delete")
        client.post(f"/mvb/trips/{trip.id}/delete")
        client.post(f"/mvb/pallets/{pallet.id}/delete")
    assert db.session.get(MvbOrder, order.id) and db.session.get(MvbTrip, trip.id) and db.session.get(MvbPallet, pallet.id)
    _login(client, operator)
    assert "🗑 Удалить" not in client.get(f"/mvb/orders/{order.id}").get_data(as_text=True)

    _login(client, admin)
    client.post(f"/mvb/trips/{trip.id}/delete")
    db.session.expire_all()
    assert db.session.get(MvbTrip, trip.id) is None
    box = db.session.get(MvbBox, order.boxes[0].id)
    assert box.status == "received" and box.trip_id is None and box.trip_stop_id is None

    client.post(f"/mvb/pallets/{pallet.id}/delete")
    db.session.expire_all()
    assert db.session.get(MvbPallet, pallet.id) is None
    assert {b.pallet_id for b in second.boxes} == {None}

    box_ids = [b.id for b in order.boxes]
    client.post(f"/mvb/orders/{order.id}/delete")
    db.session.expire_all()
    assert db.session.get(MvbOrder, order.id) is None
    assert MvbBox.query.filter(MvbBox.id.in_(box_ids)).count() == 0
    assert "удалена" in client.get("/mvb/orders").get_data(as_text=True)


def test_default_prices_seeded_once_and_pickup_zones(db, client):
    """Прайс MWB по умолчанию: города с прайсом, пункты WB/Ozon, зоны
    забора; заносится один раз и не трогает уже заведенное."""
    from wms import _seed_mvb_default_prices
    from wms.models import MvbCity, MvbDestination, MvbPickupZone

    db.session.add(MvbCity(name="казань"))
    db.session.commit()
    _seed_mvb_default_prices()
    kazan = MvbCity.find("Казань")
    assert MvbCity.query.filter(db.func.lower(MvbCity.name) == "казань").count() == 1
    assert MvbPriceTier.price_for("sc", 1, kazan.id) == 600 and MvbPriceTier.price_for("sc", 11, kazan.id) == 420
    novosib = MvbCity.find("Новосибирск")
    assert [MvbPriceTier.price_for("sc", n, novosib.id) for n in (1, 11, 16)] == [1050, 870, 790]
    assert MvbDestination.find("ozon", "Хоругвино") and MvbDestination.find("wb", "Подольск 4")
    assert MvbDestination.find("wb", "Хоругвино") is None
    assert {(z.name, z.price) for z in MvbPickupZone.active()} == {
        ("Черкесск", 1000), ("Регионы", 1500), ("Хабез", 1800), ("Отрадная", 2800),
    }
    # второй раз не срабатывает: удаленное оператором не возвращается
    MvbPriceTier.query.filter_by(city_id=kazan.id).delete()
    db.session.commit()
    _seed_mvb_default_prices()
    assert MvbPriceTier.query.filter_by(city_id=kazan.id).count() == 0

    # заявка: зона забора обязательна, цена забора — по зоне
    seller = _user("seller", "mvb_client", _mvb_client())
    _login(client, seller)
    html = client.post("/mvb/orders/new", data=_multi_order_form(
        line_marketplace=["ozon", "wb"], line_destination=["Хоругвино", "Подольск 4"])).get_data(as_text=True)
    assert "Выберите зону забора" in html and MvbOrder.query.count() == 0
    zone = MvbPickupZone.query.filter_by(name="Регионы").one()
    client.post("/mvb/orders/new", data=_multi_order_form(
        line_marketplace=["ozon", "wb"], line_destination=["Хоругвино", "Подольск 4"], pickup_zone_id=str(zone.id)))
    order = MvbOrder.query.one()
    client.post(f"/mvb/orders/{order.id}/confirm")
    order = db.session.get(MvbOrder, order.id)
    assert order.pickup_cost == 1500 and order.sc_cost == 3 * 550 + 2 * 550
    assert "Зона забора" in client.get(f"/mvb/orders/{order.id}").get_data(as_text=True)

    # оператор правит зону
    _login(client, _user("operator", "mvb_staff"))
    client.post("/mvb/prices", data={"action": "zone_save", "zone_id": str(zone.id), "name": "Регионы", "price": "1600"})
    assert db.session.get(MvbPickupZone, zone.id).price == 1600
    assert "Зона *" in client.get("/mvb/prices").get_data(as_text=True)


def test_pallet_cost_shared_by_box_share(db, client):
    """Палетирование: паллета одного клиента — вся цена ему; сборная —
    делится по доле коробов каждого клиента на паллете."""
    staff = _user("staff1", "mvb_admin")
    a = _received_order(client, _user("c1", "mvb_client", _mvb_client("ИП А")), staff, box_count="3")
    b = _received_order(client, _user("c2", "mvb_client", _mvb_client("ИП Б")), staff, box_count="1")
    client.post("/mvb/prices", data={"action": "pallet_price", "pallet_price": "500"})
    client.post("/mvb/pallets", data={"marketplace": "wb", "destination": "Коледино"})
    pallet = MvbPallet.query.one()
    for box in a.boxes:
        client.post(f"/mvb/pallets/{pallet.id}/scan", data={"barcode": box.barcode})
    db.session.expire_all()
    assert db.session.get(MvbOrder, a.id).pallet_cost == 500
    client.post(f"/mvb/pallets/{pallet.id}/scan", data={"barcode": b.boxes[0].barcode})
    db.session.expire_all()
    a, b = db.session.get(MvbOrder, a.id), db.session.get(MvbOrder, b.id)
    assert a.pallet_cost == 375 and b.pallet_cost == 125
    assert "палетирование 375.00" in client.get(f"/mvb/orders/{a.id}").get_data(as_text=True)

    # короб снят с паллеты — доли пересчитываются
    client.post(f"/mvb/pallets/{pallet.id}/remove/{b.boxes[0].id}")
    db.session.expire_all()
    assert db.session.get(MvbOrder, a.id).pallet_cost == 500
    assert db.session.get(MvbOrder, b.id).pallet_cost is None
    # отчет
    page = client.get("/mvb/reports?date_from=2000-01-01&date_to=2100-01-01").get_data(as_text=True)
    assert "Палетирование" in page


def test_client_sees_driver_data_for_pass(db, client):
    """В заявке клиента — данные водителя для пропуска; пока рейс не
    назначен — заглушка."""
    seller = _user("seller", "mvb_client", _mvb_client())
    staff = _user("staff1", "mvb_admin")
    order = _received_order(client, seller, staff, box_count="2")
    _login(client, seller)
    page = client.get(f"/mvb/orders/{order.id}").get_data(as_text=True)
    assert "Здесь появятся данные водителя для пропуска" in page

    _login(client, staff)
    trip = _new_trip(client, ("wb", "Коледино"))
    client.post(f"/mvb/trips/{trip.id}/plan", data={
        "car_plate": "А777АА09", "car_model": "Газель", "driver_name": "Иванов Иван", "driver_phone": "+79991112233",
        "planned_arrival_at": "2026-10-06T09:00",
    })
    _login(client, seller)
    page = client.get(f"/mvb/orders/{order.id}").get_data(as_text=True)
    assert "Здесь появятся данные водителя" not in page
    assert "Иванов Иван" in page and "А777АА09" in page and "Газель" in page and "+79991112233" in page


def test_mvb_guide_for_every_role(db, client):
    import os
    for name, role in [("c", "mvb_client"), ("d", "mvb_driver"), ("o", "mvb_staff"), ("k", "mvb_storekeeper")]:
        _login(client, _user(name, role, _mvb_client(name) if role == "mvb_client" else None))
        page = client.get("/mvb/guide").get_data(as_text=True)
        assert "Как работает доставка коробов на СЦ" in page and "🎓 Обучение" in page
        assert 'class="mine"' in page
    static = os.path.join(os.path.dirname(__file__), "..", "wms", "static", "mvb_onboarding")
    import re
    for img in set(re.findall(r"mvb_onboarding/([\w.]+\.png)", page)):
        assert os.path.exists(os.path.join(static, img)), img


def test_direction_cannot_be_planned_into_two_trips(db, client):
    """Направление, уже запланированное в рейс, пропадает из «К отправке» и
    второй рейс на него не создается (баг: одна заявка — два водителя)."""
    staff = _user("staff1", "mvb_admin")
    order = _received_order(client, _user("c1", "mvb_client", _mvb_client()), staff,
                            marketplace="ozon", destination="Волгоград", box_count="1")
    key = f"ozon|Волгоград|{order.lines[0].slot_date.isoformat()}"
    client.post("/mvb/trips/new", data={"dir": [key]})
    client.post("/mvb/trips/new", data={"dir": [key]})
    assert MvbTrip.query.count() == 1
    client.get("/mvb/dispatch")
    page = client.get("/mvb/dispatch").get_data(as_text=True)
    assert "Волгоград" not in page and order.number not in page
    # рейс отменен — направление снова можно отправить
    trip = MvbTrip.query.one()
    client.post(f"/mvb/trips/{trip.id}/cancel")
    client.post("/mvb/trips/new", data={"dir": [key]})
    assert MvbTrip.query.count() == 2


def test_sc_driver_link_expires_day_after_route_finished(db, client):
    from datetime import timedelta
    staff = _user("staff1", "mvb_admin")
    _login(client, staff)
    trip = _new_trip(client, ("wb", "Коледино"))
    url = f"/mvb/t/{trip.access_token}"
    assert client.get(url).status_code == 200
    trip.status, trip.delivered_at = "delivered", datetime.utcnow() - timedelta(hours=2)
    db.session.commit()
    assert client.get(url).status_code == 200
    trip.delivered_at = datetime.utcnow() - timedelta(days=2)
    db.session.commit()
    resp = client.get(url)
    assert resp.status_code == 410 and "Ссылка больше не действует" in resp.get_data(as_text=True)


def test_screens_by_role_driver_and_keeper_do_not_see_orders(db, client):
    """Водитель не видит раздел заявок (свои экраны забора); кладовщик
    смотрит карточку заявки без стоимости, но не список и не правку."""
    seller = _user("seller", "mvb_client", _mvb_client())
    _login(client, seller)
    order = _confirmed_order(client)
    driver, keeper = _user("drv", "mvb_driver"), _user("kpr", "mvb_storekeeper")
    _login(client, driver)
    for url in ("/mvb/orders", f"/mvb/orders/{order.id}", "/mvb/orders/new", f"/mvb/orders/{order.id}/labels.pdf"):
        assert client.get(url).status_code == 302, url
    assert client.post(f"/mvb/orders/{order.id}/cancel").status_code == 302
    assert db.session.get(MvbOrder, order.id).status == "confirmed"
    _login(client, keeper)
    assert client.get("/mvb/orders").status_code == 302
    page = client.get(f"/mvb/orders/{order.id}").get_data(as_text=True)
    assert order.number in page and "Стоимость:" not in page and "Отменить</button>" not in page


def test_partial_pickup_when_boxes_do_not_fit(db, client):
    """Все короба не влезли: водитель забирает часть, остаток возвращается
    в ленту забора и его берет другой водитель."""
    _login(client, _user("seller", "mvb_client", _mvb_client()))
    order = _confirmed_order(client, box_count="5")
    d1, d2 = _user("driver1", "mvb_driver"), _user("driver2", "mvb_driver")
    _login(client, d1)
    client.post(f"/mvb/driver/orders/{order.id}/take")
    # без сканов «часть» не закрыть
    client.post(f"/mvb/driver/orders/{order.id}/done", data={"partial": "1"})
    assert db.session.get(MvbOrder, order.id).driver_id == d1.id
    for box in order.boxes[:3]:
        client.post("/mvb/scan/pickup", data={"barcode": box.barcode})
    page = client.get(f"/mvb/scan/pickup?order={order.id}").get_data(as_text=True)
    assert "Не влезает — забрал часть" in page
    resp = client.post(f"/mvb/driver/orders/{order.id}/done", data={"partial": "1"})
    assert resp.headers["Location"].endswith("/mvb/driver")
    order = db.session.get(MvbOrder, order.id)
    assert order.driver_id is None and order.pickup_done_at is None
    assert sum(1 for b in order.boxes if b.status == "picked_up") == 3

    _login(client, d2)
    feed = client.get("/mvb/driver").get_data(as_text=True)
    assert order.number in feed or order.client.name in feed
    assert "2 (из 5)" in feed
    client.post(f"/mvb/driver/orders/{order.id}/take")
    for box in order.boxes[3:]:
        client.post("/mvb/scan/pickup", data={"barcode": box.barcode})
    client.post(f"/mvb/driver/orders/{order.id}/done")
    assert db.session.get(MvbOrder, order.id).pickup_done_at is not None
